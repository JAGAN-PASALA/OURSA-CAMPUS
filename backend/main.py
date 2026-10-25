import json
import hashlib
import hmac
import base64
import struct
import logging
import secrets
import time
import csv
import io
import zipfile
from collections import defaultdict, deque
from contextlib import asynccontextmanager
from datetime import date, datetime, timedelta, timezone
from pathlib import Path
from uuid import uuid4

import cv2
import numpy as np
from fastapi import Depends, FastAPI, File, Form, HTTPException, Request, Response, UploadFile, status
from fastapi.concurrency import run_in_threadpool
from fastapi.middleware.cors import CORSMiddleware
from fastapi.security import HTTPAuthorizationCredentials, HTTPBearer
from fastapi.responses import FileResponse, JSONResponse, RedirectResponse, StreamingResponse
from starlette.middleware.httpsredirect import HTTPSRedirectMiddleware
from starlette.middleware.trustedhost import TrustedHostMiddleware
from email_validator import EmailNotValidError, validate_email
from sqlalchemy import func, or_, select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

import database as db
import email_service
import payment_service
import storage_service
from distributed import fixed_window_allow, redis_ready, revoke_session, session_is_revoked, start_session, touch_session
from auth import create_token, decode_token, hash_password, verify_password
from config import get_settings
from ml_service import vision
from timetable_service import timetable_access, semester_dates, campus_today, register_routes as register_timetable_routes
from models import (
    AcademicTerm,
    Announcement,
    Assessment,
    AssessmentAttempt,
    AssessmentResult,
    AssessmentTopic,
    AuditLog,
    AttendanceCorrection,
    BehaviourModelVersion,
    Department,
    ClassSection,
    ClassSession,
    ClassMentor,
    Course,
    Enrollment,
    FaceEmbedding,
    FeeRecord,
    FeePayment,
    Incident,
    Institution,
    InstitutionDomain,
    LearningAssignment,
    LearningSubmission,
    LearningSubmissionAttachment,
    NotificationDelivery,
    OnDutyRequest,
    OnDutyReview,
    PaymentWebhookEvent,
    PasswordResetToken,
    PushDevice,
    Question,
    RaisedHandAttendanceRequest,
    SectionMembership,
    SessionAttendance,
    StaffCareerIdentity,
    StaffCareerServiceRecord,
    StaffFaceEmbedding,
    StaffProfile,
    StaffTransferRequest,
    SupportSession,
    Student,
    StudentAnswer,
    StudentProfile,
    StudentDocument,
    SyllabusTopic,
    SyllabusUnit,
    TeacherAssignment,
    TeacherFeedback,
    TimetableSlot,
    TopicCoverage,
    User,
    UserLegalAcceptance,
    ResultImportBatch,
)
from schemas import (
    AcademicTermCreate,
    AnnouncementCreate,
    AssessmentCreate,
    AssessmentUpdate,
    AssessmentStatusUpdate,
    AcademicTermUpdate,
    AttendanceCorrectionCreate,
    AttendanceOverride,
    AttemptAutosave,
    AttemptSubmit,
    CareerProfileUpdate,
    CareerTransferCreate,
    CareerTransferReview,
    BehaviourModelPromotion,
    ClassSectionCreate,
    ClassSessionCreate,
    ClassroomChangeCreate,
    ConsentUpdate,
    CourseCreate,
    CourseUpdate,
    DepartmentCreate,
    EnrollmentCreate,
    EnrollmentUpdate,
    FeeRecordUpsert,
    OnlinePaymentCreate,
    OfflinePaymentReview,
    InstitutionCreate,
    InstitutionLifecycleRequest,
    InstitutionUpdate,
    InstitutionMonitoringPolicy,
    LearningAssignmentCreate,
    LearningSubmissionCreate,
    LearningSubmissionGrade,
    LoginRequest,
    OfflineResultCreate,
    PasswordChange,
    PasswordResetConfirm,
    ForgotPasswordRequest,
    TemporaryPasswordReset,
    AdministratorCreate,
    PushDeviceCreate,
    PushPreferenceUpdate,
    QuestionCreate,
    RaisedHandReview,
    SectionMembershipCreate,
    StudentCreate,
    StudentBiometricConsentUpdate,
    StudentProfileUpdate,
    StudentProfilePermissionUpdate,
    OnDutyReviewUpdate,
    StudentUpdate,
    StaffCreate,
    SupportSessionCreate,
    SupportSessionReview,
    SyllabusTopicCreate,
    SyllabusUnitCreate,
    TeacherAssignmentCreate,
    TeacherFeedbackCreate,
    TeacherReplacement,
    TeacherCreate,
    TimetableSlotCreate,
    TimetableReview,
    TopicCoverageUpdate,
)


settings = get_settings()
logger = logging.getLogger(__name__)
security = HTTPBearer(auto_error=False)
login_windows: dict[str, deque[float]] = defaultdict(deque)
SUBMISSION_ROOT = Path(__file__).resolve().parent / "data" / "submissions"
PAYMENT_RECEIPT_ROOT = Path(__file__).resolve().parent / "data" / "payment_receipts"
OD_ATTACHMENT_ROOT = Path(__file__).resolve().parent / "data" / "on_duty"


@asynccontextmanager
async def lifespan(_: FastAPI):
    db.init_database()
    yield


app = FastAPI(
    title="OUR'S Smart Classroom API",
    version="8.2.0",
    docs_url="/docs" if not settings.is_production else None,
    redoc_url=None,
    openapi_url=None if settings.is_production else "/openapi.json",
    lifespan=lifespan,
)
app.add_middleware(TrustedHostMiddleware, allowed_hosts=settings.allowed_hosts)
if settings.is_production:
    app.add_middleware(HTTPSRedirectMiddleware)
app.add_middleware(
    CORSMiddleware,
    allow_origins=settings.allowed_origins + (settings.native_app_origins.split(",") if not settings.is_production else ["capacitor://localhost", "https://localhost"]),
    allow_credentials=True,
    allow_methods=["GET", "POST", "PATCH", "DELETE"],
    allow_headers=["Authorization", "Content-Type", "X-Request-ID", "X-Institution-Domain", "X-Institution-Slug", "X-Client-Platform"],
)


@app.middleware("http")
async def security_headers(request: Request, call_next):
    content_length = request.headers.get("content-length")
    if content_length:
        try:
            if int(content_length) > settings.max_upload_mb * 1024 * 1024:
                return JSONResponse(status_code=413, content={"detail": "Upload is too large"})
        except ValueError as exc:
            return JSONResponse(status_code=400, content={"detail": "Invalid Content-Length header"})
    request_id = request.headers.get("x-request-id") or str(uuid4())
    if request.method in {"POST", "PUT", "PATCH", "DELETE"} and request.cookies.get(settings.session_cookie_name) and not request.headers.get("authorization"):
        origin = request.headers.get("origin")
        allowed = {item.rstrip("/") for item in settings.allowed_origins}
        same_origin = f"{request.url.scheme}://{request.headers.get('host', '')}".rstrip("/")
        if (settings.is_production and not origin) or (origin and origin.rstrip("/") not in allowed and origin.rstrip("/") != same_origin):
            return JSONResponse(status_code=403, content={"detail": "The request origin is not allowed"}, headers={"X-Request-ID": request_id})
    response = await call_next(request)
    response.headers["X-Request-ID"] = request_id
    response.headers["X-Content-Type-Options"] = "nosniff"
    response.headers["X-Frame-Options"] = "DENY"
    response.headers["Referrer-Policy"] = "no-referrer"
    response.headers["Permissions-Policy"] = "camera=(self), microphone=(), geolocation=()"
    response.headers["Content-Security-Policy"] = "default-src 'none'; frame-ancestors 'none'; base-uri 'none'"
    if settings.is_production:
        response.headers["Strict-Transport-Security"] = "max-age=31536000; includeSubDomains"
    response.headers["Cache-Control"] = "no-store" if request.url.path.startswith("/api/") else "no-cache"
    return response


def request_domain(request: Request) -> str:
    if not settings.is_production:
        return request.headers.get("x-institution-domain") or request.headers.get("host", "")
    return request.headers.get("host", "")


def resolve_request_institution(request: Request, session: Session) -> tuple[Institution | None, str | None]:
    slug = request.headers.get("x-institution-slug")
    host = db.normalize_domain(request_domain(request))
    if settings.app_env.lower() == "test":
        if slug:
            return db.get_institution_by_slug(session, slug), None
        return db.get_institution_by_domain(session, host)
    if not settings.is_production:
        if slug:
            return db.get_institution_by_slug(session, slug), None
        return None, host
    if host in settings.portal_hosts:
        return (db.get_institution_by_slug(session, slug), None) if slug else (None, host)
    institution, domain = db.get_institution_by_domain(session, host)
    return institution, domain


def current_institution(request: Request, session: Session = Depends(db.get_db)) -> Institution:
    institution, _ = resolve_request_institution(request, session)
    if not institution:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "This domain is not assigned to an active institution")
    return institution


def current_user(
    request: Request,
    credentials: HTTPAuthorizationCredentials | None = Depends(security),
    session: Session = Depends(db.get_db),
) -> User:
    token = credentials.credentials if credentials and credentials.scheme.lower() == "bearer" else request.cookies.get(settings.session_cookie_name)
    if not token:
        raise HTTPException(status.HTTP_401_UNAUTHORIZED, "Sign in is required")
    payload = decode_token(token)
    if not payload or session_is_revoked(payload.get("jti")):
        raise HTTPException(status.HTTP_401_UNAUTHORIZED, "Your session is invalid or expired")
    if payload.get("jti") and not touch_session(payload["jti"], IDLE_SESSION_SECONDS):
        raise HTTPException(status.HTTP_401_UNAUTHORIZED, "Your session expired after inactivity. Sign in again.")
    user = db.get_user(session, int(payload["sub"]))
    if not user or not user.active or user.token_version != payload.get("ver"):
        raise HTTPException(status.HTTP_401_UNAUTHORIZED, "User account or session is no longer active")
    institution, _ = resolve_request_institution(request, session)
    if not institution or user.institution_id != institution.id or payload.get("institution_id") != institution.id:
        raise HTTPException(status.HTTP_401_UNAUTHORIZED, "This session belongs to a different institution")
    if settings.app_env.lower() != "test" and user.role == "admin" and not user.mfa_enabled and request.url.path not in {"/api/auth/mfa/setup", "/api/auth/mfa/confirm", "/api/auth/me", "/api/auth/logout"}:
        raise HTTPException(status.HTTP_428_PRECONDITION_REQUIRED, "Administrator two-step authentication setup is required")
    if settings.app_env.lower() != "test" and user.must_change_password and request.url.path not in {"/api/auth/me", "/api/auth/change-password", "/api/auth/logout", "/api/auth/mfa/setup", "/api/auth/mfa/confirm"}:
        raise HTTPException(status.HTTP_428_PRECONDITION_REQUIRED, "You must change the temporary password before using the application")
    return user


def teacher_only(user: User = Depends(current_user)) -> User:
    if user.role not in ("teacher", "admin") or user.platform_admin:
        raise HTTPException(status.HTTP_403_FORBIDDEN, "Teacher or administrator access is required")
    return user


def admin_only(user: User = Depends(current_user)) -> User:
    if user.role != "admin" or user.platform_admin:
        raise HTTPException(status.HTTP_403_FORBIDDEN, "Administrator access is required")
    return user


def platform_admin_only(user: User = Depends(current_user)) -> User:
    if not user.platform_admin:
        raise HTTPException(status.HTTP_403_FORBIDDEN, "Platform administrator access is required")
    return user


def allocation_manager(
    user: User = Depends(current_user),
    session: Session = Depends(db.get_db),
) -> User:
    """Allow college administrators and HODs to manage teaching allocations."""
    if user.platform_admin:
        raise HTTPException(status.HTTP_403_FORBIDDEN, "Platform accounts cannot modify college academic records")
    if user.role == "admin":
        return user
    profile = session.scalar(select(StaffProfile).where(
        StaffProfile.user_id == user.id,
        StaffProfile.institution_id == user.institution_id,
        StaffProfile.active.is_(True),
    ))
    if user.role == "teacher" and profile and profile.designation == "hod":
        return user
    raise HTTPException(status.HTTP_403_FORBIDDEN, "Institution administrator or HOD allocation access is required")


async def decode_image(image: UploadFile) -> np.ndarray:
    if image.content_type not in {"image/jpeg", "image/png", "image/webp"}:
        raise HTTPException(status.HTTP_415_UNSUPPORTED_MEDIA_TYPE, "Upload a JPEG, PNG or WebP image")
    payload = await image.read(settings.max_upload_mb * 1024 * 1024 + 1)
    if len(payload) > settings.max_upload_mb * 1024 * 1024:
        raise HTTPException(status.HTTP_413_REQUEST_ENTITY_TOO_LARGE, "Image exceeds the upload limit")
    frame = cv2.imdecode(np.frombuffer(payload, dtype=np.uint8), cv2.IMREAD_COLOR)
    if frame is None:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "The uploaded file is not a valid image")
    return frame


def commit_or_conflict(session: Session, message: str) -> None:
    try:
        session.commit()
    except IntegrityError as exc:
        session.rollback()
        raise HTTPException(status.HTTP_409_CONFLICT, message) from exc


def flush_or_conflict(session: Session, message: str) -> None:
    try:
        session.flush()
    except IntegrityError as exc:
        session.rollback()
        raise HTTPException(status.HTTP_409_CONFLICT, message) from exc


def department_for_student(session: Session, institution_id: int, department_id: int | None) -> Department | None:
    if not department_id:
        return None
    department = session.get(Department, department_id)
    if not department or department.institution_id != institution_id or not department.active:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Department not found for this institution")
    return department


def staff_profile_for(session: Session, user: User) -> StaffProfile:
    profile = session.scalar(select(StaffProfile).where(StaffProfile.user_id == user.id))
    if not profile:
        profile = StaffProfile(institution_id=user.institution_id, user_id=user.id, designation="teacher" if user.role == "teacher" else "administrator")
        session.add(profile)
        session.flush()
    return profile


def student_is_minor(profile: StudentProfile | None) -> bool | None:
    if not profile or not profile.date_of_birth:
        return None
    today = campus_today()
    years = today.year - profile.date_of_birth.year - ((today.month, today.day) < (profile.date_of_birth.month, profile.date_of_birth.day))
    return years < 18


def notify_section_students(
    session: Session,
    *,
    institution_id: int,
    section_id: int,
    term_id: int,
    template: str,
    message: str,
    dedup_prefix: str,
    deep_link: str,
) -> int:
    """Notify only active students in the exact class and term."""
    rows = session.execute(
        select(Student, User)
        .join(User, User.id == Student.user_id)
        .join(SectionMembership, SectionMembership.student_id == Student.id)
        .where(
            Student.institution_id == institution_id,
            SectionMembership.section_id == section_id,
            SectionMembership.term_id == term_id,
            SectionMembership.active.is_(True),
            User.active.is_(True),
        )
    ).all()
    for student, account in rows:
        db.add_notification(
            session,
            institution_id=institution_id,
            user_id=account.id,
            student_id=student.id,
            template=template,
            message=message,
            dedup_key=f"{dedup_prefix}:{student.id}",
            deep_link=deep_link,
        )
    return len(rows)


def notify_announcement_audience(session: Session, row: Announcement) -> int:
    roles = {
        "students": {"student"},
        "teachers": {"teacher"},
        "administrators": {"admin"},
        "all": {"student", "teacher", "admin"},
    }[row.audience]
    statement = (
        select(User, Student)
        .outerjoin(Student, Student.user_id == User.id)
        .where(User.institution_id == row.institution_id, User.role.in_(roles), User.active.is_(True))
    )
    if row.section_id:
        statement = (
            statement.join(SectionMembership, SectionMembership.student_id == Student.id)
            .where(
                User.role == "student",
                SectionMembership.section_id == row.section_id,
                SectionMembership.active.is_(True),
            )
        )
    recipients = session.execute(statement).all()
    for account, student in recipients:
        db.add_notification(
            session,
            institution_id=row.institution_id,
            user_id=account.id,
            student_id=student.id if student else None,
            template="announcement",
            message=f"{row.title}: {row.body[:240]}",
            dedup_key=f"announcement:{row.id}:{account.id}",
            deep_link="/portal?view=announcements",
            push_category="security" if row.audience == "administrators" else "attendance",
        )
    return len(recipients)


def ensure_section_enrollments(session: Session, student_id: int, section_id: int, term_id: int) -> int:
    course_ids = session.scalars(select(TeacherAssignment.course_id).where(TeacherAssignment.section_id == section_id, TeacherAssignment.term_id == term_id, TeacherAssignment.active.is_(True))).all()
    created = 0
    for course_id in set(course_ids):
        if not session.scalar(select(Enrollment.id).where(Enrollment.student_id == student_id, Enrollment.course_id == course_id)):
            session.add(Enrollment(student_id=student_id, course_id=course_id))
            created += 1
    return created


def normalized_student_email(raw_value: str, department: Department | None) -> str:
    value = raw_value.strip().lower()
    if "@" not in value:
        if not department:
            raise HTTPException(status.HTTP_400_BAD_REQUEST, "Choose a department or enter the student's full email address")
        if department.email_domain.endswith(".invalid"):
            raise HTTPException(status.HTTP_400_BAD_REQUEST, "Configure this existing department's real email domain, or enter the student's full email address")
        value = f"{value}@{department.email_domain}"
    try:
        return validate_email(value, check_deliverability=False).normalized.lower()
    except EmailNotValidError as exc:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, f"Invalid student email: {exc}") from exc


def ensure_unique_student_identity(session: Session, institution_id: int, name: str, email: str, roll_no: str, exclude_student_id: int | None = None) -> None:
    user_id_to_exclude = None
    if exclude_student_id:
        existing_student = session.get(Student, exclude_student_id)
        user_id_to_exclude = existing_student.user_id if existing_student else None
    email_query = select(User.id).where(User.institution_id == institution_id, func.lower(User.email) == email.lower())
    if user_id_to_exclude:
        email_query = email_query.where(User.id != user_id_to_exclude)
    if session.scalar(email_query):
        raise HTTPException(status.HTTP_409_CONFLICT, "This email address already exists in the institution")
    roll_query = select(Student.id).where(Student.institution_id == institution_id, func.lower(Student.roll_no) == roll_no.lower())
    if exclude_student_id:
        roll_query = roll_query.where(Student.id != exclude_student_id)
    if session.scalar(roll_query):
        raise HTTPException(status.HTTP_409_CONFLICT, "This register/roll number already exists")
    duplicate_query = select(Student.id).join(User, User.id == Student.user_id).where(Student.institution_id == institution_id, func.lower(Student.roll_no) == roll_no.lower(), func.lower(User.name) == name.lower())
    if exclude_student_id:
        duplicate_query = duplicate_query.where(Student.id != exclude_student_id)
    if session.scalar(duplicate_query):
        raise HTTPException(status.HTTP_409_CONFLICT, "A student with the same name and register number already exists")


def require_assignment_access(session: Session, actor: User, assignment_id: int) -> TeacherAssignment:
    assignment = session.get(TeacherAssignment, assignment_id)
    course = session.get(Course, assignment.course_id) if assignment else None
    section = session.get(ClassSection, assignment.section_id) if assignment else None
    if not assignment or not course or not section or course.institution_id != actor.institution_id or section.institution_id != actor.institution_id:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Teacher assignment not found")
    if actor.role != "admin" and (assignment.teacher_id != actor.id or not assignment.active):
        raise HTTPException(status.HTTP_403_FORBIDDEN, "This class is assigned to another teacher")
    return assignment


def require_assessment_access(session: Session, actor: User, assessment_id: int) -> Assessment:
    assessment = session.get(Assessment, assessment_id)
    course = session.get(Course, assessment.course_id) if assessment else None
    if not assessment or not course or course.institution_id != actor.institution_id:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Assessment not found")
    if actor.role != "admin":
        teaching = session.get(TeacherAssignment, assessment.teacher_assignment_id) if assessment.teacher_assignment_id else None
        if not teaching or teaching.teacher_id != actor.id or not teaching.active:
            raise HTTPException(status.HTTP_403_FORBIDDEN, "This assessment belongs to another teacher")
    return assessment


def require_coursework_access(session: Session, actor: User, coursework_id: int) -> LearningAssignment:
    coursework = session.get(LearningAssignment, coursework_id)
    course = session.get(Course, coursework.course_id) if coursework else None
    if not coursework or not course or course.institution_id != actor.institution_id:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Course assignment not found")
    if actor.role != "admin":
        teaching = session.get(TeacherAssignment, coursework.teacher_assignment_id)
        if not teaching or teaching.teacher_id != actor.id or not teaching.active:
            raise HTTPException(status.HTTP_403_FORBIDDEN, "This course assignment belongs to another teacher")
    return coursework


def ensure_timetable_has_no_conflict(session: Session, assignment: TeacherAssignment, payload: TimetableSlotCreate, exclude_slot_id: int | None = None) -> None:
    college_id = session.get(Course, assignment.course_id).institution_id
    session.scalar(select(Institution).where(Institution.id == college_id).with_for_update())
    overlapping = session.execute(
        select(TimetableSlot, TeacherAssignment)
        .join(TeacherAssignment, TeacherAssignment.id == TimetableSlot.assignment_id)
        .join(Course, Course.id == TeacherAssignment.course_id)
        .where(
            Course.institution_id == session.get(Course, assignment.course_id).institution_id,
            TimetableSlot.active.is_(True),
            TimetableSlot.publication_status.in_(["pending", "published"]),
            TimetableSlot.weekday == payload.weekday,
            TimetableSlot.effective_from <= payload.effective_to,
            TimetableSlot.effective_to >= payload.effective_from,
            TimetableSlot.start_time < payload.end_time,
            TimetableSlot.end_time > payload.start_time,
        )
    ).all()
    for slot, existing_assignment in overlapping:
        if exclude_slot_id and slot.id == exclude_slot_id:
            continue
        if slot.section_id == assignment.section_id:
            raise HTTPException(status.HTTP_409_CONFLICT, "This class section already has an overlapping period")
        if existing_assignment.teacher_id == assignment.teacher_id:
            raise HTTPException(status.HTTP_409_CONFLICT, "This teacher already has an overlapping period")
        if slot.classroom.strip().lower() == payload.classroom.strip().lower():
            raise HTTPException(status.HTTP_409_CONFLICT, "This classroom is already occupied during that time")


def generate_timetable_sessions(session: Session, slot: TimetableSlot, actor_id: int) -> int:
    generated = 0
    cursor = max(slot.effective_from, campus_today())
    final_date = slot.effective_to
    while cursor <= final_date:
        if cursor.weekday() == slot.weekday:
            existing = session.scalar(
                select(ClassSession).where(
                    ClassSession.section_id == slot.section_id,
                    ClassSession.session_date == cursor,
                    ClassSession.period_number == slot.period_number,
                )
            )
            if existing and existing.status == "cancelled" and not existing.actual_start and existing.attendance_state != "published":
                existing.assignment_id, existing.timetable_slot_id = slot.assignment_id, slot.id
                existing.classroom, existing.status = slot.classroom, "scheduled"
                existing.created_by = actor_id
                generated += 1
            elif not existing:
                session.add(ClassSession(assignment_id=slot.assignment_id, timetable_slot_id=slot.id, section_id=slot.section_id, session_date=cursor, period_number=slot.period_number, classroom=slot.classroom, created_by=actor_id))
                generated += 1
        cursor += timedelta(days=1)
    return generated


def aware_utc(value: datetime) -> datetime:
    return value if value.tzinfo else value.replace(tzinfo=timezone.utc)


def attempt_deadline(attempt: AssessmentAttempt, assessment: Assessment) -> datetime:
    deadline = aware_utc(attempt.started_at) + timedelta(minutes=assessment.duration_minutes)
    if assessment.available_until:
        deadline = min(deadline, aware_utc(assessment.available_until))
    return deadline


def score_assessment_attempt(session: Session, attempt: AssessmentAttempt, assessment: Assessment, student: Student, submitted: dict[int, str], user_id: int, timed_out: bool) -> dict:
    questions = {row.id: row for row in session.scalars(select(Question).where(Question.assessment_id == assessment.id)).all()}
    score = 0.0
    answer_review = []
    for question_id, question in questions.items():
        answer = submitted.get(question_id, "")
        is_correct = answer.strip().casefold() == (question.correct_answer or "").strip().casefold()
        awarded = question.marks if is_correct else 0
        score += awarded
        answer_row = session.scalar(select(StudentAnswer).where(StudentAnswer.attempt_id == attempt.id, StudentAnswer.question_id == question_id))
        if not answer_row:
            answer_row = StudentAnswer(attempt_id=attempt.id, question_id=question_id)
            session.add(answer_row)
        answer_row.answer_text = answer
        answer_row.awarded_marks = awarded
        answer_review.append({"question_id": question_id, "question_text": question.question_text, "selected_answer": answer, "correct_answer": question.correct_answer if assessment.show_correct_answers else None, "is_correct": is_correct, "awarded_marks": awarded, "marks": question.marks})
    attempt.submitted_at = datetime.now(timezone.utc)
    attempt.status = "auto_submitted" if timed_out else "scored"
    attempt.score = score
    result = session.scalar(select(AssessmentResult).where(AssessmentResult.assessment_id == assessment.id, AssessmentResult.student_id == student.id))
    if not result:
        result = AssessmentResult(assessment_id=assessment.id, student_id=student.id, marks=score, status=attempt.status, evaluated_by=assessment.created_by)
        session.add(result)
    result.marks = score
    result.status = attempt.status
    db.add_notification(session, institution_id=student.institution_id, user_id=student.user_id, student_id=student.id, template="assessment_result", message=f"Your result for {assessment.title} is {score:g}/{assessment.total_marks:g}.", dedup_key=f"assessment-result:{assessment.id}:{student.id}", deep_link="/portal?view=results")
    db.write_audit(session, user_id, "assessment.submitted", "assessment_attempt", attempt.id, {"score": score, "timed_out": timed_out})
    return {"attempt_id": attempt.id, "status": attempt.status, "score": score, "total_marks": assessment.total_marks, "timed_out": timed_out, "answers": answer_review}


@app.get("/api/health")
def health():
    return {"status": "ok", "service": settings.app_name, "version": "8.2.0"}


@app.get("/api/ready")
def ready():
    if not db.database_ready():
        raise HTTPException(status.HTTP_503_SERVICE_UNAVAILABLE, "Database is unavailable")
    dependencies = {"database": True, "redis": redis_ready(), "object_storage": storage_service.storage_ready()}
    if settings.is_production and not all(dependencies.values()):
        raise HTTPException(status.HTTP_503_SERVICE_UNAVAILABLE, {"message": "A required dependency is unavailable", "dependencies": dependencies})
    return {"status": "ready", "dependencies": dependencies}


@app.get("/api/public/tenant")
def public_tenant(request: Request, session: Session = Depends(db.get_db)):
    institution, domain = resolve_request_institution(request, session)
    if not institution:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "No institution is configured for this domain")
    return db.public_institution(institution, domain)


@app.get("/api/public/institutions")
def public_institutions(request: Request, q: str = "", session: Session = Depends(db.get_db)):
    ip = request.client.host if request.client else "unknown"
    if not fixed_window_allow(f"institution-search:{ip}", 30, 60):
        raise HTTPException(status.HTTP_429_TOO_MANY_REQUESTS, "Too many searches. Please wait and try again.")
    return db.search_public_institutions(session, q[:100])


@app.get("/api/platform/institutions")
def platform_institutions(_: User = Depends(platform_admin_only), session: Session = Depends(db.get_db)):
    return db.list_institutions(session)


@app.get("/api/public/institutions/{institution_id}/logo")
def public_institution_logo(institution_id: int, session: Session = Depends(db.get_db)):
    institution = session.get(Institution, institution_id)
    if not institution or not institution.active or not institution.logo_storage_key:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Institution logo was not found")
    local = storage_service.local_path(institution.logo_storage_key)
    if local:
        return FileResponse(local)
    return RedirectResponse(storage_service.download_url(institution.logo_storage_key, 300), status_code=307)


@app.post("/api/platform/institutions/{institution_id}/logo")
async def upload_institution_logo(institution_id: int, logo: UploadFile = File(...), actor: User = Depends(platform_admin_only), session: Session = Depends(db.get_db)):
    institution = session.get(Institution, institution_id)
    if not institution:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Institution was not found")
    if logo.content_type not in {"image/png", "image/jpeg", "image/webp"}:
        raise HTTPException(status.HTTP_415_UNSUPPORTED_MEDIA_TYPE, "Upload a PNG, JPEG or WebP logo")
    payload = await logo.read(2 * 1024 * 1024 + 1)
    if not payload or len(payload) > 2 * 1024 * 1024:
        raise HTTPException(status.HTTP_413_REQUEST_ENTITY_TOO_LARGE, "Logo must be smaller than 2 MB")
    key = f"institutions/{institution.id}/branding/logo-{uuid4().hex}{Path(logo.filename or '.png').suffix.lower()}"
    try:
        stored = await run_in_threadpool(storage_service.put_private, key, payload, logo.content_type)
    except ValueError as exc:
        raise HTTPException(status.HTTP_422_UNPROCESSABLE_ENTITY, str(exc)) from exc
    institution.logo_storage_key = stored.key
    db.write_audit(session, actor.id, "institution.logo_updated", "institution", institution.id, {"checksum": hashlib.sha256(payload).hexdigest()}, institution_id=institution.id)
    session.commit()
    return db.public_institution(institution)


@app.get("/api/platform/overview")
def platform_overview(actor: User = Depends(platform_admin_only), session: Session = Depends(db.get_db)):
    """Return platform-level operations data without exposing tenant records."""
    customer_filter = Institution.id != actor.institution_id
    institutions_total = session.scalar(select(func.count()).select_from(Institution).where(customer_filter)) or 0
    institutions_active = session.scalar(select(func.count()).select_from(Institution).where(customer_filter, Institution.active.is_(True))) or 0
    customer_users = session.scalar(
        select(func.count()).select_from(User).join(Institution, Institution.id == User.institution_id).where(customer_filter)
    ) or 0
    pending_support = session.scalar(
        select(func.count()).select_from(SupportSession).where(
            SupportSession.platform_user_id == actor.id,
            SupportSession.status == "pending",
            SupportSession.expires_at > datetime.now(timezone.utc),
        )
    ) or 0
    active_model = session.scalar(
        select(BehaviourModelVersion).where(BehaviourModelVersion.status == "active").order_by(BehaviourModelVersion.promoted_at.desc())
    )
    return {
        "institutions_total": institutions_total,
        "institutions_active": institutions_active,
        "institutions_archived": institutions_total - institutions_active,
        "customer_users": customer_users,
        "pending_support_requests": pending_support,
        "active_model": db.behaviour_model_version_payload(active_model) if active_model else None,
        "scope_notice": "Platform accounts manage institutions, platform operators, support consent and shared model releases. College records remain tenant-controlled.",
    }


@app.get("/api/platform/system-health")
def platform_system_health(_: User = Depends(platform_admin_only), session: Session = Depends(db.get_db)):
    meta = session.get(db.ModelMeta, 1)
    return {
        "api": "ready",
        "database": "ready" if db.database_ready() else "unavailable",
        "email": "configured" if settings.email_enabled else "not_configured",
        "payments": "configured" if settings.payment_enabled else "disabled",
        "sms_whatsapp": "configured" if settings.messaging_provider_enabled else "not_configured",
        "behaviour_model": "active" if meta and meta.trained else "safe_mode",
        "model_samples": meta.samples if meta else 0,
        "environment": settings.app_env,
    }


@app.post("/api/platform/institutions", status_code=status.HTTP_201_CREATED)
def create_institution(payload: InstitutionCreate, actor: User = Depends(platform_admin_only), session: Session = Depends(db.get_db)):
    if session.scalar(select(Institution.id).where(Institution.slug == payload.slug)) or session.scalar(select(InstitutionDomain.id).where(InstitutionDomain.domain == payload.domain)):
        raise HTTPException(status.HTTP_409_CONFLICT, "Institution slug or domain is already registered")
    institution = Institution(name=payload.name.strip(), slug=payload.slug, organization_type=payload.organization_type, brand_name=payload.brand_name.strip(), primary_color=payload.primary_color)
    session.add(institution)
    session.flush()
    session.add(InstitutionDomain(institution_id=institution.id, domain=payload.domain, is_primary=True))
    admin = User(institution_id=institution.id, name=payload.admin_name.strip(), email=str(payload.admin_email).lower(), password_hash=hash_password(payload.admin_password), role="admin", must_change_password=True)
    session.add(admin)
    flush_or_conflict(session, "Institution domain or administrator email is already registered")
    session.add(StaffProfile(institution_id=institution.id, user_id=admin.id, designation="administrator"))
    db.write_audit(session, actor.id, "institution.created", "institution", institution.id, {"name": institution.name, "domain": payload.domain}, institution_id=actor.institution_id)
    session.commit()
    return {**db.public_institution(institution, payload.domain), "admin_email": admin.email}


@app.patch("/api/platform/institutions/{institution_id}")
def update_platform_institution(institution_id: int, payload: InstitutionUpdate, actor: User = Depends(platform_admin_only), session: Session = Depends(db.get_db)):
    institution = session.get(Institution, institution_id)
    if not institution:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Institution not found")
    institution.name = payload.name.strip()
    institution.brand_name = payload.brand_name.strip()
    institution.primary_color = payload.primary_color
    db.write_audit(session, actor.id, "institution.updated", "institution", institution.id, payload.model_dump())
    session.commit()
    return db.public_institution(institution)


@app.post("/api/platform/institutions/{institution_id}/administrators/{administrator_id}/reset-password")
def reset_customer_administrator(institution_id: int, administrator_id: int, payload: TemporaryPasswordReset, actor: User = Depends(platform_admin_only), session: Session = Depends(db.get_db)):
    administrator = session.get(User, administrator_id)
    if not administrator or administrator.institution_id != institution_id or administrator.role != "admin" or administrator.platform_admin:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Institution administrator not found")
    administrator.password_hash = hash_password(payload.temporary_password)
    administrator.must_change_password = True
    administrator.failed_login_count = 0
    administrator.locked_until = None
    administrator.token_version += 1
    db.write_audit(session, actor.id, "institution_admin.password_reset", "user", administrator.id, {"institution_id": institution_id})
    session.commit()
    return {"message": "Temporary password issued; all previous sessions were signed out"}


@app.get("/api/platform/team")
def platform_team(actor: User = Depends(platform_admin_only), session: Session = Depends(db.get_db)):
    rows = session.scalars(select(User).where(User.institution_id == actor.institution_id, User.role == "admin", User.platform_admin.is_(True), User.active.is_(True)).order_by(User.name)).all()
    return [db.public_user(row, session) for row in rows]


@app.get("/api/platform/audit-logs")
def platform_audit_logs(limit: int = 100, actor: User = Depends(platform_admin_only), session: Session = Depends(db.get_db)):
    return db.list_audit_logs(session, actor.institution_id, limit)


@app.post("/api/platform/team", status_code=status.HTTP_201_CREATED)
def create_platform_team_member(payload: AdministratorCreate, actor: User = Depends(platform_admin_only), session: Session = Depends(db.get_db)):
    member = User(institution_id=actor.institution_id, name=payload.name.strip(), email=str(payload.email).lower(), password_hash=hash_password(payload.password), role="admin", platform_admin=True, must_change_password=True)
    session.add(member)
    flush_or_conflict(session, "A platform team member already uses this email")
    session.add(StaffProfile(institution_id=actor.institution_id, user_id=member.id, designation="administrator"))
    db.write_audit(session, actor.id, "platform_team.created", "user", member.id)
    session.commit()
    return db.public_user(member)


def support_session_payload(session: Session, row: SupportSession) -> dict:
    now = datetime.now(timezone.utc)
    expires_at = aware_utc(row.expires_at)
    status_value = "expired" if row.status in {"pending", "approved"} and expires_at <= now else row.status
    institution = session.get(Institution, row.target_institution_id)
    administrator = session.get(User, row.represented_admin_id) if row.represented_admin_id else None
    platform_user = session.get(User, row.platform_user_id)
    return {
        "id": row.id,
        "institution_id": row.target_institution_id,
        "institution_name": institution.name if institution else "Unavailable institution",
        "represented_admin_id": row.represented_admin_id,
        "represented_admin_name": administrator.name if administrator else None,
        "platform_user_name": platform_user.name if platform_user else "Platform support",
        "reason": row.reason,
        "ticket_reference": row.ticket_reference,
        "status": status_value,
        "read_only": row.read_only,
        "requested_at": row.requested_at.isoformat(),
        "approved_at": row.approved_at.isoformat() if row.approved_at else None,
        "expires_at": row.expires_at.isoformat(),
        "revoked_at": row.revoked_at.isoformat() if row.revoked_at else None,
    }


@app.post("/api/platform/support-sessions", status_code=status.HTTP_201_CREATED)
def request_support_session(payload: SupportSessionCreate, actor: User = Depends(platform_admin_only), session: Session = Depends(db.get_db)):
    institution = session.get(Institution, payload.institution_id)
    if not institution or institution.id == actor.institution_id:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Customer institution not found")
    administrator = session.get(User, payload.represented_admin_id) if payload.represented_admin_id else session.scalar(
        select(User).where(User.institution_id == institution.id, User.role == "admin", User.active.is_(True)).order_by(User.id)
    )
    if not administrator or administrator.institution_id != institution.id or administrator.role != "admin" or administrator.platform_admin:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Active college administrator not found")
    row = SupportSession(
        platform_user_id=actor.id,
        target_institution_id=institution.id,
        represented_admin_id=administrator.id,
        reason=payload.reason.strip(),
        ticket_reference=(payload.ticket_reference or "").strip() or None,
        expires_at=datetime.now(timezone.utc) + timedelta(minutes=payload.duration_minutes),
        read_only=True,
    )
    session.add(row)
    session.flush()
    db.write_audit(session, actor.id, "support.requested", "support_session", row.id, {"target_institution_id": institution.id, "reason": row.reason, "ticket_reference": row.ticket_reference}, institution_id=actor.institution_id)
    db.write_audit(session, actor.id, "support.request_received", "support_session", row.id, {"platform_actor_id": actor.id, "represented_admin_id": administrator.id, "reason": row.reason}, institution_id=institution.id)
    session.commit()
    return support_session_payload(session, row)


@app.get("/api/platform/support-sessions")
def platform_support_sessions(actor: User = Depends(platform_admin_only), session: Session = Depends(db.get_db)):
    rows = session.scalars(select(SupportSession).where(SupportSession.platform_user_id == actor.id).order_by(SupportSession.requested_at.desc())).all()
    return [support_session_payload(session, row) for row in rows]


@app.get("/api/admin/support-sessions")
def college_support_sessions(actor: User = Depends(admin_only), session: Session = Depends(db.get_db)):
    rows = session.scalars(select(SupportSession).where(SupportSession.target_institution_id == actor.institution_id).order_by(SupportSession.requested_at.desc())).all()
    return [support_session_payload(session, row) for row in rows]


@app.patch("/api/admin/support-sessions/{support_session_id}/review")
def review_support_session(support_session_id: int, payload: SupportSessionReview, actor: User = Depends(admin_only), session: Session = Depends(db.get_db)):
    row = session.get(SupportSession, support_session_id)
    if not row or row.target_institution_id != actor.institution_id:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Support request not found")
    if row.represented_admin_id != actor.id:
        raise HTTPException(status.HTTP_403_FORBIDDEN, "Only the named college administrator can approve or reject this support request")
    if aware_utc(row.expires_at) <= datetime.now(timezone.utc):
        row.status = "expired"
        session.commit()
        raise HTTPException(status.HTTP_409_CONFLICT, "This support request has expired")
    if row.status not in {"pending", "approved"}:
        raise HTTPException(status.HTTP_409_CONFLICT, "This support request is already closed")
    if payload.decision == "approved":
        row.status = "approved"
        row.approved_by = actor.id
        row.approved_at = datetime.now(timezone.utc)
    elif payload.decision == "rejected":
        row.status = "rejected"
        row.revoked_at = datetime.now(timezone.utc)
    else:
        row.status = "revoked"
        row.revoked_at = datetime.now(timezone.utc)
    db.write_audit(session, actor.id, f"support.{row.status}", "support_session", row.id, {"platform_actor_id": row.platform_user_id, "reason": payload.reason})
    session.commit()
    return support_session_payload(session, row)


@app.get("/api/platform/support-sessions/{support_session_id}/administration")
def open_read_only_support_administration(support_session_id: int, actor: User = Depends(platform_admin_only), session: Session = Depends(db.get_db)):
    row = session.get(SupportSession, support_session_id)
    if not row or row.platform_user_id != actor.id:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Support session not found")
    if row.status != "approved" or row.revoked_at or aware_utc(row.expires_at) <= datetime.now(timezone.utc):
        raise HTTPException(status.HTTP_403_FORBIDDEN, "Approved, unexpired college consent is required")
    institution = session.get(Institution, row.target_institution_id)
    if not institution or not institution.active:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Customer institution is unavailable")
    counts = {
        "users": session.scalar(select(func.count()).select_from(User).where(User.institution_id == institution.id)) or 0,
        "students": session.scalar(select(func.count()).select_from(Student).where(Student.institution_id == institution.id)) or 0,
        "staff": session.scalar(select(func.count()).select_from(User).where(User.institution_id == institution.id, User.role.in_(["teacher", "staff"]))) or 0,
        "departments": session.scalar(select(func.count()).select_from(Department).where(Department.institution_id == institution.id, Department.active.is_(True))) or 0,
        "classes": session.scalar(select(func.count()).select_from(ClassSection).where(ClassSection.institution_id == institution.id, ClassSection.active.is_(True))) or 0,
        "courses": session.scalar(select(func.count()).select_from(Course).where(Course.institution_id == institution.id, Course.active.is_(True))) or 0,
    }
    db.write_audit(session, actor.id, "support.read_only_administration_opened", "support_session", row.id, {"real_platform_actor_id": actor.id, "represented_admin_id": row.represented_admin_id, "read_only": True}, institution_id=institution.id)
    session.commit()
    return {
        "support": support_session_payload(session, row),
        "institution": db.public_institution(institution),
        "counts": counts,
        "administrators": [db.public_user(admin, session) for admin in session.scalars(select(User).where(User.institution_id == institution.id, User.role == "admin", User.active.is_(True)).order_by(User.name)).all()],
        "recent_audit": db.list_audit_logs(session, institution.id, 20),
        "notice": "Read-only support view. Credentials, biometric templates, payment changes, exports, and destructive actions are blocked.",
    }


@app.post("/api/platform/support-sessions/{support_session_id}/revoke")
def revoke_support_session(support_session_id: int, actor: User = Depends(platform_admin_only), session: Session = Depends(db.get_db)):
    row = session.get(SupportSession, support_session_id)
    if not row or row.platform_user_id != actor.id:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Support session not found")
    row.status = "revoked"
    row.revoked_at = datetime.now(timezone.utc)
    db.write_audit(session, actor.id, "support.revoked", "support_session", row.id, {"target_institution_id": row.target_institution_id})
    session.commit()
    return support_session_payload(session, row)


@app.post("/api/platform/institutions/{institution_id}/lifecycle")
def institution_lifecycle(institution_id: int, payload: InstitutionLifecycleRequest, actor: User = Depends(platform_admin_only), session: Session = Depends(db.get_db)):
    institution = session.get(Institution, institution_id)
    if not institution or institution.id == actor.institution_id:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Customer institution not found")
    if payload.typed_name.strip() != institution.name:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "Type the exact institution name to confirm")
    if not verify_password(payload.current_password, actor.password_hash):
        raise HTTPException(status.HTTP_403_FORBIDDEN, "Recent platform-owner password verification failed")
    if payload.action == "restore":
        if institution.active:
            raise HTTPException(status.HTTP_409_CONFLICT, "This institution is already active")
        institution.active = True
        institution.archived_at = None
        institution.archived_by = None
        session.query(InstitutionDomain).filter(InstitutionDomain.institution_id == institution.id).update({InstitutionDomain.active: True}, synchronize_session=False)
    elif payload.action == "archive":
        if not institution.active:
            raise HTTPException(status.HTTP_409_CONFLICT, "This institution is already archived")
        institution.active = False
        institution.archived_at = datetime.now(timezone.utc)
        institution.archived_by = actor.id
        session.query(InstitutionDomain).filter(InstitutionDomain.institution_id == institution.id).update({InstitutionDomain.active: False}, synchronize_session=False)
        session.query(User).filter(User.institution_id == institution.id).update({User.token_version: User.token_version + 1}, synchronize_session=False)
    else:
        if institution.active:
            raise HTTPException(status.HTTP_409_CONFLICT, "Archive the institution before permanent deletion")
        total_users = session.scalar(select(func.count()).select_from(User).where(User.institution_id == institution.id)) or 0
        protected_counts = {
            "additional_users": max(0, total_users - 1),
            "students": session.scalar(select(func.count()).select_from(Student).where(Student.institution_id == institution.id)) or 0,
            "departments": session.scalar(select(func.count()).select_from(Department).where(Department.institution_id == institution.id)) or 0,
            "courses": session.scalar(select(func.count()).select_from(Course).where(Course.institution_id == institution.id)) or 0,
            "classes": session.scalar(select(func.count()).select_from(ClassSection).where(ClassSection.institution_id == institution.id)) or 0,
            "fees": session.scalar(select(func.count()).select_from(FeeRecord).where(FeeRecord.institution_id == institution.id)) or 0,
            "assessments": session.scalar(select(func.count()).select_from(Assessment).join(Course, Course.id == Assessment.course_id).where(Course.institution_id == institution.id)) or 0,
            "announcements": session.scalar(select(func.count()).select_from(Announcement).where(Announcement.institution_id == institution.id)) or 0,
            "incidents": session.scalar(select(func.count()).select_from(Incident).where(Incident.institution_id == institution.id)) or 0,
            "support_sessions": session.scalar(select(func.count()).select_from(SupportSession).where(SupportSession.target_institution_id == institution.id)) or 0,
            "audit_events": session.scalar(select(func.count()).select_from(AuditLog).where(AuditLog.institution_id == institution.id)) or 0,
        }
        if any(protected_counts.values()):
            raise HTTPException(status.HTTP_409_CONFLICT, {"message": "A non-empty institution must be archived, not permanently deleted", "dependencies": protected_counts})
        if not payload.backup_confirmed:
            raise HTTPException(status.HTTP_409_CONFLICT, "Confirm the backup check before permanent deletion")
        name = institution.name
        session.delete(institution)
        db.write_audit(session, actor.id, "institution.permanently_deleted", "institution", institution_id, {"name": name, "reason": payload.reason, "backup_confirmed": True}, institution_id=actor.institution_id)
        session.commit()
        return {"id": institution_id, "status": "deleted"}
    db.write_audit(session, actor.id, f"institution.{payload.action}d", "institution", institution.id, {"reason": payload.reason}, institution_id=actor.institution_id)
    session.commit()
    return {**db.public_institution(institution), "active": institution.active, "archived_at": institution.archived_at.isoformat() if institution.archived_at else None}


LEGAL_DOCUMENT_VERSION = "2026-10-10-v1"
IDLE_SESSION_SECONDS = 30 * 60

def _mfa_fernet():
    from cryptography.fernet import Fernet
    if settings.mfa_encryption_key:
        return Fernet(settings.mfa_encryption_key.encode())
    # Development-only fallback. Production startup requires a dedicated key.
    key = base64.urlsafe_b64encode(hashlib.sha256((settings.app_secret + ":ours-mfa-secret-v1").encode()).digest())
    return Fernet(key)

def _totp(secret: str, at_time: int | None = None) -> tuple[str, int]:
    counter = int((at_time if at_time is not None else time.time()) // 30)
    key = base64.b32decode(secret.upper() + "=" * ((8 - len(secret) % 8) % 8))
    digest = hmac.new(key, struct.pack(">Q", counter), hashlib.sha1).digest()
    offset = digest[-1] & 0x0F
    value = (struct.unpack(">I", digest[offset:offset + 4])[0] & 0x7fffffff) % 1_000_000
    return f"{value:06d}", counter

def _verify_totp(user: User, supplied: str | None) -> bool:
    if not supplied or not user.mfa_secret_encrypted:
        return False
    try:
        secret = _mfa_fernet().decrypt(user.mfa_secret_encrypted.encode()).decode()
    except Exception:
        return False
    now_counter = int(time.time() // 30)
    for counter in (now_counter - 1, now_counter, now_counter + 1):
        code, _ = _totp(secret, counter * 30)
        if hmac.compare_digest(code, supplied.strip()) and (user.mfa_last_counter is None or counter > user.mfa_last_counter):
            user.mfa_last_counter = counter
            return True
    return False

def _store_terms_acceptance(session: Session, user: User, client_ip: str | None) -> None:
    exists = session.scalar(select(UserLegalAcceptance.id).where(UserLegalAcceptance.user_id == user.id, UserLegalAcceptance.document_version == LEGAL_DOCUMENT_VERSION))
    if not exists:
        session.add(UserLegalAcceptance(user_id=user.id, institution_id=user.institution_id, document_version=LEGAL_DOCUMENT_VERSION, accepted_at=datetime.now(timezone.utc)))

@app.get("/api/legal/terms")
def legal_terms():
    return {"version": LEGAL_DOCUMENT_VERSION, "title": "OUR’S Smart Classroom Terms of Use", "notice": "Use of this service is subject to your institution’s agreement and applicable law. The service may process account details, academic records, attendance, assignment files, camera-derived data and optional biometric data to provide enabled features. Automated detections can be wrong and must be reviewable. Do not upload unlawful, harmful or third-party data without authority. Access is role-based; misuse, unauthorized access and attempts to bypass controls are prohibited. Data is retained according to the institution’s configured retention policy and applicable law. Contact your institution administrator or designated privacy/grievance contact for access, correction, deletion, consent withdrawal where applicable, or complaints. Optional camera, biometric, notification and device permissions are handled separately. Acceptance does not replace any separate consent required by law. This notice is a deployment template and requires legal review before public launch.", "requires_legal_review": True}

@app.get("/api/legal/privacy")
def legal_privacy():
    return {"version": LEGAL_DOCUMENT_VERSION, "title": "OUR’S Smart Classroom Privacy Notice", "notice": "The platform may process account identifiers, institution and academic records, attendance, submitted files, security logs, and camera-derived information when the institution enables those functions. Face recognition and behaviour analysis require separate appropriate consent and safeguards. Data may be processed by hosting, storage, email, messaging and security providers under applicable arrangements. Access is restricted by role and institution. Retention, rights requests, grievance contacts, processor details and jurisdiction-specific disclosures must be completed by the deploying institution before production. This is a template notice, not a legal compliance guarantee.", "requires_legal_review": True}

@app.post("/api/auth/mfa/setup")
def mfa_setup(user: User = Depends(current_user), session: Session = Depends(db.get_db)):
    if user.role != "admin":
        raise HTTPException(status.HTTP_403_FORBIDDEN, "MFA enrollment endpoint is reserved for administrators")
    if user.mfa_enabled:
        raise HTTPException(status.HTTP_409_CONFLICT, "MFA is already enabled")
    secret = base64.b32encode(secrets.token_bytes(20)).decode().rstrip("=")
    user.mfa_secret_encrypted = _mfa_fernet().encrypt(secret.encode()).decode()
    session.commit()
    label = f"OURS:{user.email}"
    return {"secret": secret, "otpauth_uri": f"otpauth://totp/{label}?secret={secret}&issuer=OURS%20Smart%20Classroom&digits=6&period=30", "message": "Add this secret to an authenticator app, then confirm with the current six-digit code. Do not share the secret."}

@app.post("/api/auth/mfa/confirm")
def mfa_confirm(payload: dict, request: Request, response: Response, user: User = Depends(current_user), session: Session = Depends(db.get_db)):
    ip = request.client.host if request.client else "unknown"
    if not fixed_window_allow(f"mfa-confirm:{user.id}:{ip}", 5, 300):
        raise HTTPException(status.HTTP_429_TOO_MANY_REQUESTS, "Too many authenticator attempts. Try again later.")
    code = str(payload.get("code", ""))[:8]
    if user.role != "admin" or not user.mfa_secret_encrypted:
        raise HTTPException(status.HTTP_409_CONFLICT, "Start MFA setup first")
    try:
        secret = _mfa_fernet().decrypt(user.mfa_secret_encrypted.encode()).decode()
    except Exception as exc:
        raise HTTPException(status.HTTP_503_SERVICE_UNAVAILABLE, "MFA setup could not be verified") from exc
    matched = None
    for counter in (int(time.time() // 30) - 1, int(time.time() // 30), int(time.time() // 30) + 1):
        expected, _ = _totp(secret, counter * 30)
        if hmac.compare_digest(expected, code):
            matched = counter
    if matched is None:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "Invalid authenticator code")
    user.mfa_enabled = True
    user.mfa_last_counter = matched
    # Invalidate the short-lived setup token and issue a fresh full session.
    user.token_version += 1
    db.write_audit(session, user.id, "auth.mfa_enabled", "user", user.id)
    session.commit()
    token = create_token(user)
    token_payload = decode_token(token) or {}
    if token_payload.get("jti") and not start_session(token_payload["jti"], settings.access_token_minutes * 60):
        raise HTTPException(status.HTTP_503_SERVICE_UNAVAILABLE, "Secure session service is unavailable")
    response.set_cookie(settings.session_cookie_name, token, max_age=settings.access_token_minutes * 60, httponly=True, secure=settings.is_production, samesite="lax", path="/api", domain=settings.session_cookie_domain)
    return {"mfa_enabled": True, "user": db.public_user(user, session), "message": "Two-step authentication enabled. Keep your authenticator available for future logins."}

@app.post("/api/auth/login")
def login(request: LoginRequest, http_request: Request, response: Response, institution: Institution = Depends(current_institution), session: Session = Depends(db.get_db)):
    client = http_request.client.host if http_request.client else "unknown"
    rate_key = f"{client}:{institution.id}"
    if not fixed_window_allow(f"login:{rate_key}", settings.login_attempts_per_minute, 60):
        raise HTTPException(status.HTTP_429_TOO_MANY_REQUESTS, "Too many sign-in attempts. Wait one minute.")
    if not request.terms_accepted:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "Read and accept the Terms of Use and Privacy Notice before signing in")
    now = time.monotonic()
    attempts = login_windows[rate_key]
    while attempts and now - attempts[0] > 60:
        attempts.popleft()
    if len(attempts) >= settings.login_attempts_per_minute:
        raise HTTPException(status.HTTP_429_TOO_MANY_REQUESTS, "Too many sign-in attempts. Wait one minute.")
    attempts.append(now)
    user = db.get_user_by_email(session, str(request.email), institution.id)
    selected_role_matches = bool(user) and user.role == request.role
    password_ok = bool(user and verify_password(request.password, user.password_hash))
    if not user or not user.active or not selected_role_matches or not password_ok:
        if user:
            # Do not hard-lock an account based only on attacker-controlled failures.
            # IP/distributed throttles and edge abuse controls protect against guessing.
            user.failed_login_count = min(int(user.failed_login_count or 0) + 1, 1_000_000)
            session.commit()
        raise HTTPException(status.HTTP_401_UNAUTHORIZED, "Incorrect email or password")
    attempts.clear()
    user.failed_login_count = 0
    user.locked_until = None
    _store_terms_acceptance(session, user, client)
    if settings.app_env.lower() != "test" and user.role == "admin" and not user.mfa_enabled:
        token = create_token(user)
        token_payload = decode_token(token) or {}
        if token_payload.get("jti") and not start_session(token_payload["jti"], 10 * 60):
            raise HTTPException(status.HTTP_503_SERVICE_UNAVAILABLE, "Secure session service is unavailable")
        response.set_cookie(settings.session_cookie_name, token, max_age=10 * 60, httponly=True, secure=settings.is_production, samesite="lax", path="/api", domain=settings.session_cookie_domain)
        db.write_audit(session, user.id, "auth.admin_mfa_setup_required", "user", user.id, ip_address=client)
        session.commit()
        return {"mfa_setup_required": True, "user": db.public_user(user, session), "message": "Administrator accounts must enable two-step authentication before using the app."}
    if user.role == "admin" and user.mfa_enabled:
        if not _verify_totp(user, request.mfa_code):
            session.commit()
            return JSONResponse(status_code=202, content={"mfa_required": True, "message": "Enter the current code from your authenticator app."})
    db.write_audit(session, user.id, "auth.login", "user", user.id, ip_address=client)
    session.commit()
    token = create_token(user)
    token_payload = decode_token(token) or {}
    if token_payload.get("jti") and not start_session(token_payload["jti"], settings.access_token_minutes * 60):
        raise HTTPException(status.HTTP_503_SERVICE_UNAVAILABLE, "Secure session service is unavailable")
    response.set_cookie(settings.session_cookie_name, token, max_age=settings.access_token_minutes * 60, httponly=True, secure=settings.is_production, samesite="lax", path="/api", domain=settings.session_cookie_domain)
    result = {"user": db.public_user(user, session), "expires_in": settings.access_token_minutes * 60}
    if http_request.headers.get("x-client-platform") in {"android", "ios"}:
        result["access_token"] = token
    return result


@app.post("/api/auth/forgot-password")
def forgot_password(payload: ForgotPasswordRequest, request: Request, institution: Institution = Depends(current_institution), session: Session = Depends(db.get_db)):
    ip = request.client.host if request.client else "unknown"
    if not fixed_window_allow(f"forgot-password:{ip}:{institution.id}", 5, 3600):
        return {"message": "If that account exists, a password-reset link has been sent. Contact your institution administrator if email delivery is not configured."}
    user = db.get_user_by_email(session, str(payload.email), institution.id)
    if user and user.active and user.role == payload.role:
        raw_token = secrets.token_urlsafe(48)
        token_hash = hashlib.sha256(raw_token.encode()).hexdigest()
        session.add(PasswordResetToken(user_id=user.id, token_hash=token_hash, expires_at=datetime.now(timezone.utc) + timedelta(minutes=30)))
        delivered, error = email_service.send_password_reset(user.email, user.name, institution.name, institution.slug, raw_token)
        session.add(NotificationDelivery(institution_id=institution.id, user_id=user.id, channel="email", destination_masked=f"***{user.email[-6:]}", template="password_reset", status="sent" if delivered else "failed", error=error, sent_at=datetime.now(timezone.utc) if delivered else None))
        db.write_audit(session, user.id, "auth.password_reset_requested", "user", user.id, {"email_delivery": "sent" if delivered else "failed"})
        session.commit()
    return {"message": "If that account exists, a password-reset link has been sent. Contact your institution administrator if email delivery is not configured."}


@app.post("/api/auth/reset-password")
def reset_password(payload: PasswordResetConfirm, request: Request, institution: Institution = Depends(current_institution), session: Session = Depends(db.get_db)):
    ip = request.client.host if request.client else "unknown"
    if not fixed_window_allow(f"reset-password:{ip}:{institution.id}", 10, 3600):
        raise HTTPException(status.HTTP_429_TOO_MANY_REQUESTS, "Too many password-reset attempts. Try again later.")
    token_hash = hashlib.sha256(payload.token.encode()).hexdigest()
    row = session.scalar(select(PasswordResetToken).where(PasswordResetToken.token_hash == token_hash, PasswordResetToken.used_at.is_(None)))
    user = session.get(User, row.user_id) if row else None
    if not row or not user or user.institution_id != institution.id:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "Reset link is invalid or already used")
    expiry = row.expires_at if row.expires_at.tzinfo else row.expires_at.replace(tzinfo=timezone.utc)
    if expiry < datetime.now(timezone.utc):
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "Reset link has expired")
    user.password_hash = hash_password(payload.new_password)
    user.must_change_password = False
    user.failed_login_count = 0
    user.locked_until = None
    user.token_version += 1
    row.used_at = datetime.now(timezone.utc)
    db.write_audit(session, user.id, "auth.password_reset", "user", user.id)
    session.commit()
    return {"message": "Password reset successfully. You can now sign in."}


@app.get("/api/auth/me")
def authenticated_user(user: User = Depends(current_user), session: Session = Depends(db.get_db)):
    return {"user": db.public_user(user, session)}


def require_student_document_access(session: Session, actor: User, student_id: int) -> Student:
    student = session.get(Student, student_id)
    if not student or student.institution_id != actor.institution_id:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Student was not found")
    if actor.role == "student":
        if student.user_id != actor.id:
            raise HTTPException(status.HTTP_403_FORBIDDEN, "You can access only your own documents")
        return student
    if actor.role == "admin":
        return student
    profile = session.scalar(select(StaffProfile).where(StaffProfile.user_id == actor.id, StaffProfile.active.is_(True)))
    assigned = session.scalar(
        select(TeacherAssignment.id)
        .join(SectionMembership, SectionMembership.section_id == TeacherAssignment.section_id)
        .where(TeacherAssignment.teacher_id == actor.id, TeacherAssignment.active.is_(True), SectionMembership.student_id == student.id, SectionMembership.active.is_(True))
        .limit(1)
    )
    if assigned or (profile and profile.designation in {"hod", "dean"} and (not profile.department_id or profile.department_id == student.department_id)):
        return student
    raise HTTPException(status.HTTP_403_FORBIDDEN, "Only assigned teachers, HODs, deans and administrators can access this document")


@app.get("/api/students/{student_id}/documents")
def list_student_documents(student_id: int, actor: User = Depends(current_user), session: Session = Depends(db.get_db)):
    require_student_document_access(session, actor, student_id)
    rows = session.scalars(select(StudentDocument).where(StudentDocument.student_id == student_id, StudentDocument.institution_id == actor.institution_id).order_by(StudentDocument.created_at.desc()).limit(200)).all()
    return [{"id": row.id, "category": row.category, "title": row.title, "file_name": row.file_name, "mime_type": row.mime_type, "size_bytes": row.size_bytes, "scan_status": row.scan_status, "created_at": db.campus_datetime_iso(row.created_at)} for row in rows]


@app.post("/api/students/{student_id}/documents", status_code=status.HTTP_201_CREATED)
async def upload_student_document(student_id: int, title: str = Form(...), category: str = Form(...), file: UploadFile = File(...), actor: User = Depends(current_user), session: Session = Depends(db.get_db)):
    require_student_document_access(session, actor, student_id)
    if not fixed_window_allow(f"upload:student-document:{actor.id}", 20, 3600):
        raise HTTPException(status.HTTP_429_TOO_MANY_REQUESTS, "Too many document uploads. Try again later.")
    allowed_types = {"application/pdf", "image/jpeg", "image/png"}
    if file.content_type not in allowed_types:
        raise HTTPException(status.HTTP_415_UNSUPPORTED_MEDIA_TYPE, "Upload a PDF, JPEG or PNG document")
    payload = await file.read(settings.max_upload_mb * 1024 * 1024 + 1)
    if not payload or len(payload) > settings.max_upload_mb * 1024 * 1024:
        raise HTTPException(status.HTTP_413_REQUEST_ENTITY_TOO_LARGE, "Document is empty or exceeds the upload limit")
    checksum = hashlib.sha256(payload).hexdigest()
    if session.scalar(select(StudentDocument.id).where(StudentDocument.student_id == student_id, StudentDocument.checksum_sha256 == checksum)):
        raise HTTPException(status.HTTP_409_CONFLICT, "This student document was already uploaded")
    safe_name = Path(file.filename or "document").name
    key = f"institutions/{actor.institution_id}/students/{student_id}/{uuid4().hex}-{safe_name}"
    try:
        stored = await run_in_threadpool(storage_service.put_private, key, payload, file.content_type)
    except ValueError as exc:
        raise HTTPException(status.HTTP_422_UNPROCESSABLE_ENTITY, str(exc)) from exc
    row = StudentDocument(institution_id=actor.institution_id, student_id=student_id, category=category.strip()[:60], title=title.strip()[:180], file_name=safe_name, storage_key=stored.key, mime_type=file.content_type, size_bytes=stored.size, checksum_sha256=checksum, scan_status=stored.scan_status, uploaded_by=actor.id)
    session.add(row)
    session.flush()
    db.write_audit(session, actor.id, "student_document.uploaded", "student_document", row.id, {"student_id": student_id, "category": row.category, "checksum": checksum})
    session.commit()
    return {"id": row.id, "scan_status": row.scan_status, "file_name": row.file_name}


@app.get("/api/student-documents/{document_id}/download")
def download_student_document(document_id: int, actor: User = Depends(current_user), session: Session = Depends(db.get_db)):
    row = session.get(StudentDocument, document_id)
    if not row or row.institution_id != actor.institution_id:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Document was not found")
    require_student_document_access(session, actor, row.student_id)
    if row.scan_status != "clean":
        raise HTTPException(status.HTTP_423_LOCKED, "Document is quarantined until malware scanning succeeds")
    db.write_audit(session, actor.id, "student_document.downloaded", "student_document", row.id, {"student_id": row.student_id})
    session.commit()
    local = storage_service.local_path(row.storage_key)
    if local:
        return FileResponse(local, media_type=row.mime_type, filename=row.file_name)
    return {"url": storage_service.download_url(row.storage_key, 300), "expires_in": 300}


@app.post("/api/auth/logout")
def logout(request: Request, response: Response, user: User = Depends(current_user), session: Session = Depends(db.get_db)):
    token = request.cookies.get(settings.session_cookie_name)
    payload = decode_token(token) if token else None
    if payload and payload.get("jti"):
        revoke_session(payload["jti"], settings.access_token_minutes * 60)
    user.token_version += 1
    session.query(PushDevice).filter(PushDevice.user_id == user.id).update({PushDevice.active: False})
    db.write_audit(session, user.id, "auth.logout", "user", user.id)
    session.commit()
    response.delete_cookie(settings.session_cookie_name, path="/api", domain=settings.session_cookie_domain)
    return {"message": "Signed out"}


@app.post("/api/auth/change-password")
def change_password(payload: PasswordChange, response: Response, http_request: Request, user: User = Depends(current_user), session: Session = Depends(db.get_db)):
    if not verify_password(payload.current_password, user.password_hash):
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "Current password is incorrect")
    user.password_hash = hash_password(payload.new_password)
    user.must_change_password = False
    user.failed_login_count = 0
    user.locked_until = None
    user.token_version += 1
    db.write_audit(session, user.id, "auth.password_changed", "user", user.id)
    session.commit()
    token = create_token(user)
    token_payload = decode_token(token) or {}
    if token_payload.get("jti") and not start_session(token_payload["jti"], settings.access_token_minutes * 60):
        raise HTTPException(status.HTTP_503_SERVICE_UNAVAILABLE, "Secure session service is unavailable")
    response.set_cookie(settings.session_cookie_name, token, max_age=settings.access_token_minutes * 60, httponly=True, secure=settings.is_production, samesite="lax", path="/api", domain=settings.session_cookie_domain)
    result = {"user": db.public_user(user, session)}
    if http_request.headers.get("x-client-platform") in {"android", "ios"}:
        result["access_token"] = token
    return result


@app.get("/api/dashboard")
def dashboard(user: User = Depends(current_user), session: Session = Depends(db.get_db)):
    if user.platform_admin:
        return {"platform": True}
    return db.dashboard_for(session, user)


@app.get("/api/portal")
def portal(user: User = Depends(current_user), session: Session = Depends(db.get_db)):
    return db.portal_for_user(session, user)


@app.get("/api/public/push-config")
def public_push_config():
    import push_service
    try:
        config = json.loads(settings.firebase_web_config)
    except ValueError:
        config = {}
    allowed_keys = {"apiKey", "authDomain", "projectId", "storageBucket", "messagingSenderId", "appId"}
    public_config = {k: v for k, v in config.items() if k in allowed_keys and isinstance(v, str)}
    return {"enabled": push_service.enabled(), "firebase": public_config, "vapid_key": settings.firebase_vapid_key}


@app.get("/api/me/notifications")
def my_notifications(user: User = Depends(current_user), session: Session = Depends(db.get_db)):
    rows = session.scalars(select(NotificationDelivery).where(NotificationDelivery.user_id == user.id, NotificationDelivery.institution_id == user.institution_id, NotificationDelivery.channel == "in_app").order_by(NotificationDelivery.created_at.desc()).limit(100)).all()
    return [{"id": r.id, "template": r.template, "created_at": r.created_at.isoformat(), **json.loads(r.payload_json or "{}")} for r in rows]


@app.post("/api/me/push-devices", status_code=status.HTTP_201_CREATED)
def add_push_device(payload: PushDeviceCreate, user: User = Depends(current_user), session: Session = Depends(db.get_db)):
    device = db.register_push_device(session, user, payload.token, payload.platform, payload.device_label)
    db.write_audit(session, user.id, "push_device.registered", "push_device", device.id, {"platform": device.platform})
    session.commit()
    return {"id": device.id, "platform": device.platform, "device_label": device.device_label, "active": device.active}


@app.patch("/api/me/push-devices/{device_id}")
def update_push_preferences(device_id: int, payload: PushPreferenceUpdate, user: User = Depends(current_user), session: Session = Depends(db.get_db)):
    device = session.get(PushDevice, device_id)
    if not device or device.user_id != user.id or device.institution_id != user.institution_id:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Push device not found")
    device.attendance_enabled = payload.attendance_enabled
    device.security_enabled = payload.security_enabled
    device.last_seen_at = datetime.now(timezone.utc)
    db.write_audit(session, user.id, "push_device.preferences_updated", "push_device", device.id, payload.model_dump())
    session.commit()
    return {"id": device.id, **payload.model_dump()}


@app.delete("/api/me/push-devices/{device_id}")
def remove_push_device(device_id: int, user: User = Depends(current_user), session: Session = Depends(db.get_db)):
    device = session.get(PushDevice, device_id)
    if not device or device.user_id != user.id or device.institution_id != user.institution_id:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Push device not found")
    device.active = False
    db.write_audit(session, user.id, "push_device.disabled", "push_device", device.id)
    session.commit()
    return {"message": "Push notifications disabled for this device"}


@app.patch("/api/me/profile")
def update_student_profile(payload: StudentProfileUpdate, user: User = Depends(current_user), session: Session = Depends(db.get_db)):
    student = session.scalar(select(Student).where(Student.user_id == user.id))
    if not student:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "Only a student account has a student profile")
    profile = session.scalar(select(StudentProfile).where(StudentProfile.student_id == student.id))
    if not profile:
        profile = StudentProfile(student_id=student.id)
        session.add(profile)
    if not profile.self_edit_enabled:
        raise HTTPException(status.HTTP_403_FORBIDDEN, "Your institution has not enabled personal-detail editing. Contact the administrator for a correction.")
    # Authoritative identity/guardian fields remain administrator controlled.
    for key, value in payload.model_dump(include={"phone", "address", "messaging_consent"}).items():
        setattr(profile, key, value.strip() if isinstance(value, str) else value)
    db.write_audit(session, user.id, "student_profile.updated", "student", student.id, {"messaging_consent": profile.messaging_consent})
    session.commit()
    return db.portal_for_user(session, user)["profile"]


@app.post("/api/announcements", status_code=status.HTTP_201_CREATED)
def create_announcement(payload: AnnouncementCreate, actor: User = Depends(teacher_only), session: Session = Depends(db.get_db)):
    if payload.section_id:
        section = session.get(ClassSection, payload.section_id)
        if not section or section.institution_id != actor.institution_id:
            raise HTTPException(status.HTTP_404_NOT_FOUND, "Class section not found")
    row = Announcement(institution_id=actor.institution_id, created_by=actor.id, title=payload.title.strip(), body=payload.body.strip(), audience=payload.audience, section_id=payload.section_id, published_at=datetime.now(timezone.utc), expires_at=payload.expires_at)
    session.add(row)
    session.flush()
    notified = notify_announcement_audience(session, row)
    db.write_audit(session, actor.id, "announcement.published", "announcement", row.id, {"audience": row.audience, "section_id": row.section_id, "notified": notified})
    session.commit()
    return {"id": row.id, "published_at": db.campus_datetime_iso(row.published_at)}


def od_request_dict(session: Session, row: OnDutyRequest) -> dict:
    student = session.get(Student, row.student_id)
    account = session.get(User, student.user_id) if student else None
    reviews = session.scalars(select(OnDutyReview).where(OnDutyReview.request_id == row.id).order_by(OnDutyReview.id)).all()
    return {"id": row.id, "student_id": row.student_id, "student_name": account.name if account else "Unknown", "roll_no": student.roll_no if student else "", "from_date": str(row.from_date), "to_date": str(row.to_date), "periods": json.loads(row.periods_json), "reason": row.reason, "attachment_name": row.attachment_name, "status": row.status, "current_stage": row.current_stage, "created_at": db.campus_datetime_iso(row.created_at), "resolved_at": db.campus_datetime_iso(row.resolved_at), "reviews": [{"stage": review.stage, "decision": review.decision, "reason": review.reason, "reviewed_at": db.campus_datetime_iso(review.reviewed_at)} for review in reviews]}


@app.post("/api/on-duty", status_code=201)
async def create_on_duty_request(from_date: date = Form(...), to_date: date = Form(...), periods: str = Form(...), reason: str = Form(...), letter: UploadFile = File(...), user: User = Depends(current_user), session: Session = Depends(db.get_db)):
    student = session.scalar(select(Student).where(Student.user_id == user.id))
    if not student: raise HTTPException(403, "Only students can submit an on-duty request")
    if not fixed_window_allow(f"upload:on-duty:{user.id}", 10, 3600):
        raise HTTPException(status.HTTP_429_TOO_MANY_REQUESTS, "Too many supporting-letter uploads. Try again later.")
    if to_date < from_date or (to_date - from_date).days > 31: raise HTTPException(400, "Choose a valid date range of up to 31 days")
    try: period_values = sorted({int(value) for value in periods.split(",") if value.strip()})
    except ValueError as exc: raise HTTPException(400, "Periods must be comma-separated numbers") from exc
    if not period_values or any(value < 1 or value > 30 for value in period_values): raise HTTPException(400, "Choose at least one valid period")
    if len(reason.strip()) < 5: raise HTTPException(400, "Enter the reason for on-duty attendance")
    if letter.content_type not in {"application/pdf", "image/jpeg", "image/png"}: raise HTTPException(400, "Upload a PDF, JPG or PNG letter")
    content = await letter.read(settings.max_upload_mb * 1024 * 1024 + 1)
    if not content or len(content) > settings.max_upload_mb * 1024 * 1024: raise HTTPException(status.HTTP_413_REQUEST_ENTITY_TOO_LARGE, "The supporting letter is empty or too large")
    if letter.content_type == "application/pdf" and not content.startswith(b"%PDF-"):
        raise HTTPException(400, "The uploaded file is not a valid PDF")
    if letter.content_type in {"image/jpeg", "image/png"} and cv2.imdecode(np.frombuffer(content, np.uint8), cv2.IMREAD_COLOR) is None:
        raise HTTPException(400, "The uploaded file is not a valid image")
    safe_name = Path(letter.filename or "letter").name[:255]
    membership = session.scalar(select(SectionMembership).where(SectionMembership.student_id == student.id, SectionMembership.active.is_(True)).order_by(SectionMembership.id.desc()))
    if not membership: raise HTTPException(409, "Your active class and term assignment is missing")
    mentor_id = session.scalar(select(ClassMentor.teacher_id).where(ClassMentor.section_id == membership.section_id, ClassMentor.term_id == membership.term_id))
    department_id = student.department_id
    hod_id = session.scalar(select(StaffProfile.user_id).where(StaffProfile.institution_id == user.institution_id, StaffProfile.designation == "hod", StaffProfile.department_id == department_id, StaffProfile.active.is_(True)))
    dean_id = session.scalar(select(StaffProfile.user_id).where(StaffProfile.institution_id == user.institution_id, StaffProfile.designation.in_(["dean", "principal"]), StaffProfile.active.is_(True)))
    missing = [stage for stage, reviewer in (("mentor", mentor_id), ("hod", hod_id), ("dean", dean_id)) if not reviewer]
    if missing: raise HTTPException(409, f"The institution must assign these reviewers first: {', '.join(missing)}")
    object_key = f"institutions/{user.institution_id}/on-duty/{student.id}/{uuid4().hex}{Path(safe_name).suffix.lower()}"
    try:
        stored = await run_in_threadpool(storage_service.put_private, object_key, content, letter.content_type)
    except ValueError as exc:
        raise HTTPException(status.HTTP_422_UNPROCESSABLE_ENTITY, str(exc)) from exc
    except Exception as exc:
        logger.exception("On-duty attachment storage failed")
        raise HTTPException(status.HTTP_503_SERVICE_UNAVAILABLE, "Secure upload scanning or storage is unavailable") from exc
    row = OnDutyRequest(institution_id=user.institution_id, student_id=student.id, from_date=from_date, to_date=to_date, periods_json=json.dumps(period_values), reason=reason.strip(), attachment_name=safe_name, attachment_path=stored.key, status="pending", current_stage="mentor")
    session.add(row); session.flush()
    for stage, reviewer in (("mentor", mentor_id), ("hod", hod_id), ("dean", dean_id)):
        session.add(OnDutyReview(request_id=row.id, stage=stage, reviewer_user_id=reviewer, decision="pending"))
    db.write_audit(session, user.id, "on_duty.submitted", "on_duty_request", row.id, {"from": str(from_date), "to": str(to_date), "periods": period_values})
    session.commit()
    return od_request_dict(session, row)


@app.get("/api/on-duty")
def list_on_duty_requests(user: User = Depends(current_user), session: Session = Depends(db.get_db)):
    student = session.scalar(select(Student).where(Student.user_id == user.id)) if user.role == "student" else None
    if student:
        rows = session.scalars(select(OnDutyRequest).where(OnDutyRequest.student_id == student.id).order_by(OnDutyRequest.created_at.desc())).all()
    elif user.role in {"teacher", "admin"}:
        statement = select(OnDutyRequest).where(OnDutyRequest.institution_id == user.institution_id)
        if user.role != "admin":
            statement = statement.join(OnDutyReview, OnDutyReview.request_id == OnDutyRequest.id).where(OnDutyReview.reviewer_user_id == user.id).distinct()
        rows = session.scalars(statement.order_by(OnDutyRequest.created_at.desc())).all()
    else: raise HTTPException(403, "On-duty requests are unavailable")
    return [od_request_dict(session, row) for row in rows]


@app.patch("/api/on-duty/{request_id}/review")
def review_on_duty_request(request_id: int, payload: OnDutyReviewUpdate, actor: User = Depends(teacher_only), session: Session = Depends(db.get_db)):
    row = session.scalar(select(OnDutyRequest).where(OnDutyRequest.id == request_id).with_for_update())
    if not row or row.institution_id != actor.institution_id: raise HTTPException(404, "On-duty request not found")
    if row.status != "pending": raise HTTPException(409, "This request has already been resolved")
    review = session.scalar(select(OnDutyReview).where(OnDutyReview.request_id == row.id, OnDutyReview.stage == row.current_stage))
    if not review or (review.reviewer_user_id != actor.id and actor.role != "admin"): raise HTTPException(403, "This request is waiting for another reviewer")
    review.decision, review.reason, review.reviewed_at = payload.decision, payload.reason, datetime.now(timezone.utc)
    if payload.decision == "rejected":
        row.status, row.resolved_at = "rejected", datetime.now(timezone.utc)
    else:
        stages = ["mentor", "hod", "dean"]
        index = stages.index(row.current_stage)
        if index < len(stages) - 1: row.current_stage = stages[index + 1]
        else:
            row.status, row.current_stage, row.resolved_at = "approved", "completed", datetime.now(timezone.utc)
            periods = json.loads(row.periods_json)
            sessions = session.scalars(select(ClassSession).where(ClassSession.section_id.in_(select(SectionMembership.section_id).where(SectionMembership.student_id == row.student_id, SectionMembership.active.is_(True))), ClassSession.session_date >= row.from_date, ClassSession.session_date <= row.to_date, ClassSession.period_number.in_(periods))).all()
            for class_session in sessions:
                attendance = session.scalar(select(SessionAttendance).where(SessionAttendance.class_session_id == class_session.id, SessionAttendance.student_id == row.student_id))
                if not attendance:
                    attendance = SessionAttendance(class_session_id=class_session.id, student_id=row.student_id, status="OD", decision_source="approved_od", od_request_id=row.id, marked_at=datetime.now(timezone.utc)); session.add(attendance)
                else:
                    attendance.od_original_status, attendance.od_original_source = attendance.status, attendance.decision_source
                    attendance.status, attendance.decision_source, attendance.od_request_id = "OD", "approved_od", row.id
    db.write_audit(session, actor.id, f"on_duty.{payload.decision}", "on_duty_request", row.id, {"stage": review.stage, "reason": payload.reason})
    session.commit()
    return od_request_dict(session, row)


@app.get("/api/on-duty/{request_id}/letter")
def download_on_duty_letter(request_id: int, user: User = Depends(current_user), session: Session = Depends(db.get_db)):
    row = session.get(OnDutyRequest, request_id)
    student = session.scalar(select(Student).where(Student.user_id == user.id))
    review = session.scalar(select(OnDutyReview).where(OnDutyReview.request_id == request_id, OnDutyReview.reviewer_user_id == user.id))
    if not row or row.institution_id != user.institution_id or not (user.role == "admin" or review or (student and student.id == row.student_id)): raise HTTPException(404, "Letter not found")
    key = row.attachment_path or ""
    local = storage_service.local_path(key)
    if local:
        return FileResponse(local, filename=row.attachment_name)
    if key.startswith(f"institutions/{user.institution_id}/on-duty/"):
        return RedirectResponse(storage_service.download_url(key, 300), status_code=307)
    # Backward compatibility for v8.1 records. Migrate these files to private
    # object storage before production; keep the legacy path confined to its root.
    legacy = Path(key).resolve()
    legacy_root = OD_ATTACHMENT_ROOT.resolve()
    if legacy.is_file() and legacy_root in legacy.parents:
        return FileResponse(legacy, filename=row.attachment_name)
    raise HTTPException(404, "Letter file is unavailable")


@app.post("/api/admin/fees", status_code=status.HTTP_201_CREATED)
def upsert_fee(payload: FeeRecordUpsert, actor: User = Depends(admin_only), session: Session = Depends(db.get_db)):
    student = session.get(Student, payload.student_id)
    if not student or student.institution_id != actor.institution_id:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Student not found")
    if payload.paid_amount > payload.total_amount:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "Paid amount cannot exceed the total amount")
    row = session.scalar(select(FeeRecord).where(FeeRecord.student_id == student.id, FeeRecord.term_label == payload.term_label.strip(), FeeRecord.description == payload.description.strip()))
    if not row:
        row = FeeRecord(institution_id=actor.institution_id, student_id=student.id, term_label=payload.term_label.strip(), description=payload.description.strip(), total_amount=payload.total_amount, paid_amount=payload.paid_amount, due_date=payload.due_date, updated_by=actor.id)
        session.add(row)
    else:
        row.total_amount = payload.total_amount
        row.paid_amount = payload.paid_amount
        row.due_date = payload.due_date
        row.updated_by = actor.id
    row.status = "paid" if row.paid_amount >= row.total_amount else "partial" if row.paid_amount else "due"
    session.flush()
    db.add_notification(
        session,
        institution_id=actor.institution_id,
        user_id=student.user_id,
        student_id=student.id,
        template="fee_record_updated",
        message=f"Fee update: {row.description} has a balance of {max(0, row.total_amount - row.paid_amount):.2f}.",
        dedup_key=f"fee-record:{row.id}:{row.total_amount}:{row.paid_amount}:{row.due_date}",
        deep_link="/portal?view=fees",
    )
    db.write_audit(session, actor.id, "fee_record.updated", "fee_record", row.id, {"student_id": student.id, "status": row.status})
    session.commit()
    return {"id": row.id, "status": row.status, "balance": max(0, row.total_amount - row.paid_amount)}


def apply_verified_payment(fee: FeeRecord, payment: FeePayment) -> None:
    if payment.status == "paid":
        return
    payment.status = "paid"
    payment.paid_at = datetime.now(timezone.utc)
    fee.paid_amount = min(fee.total_amount, fee.paid_amount + float(payment.amount))
    fee.status = "paid" if fee.paid_amount >= fee.total_amount else "partial"


@app.post("/api/student/fees/{fee_record_id}/payments/online", status_code=status.HTTP_201_CREATED)
def create_online_fee_payment(fee_record_id: int, payload: OnlinePaymentCreate, user: User = Depends(current_user), session: Session = Depends(db.get_db)):
    if user.role != "student":
        raise HTTPException(status.HTTP_403_FORBIDDEN, "Student access is required")
    if not fixed_window_allow(f"upload:fee-receipt:{user.id}", 10, 3600):
        raise HTTPException(status.HTTP_429_TOO_MANY_REQUESTS, "Too many receipt uploads. Try again later.")
    student = session.scalar(select(Student).where(Student.user_id == user.id))
    fee = session.get(FeeRecord, fee_record_id)
    if not student or not fee or fee.student_id != student.id or fee.institution_id != user.institution_id:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Fee record not found")
    balance = max(0.0, fee.total_amount - fee.paid_amount)
    if balance <= 0:
        raise HTTPException(status.HTTP_409_CONFLICT, "This fee is already paid")
    if payload.amount > balance + 0.001:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, f"Payment cannot exceed the current balance of {balance:.2f}")
    existing = session.scalar(select(FeePayment).where(FeePayment.institution_id == user.institution_id, FeePayment.idempotency_key == payload.idempotency_key))
    if existing:
        if existing.student_id != student.id or existing.fee_record_id != fee.id or abs(float(existing.amount) - payload.amount) > 0.001:
            raise HTTPException(status.HTTP_409_CONFLICT, "This idempotency key was already used for a different payment")
        return {"payment_id": existing.id, "status": existing.status, "provider": existing.gateway_provider, "order_id": existing.gateway_order_id, "amount": float(existing.amount), "currency": existing.currency}
    payment = FeePayment(institution_id=user.institution_id, student_id=student.id, fee_record_id=fee.id, amount=payload.amount, currency=settings.payment_currency, payment_method="online", status="creating", idempotency_key=payload.idempotency_key, gateway_provider=settings.payment_provider)
    session.add(payment)
    session.flush()
    try:
        order = payment_service.create_order(payload.amount, f"OURS-{payment.id}", {"payment_id": str(payment.id), "student_id": str(student.id), "fee_record_id": str(fee.id)})
    except (RuntimeError, ValueError) as exc:
        payment.status = "failed"
        session.commit()
        raise HTTPException(status.HTTP_503_SERVICE_UNAVAILABLE, str(exc)) from exc
    payment.status = "pending"
    payment.gateway_provider = order["provider"]
    payment.gateway_order_id = order["order_id"]
    db.write_audit(session, user.id, "fee_payment.online_order_created", "fee_payment", payment.id, {"fee_record_id": fee.id, "amount": payload.amount, "provider": order["provider"]})
    session.commit()
    return {"payment_id": payment.id, "status": payment.status, "provider": order["provider"], "order_id": order["order_id"], "public_key": order["public_key"], "amount_minor": order["amount_minor"], "amount": float(payment.amount), "currency": order["currency"]}


@app.post("/api/student/fees/{fee_record_id}/payments/offline", status_code=status.HTTP_201_CREATED)
async def submit_offline_fee_payment(fee_record_id: int, amount: float = Form(...), reference: str = Form(...), paid_on: date = Form(...), receipt: UploadFile = File(...), user: User = Depends(current_user), session: Session = Depends(db.get_db)):
    if user.role != "student":
        raise HTTPException(status.HTTP_403_FORBIDDEN, "Student access is required")
    student = session.scalar(select(Student).where(Student.user_id == user.id))
    fee = session.get(FeeRecord, fee_record_id)
    if not student or not fee or fee.student_id != student.id or fee.institution_id != user.institution_id:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Fee record not found")
    balance = max(0.0, fee.total_amount - fee.paid_amount)
    if amount <= 0 or amount > balance + 0.001:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "Enter an amount no greater than the current fee balance")
    reference = reference.strip()
    if len(reference) < 3 or len(reference) > 160:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "Enter a valid receipt or bank reference")
    allowed_types = {"application/pdf", "image/jpeg", "image/png"}
    if receipt.content_type not in allowed_types:
        raise HTTPException(status.HTTP_415_UNSUPPORTED_MEDIA_TYPE, "Upload a PDF, JPEG or PNG receipt")
    content = await receipt.read(settings.max_upload_mb * 1024 * 1024 + 1)
    if not content or len(content) > settings.max_upload_mb * 1024 * 1024:
        raise HTTPException(status.HTTP_413_REQUEST_ENTITY_TOO_LARGE, "Receipt is empty or exceeds the upload limit")
    original_name = Path(receipt.filename or "receipt").name[:255]
    if not valid_assignment_file(original_name, receipt.content_type, content):
        raise HTTPException(status.HTTP_415_UNSUPPORTED_MEDIA_TYPE, "The receipt content does not match a genuine PDF, JPEG or PNG file")
    suffix = {"application/pdf": ".pdf", "image/jpeg": ".jpg", "image/png": ".png"}[receipt.content_type]
    storage_name = f"institutions/{user.institution_id}/fee-receipts/{student.id}/{uuid4().hex}{suffix}"
    try:
        stored = await run_in_threadpool(storage_service.put_private, storage_name, content, receipt.content_type)
    except ValueError as exc:
        raise HTTPException(status.HTTP_422_UNPROCESSABLE_ENTITY, str(exc)) from exc
    except Exception as exc:
        logger.exception("Fee receipt storage failed")
        raise HTTPException(status.HTTP_503_SERVICE_UNAVAILABLE, "Secure upload scanning or storage is unavailable") from exc
    payment = FeePayment(institution_id=user.institution_id, student_id=student.id, fee_record_id=fee.id, amount=amount, currency=settings.payment_currency, payment_method="offline", status="pending_review", idempotency_key=f"offline:{uuid4().hex}", offline_reference=reference, paid_on=paid_on, receipt_file_name=original_name, receipt_storage_name=storage_name, receipt_mime_type=receipt.content_type, receipt_size_bytes=len(content), receipt_checksum_sha256=hashlib.sha256(content).hexdigest())
    session.add(payment)
    try:
        session.flush()
    except IntegrityError as exc:
        session.rollback()
        try:
            storage_service.delete_private(storage_name)
        except Exception:
            logger.exception("Failed to remove orphaned fee receipt after transaction rollback")
        raise HTTPException(status.HTTP_409_CONFLICT, "This offline payment reference was already submitted") from exc
    db.write_audit(session, user.id, "fee_payment.offline_submitted", "fee_payment", payment.id, {"fee_record_id": fee.id, "amount": amount})
    session.commit()
    return {"payment_id": payment.id, "status": payment.status, "message": "Receipt submitted for administrator verification"}


@app.get("/api/admin/fee-payments")
def list_fee_payments(payment_status: str | None = None, actor: User = Depends(admin_only), session: Session = Depends(db.get_db)):
    statement = select(FeePayment, Student.roll_no, User.name).join(Student, Student.id == FeePayment.student_id).join(User, User.id == Student.user_id).where(FeePayment.institution_id == actor.institution_id)
    if payment_status:
        statement = statement.where(FeePayment.status == payment_status)
    rows = session.execute(statement.order_by(FeePayment.created_at.desc()).limit(300)).all()
    return [{"id": payment.id, "student_id": payment.student_id, "student_name": name, "roll_no": roll_no, "fee_record_id": payment.fee_record_id, "amount": float(payment.amount), "currency": payment.currency, "payment_method": payment.payment_method, "status": payment.status, "offline_reference": payment.offline_reference, "paid_on": str(payment.paid_on) if payment.paid_on else None, "receipt_file_name": payment.receipt_file_name, "created_at": payment.created_at.isoformat()} for payment, roll_no, name in rows]


@app.patch("/api/admin/fee-payments/{payment_id}/review")
def review_offline_fee_payment(payment_id: int, payload: OfflinePaymentReview, actor: User = Depends(admin_only), session: Session = Depends(db.get_db)):
    payment = session.get(FeePayment, payment_id)
    fee = session.get(FeeRecord, payment.fee_record_id) if payment else None
    if not payment or not fee or payment.institution_id != actor.institution_id or payment.payment_method != "offline":
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Offline fee payment not found")
    if payment.status in {"paid", "rejected"}:
        return {"id": payment.id, "status": payment.status}
    if payload.decision == "rejected" and not (payload.reason or "").strip():
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "A rejection reason is required")
    if payload.decision == "approved":
        current_balance = max(0.0, fee.total_amount - fee.paid_amount)
        if float(payment.amount) > current_balance + 0.001:
            raise HTTPException(status.HTTP_409_CONFLICT, "This payment exceeds the remaining balance; review newer payments first")
        apply_verified_payment(fee, payment)
    else:
        payment.status = "rejected"
    payment.review_reason = (payload.reason or "").strip() or None
    payment.reviewed_by = actor.id
    payment.reviewed_at = datetime.now(timezone.utc)
    student = session.get(Student, payment.student_id)
    db.add_notification(session, institution_id=actor.institution_id, user_id=student.user_id if student else None, student_id=student.id if student else None, template="fee_payment_reviewed", message=f"Your offline fee payment of {float(payment.amount):.2f} was {payload.decision}.", dedup_key=f"fee-review:{payment.id}:{payment.status}", deep_link="/portal?view=fees")
    db.write_audit(session, actor.id, "fee_payment.reviewed", "fee_payment", payment.id, {"decision": payload.decision, "reason": payment.review_reason})
    session.commit()
    return {"id": payment.id, "status": payment.status, "fee_status": fee.status, "fee_paid_amount": fee.paid_amount}


@app.get("/api/fee-payments/{payment_id}/receipt")
def download_fee_receipt(payment_id: int, actor: User = Depends(current_user), session: Session = Depends(db.get_db)):
    payment = session.get(FeePayment, payment_id)
    if not payment or payment.institution_id != actor.institution_id or not payment.receipt_storage_name:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Payment receipt not found")
    if actor.role == "student":
        student = session.scalar(select(Student).where(Student.user_id == actor.id))
        if not student or payment.student_id != student.id:
            raise HTTPException(status.HTTP_403_FORBIDDEN, "You cannot access this receipt")
    elif actor.role != "admin":
        raise HTTPException(status.HTTP_403_FORBIDDEN, "Administrator or receipt owner access is required")
    key = payment.receipt_storage_name
    if not key.startswith(f"institutions/{actor.institution_id}/fee-receipts/"):
        # Legacy v8.1 receipts are still restricted to the old receipt directory.
        legacy_name = Path(key).name
        legacy = (PAYMENT_RECEIPT_ROOT / legacy_name).resolve()
        if legacy.parent == PAYMENT_RECEIPT_ROOT.resolve() and legacy.is_file():
            return FileResponse(legacy, media_type=payment.receipt_mime_type, filename=payment.receipt_file_name)
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Receipt file is unavailable")
    local = storage_service.local_path(key)
    if local:
        return FileResponse(local, media_type=payment.receipt_mime_type, filename=payment.receipt_file_name)
    return RedirectResponse(storage_service.download_url(key, 300), status_code=307)


@app.post("/api/payments/webhooks/razorpay")
async def razorpay_webhook(request: Request, session: Session = Depends(db.get_db)):
    payload = await request.body()
    if len(payload) > 1024 * 1024:
        raise HTTPException(status.HTTP_413_REQUEST_ENTITY_TOO_LARGE, "Webhook payload is too large")
    signature = request.headers.get("x-razorpay-signature")
    if not payment_service.verify_razorpay_webhook(payload, signature):
        raise HTTPException(status.HTTP_401_UNAUTHORIZED, "Invalid webhook signature")
    try:
        event = json.loads(payload)
    except json.JSONDecodeError as exc:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "Invalid webhook JSON") from exc
    event_id = request.headers.get("x-razorpay-event-id") or hashlib.sha256(payload).hexdigest()
    existing = session.scalar(select(PaymentWebhookEvent).where(PaymentWebhookEvent.provider == "razorpay", PaymentWebhookEvent.event_id == event_id))
    if existing:
        return {"status": existing.status, "duplicate": True}
    event_row = PaymentWebhookEvent(provider="razorpay", event_id=event_id, payload_hash=hashlib.sha256(payload).hexdigest())
    session.add(event_row)
    session.flush()
    event_name = str(event.get("event", ""))
    payment_entity = (((event.get("payload") or {}).get("payment") or {}).get("entity") or {})
    order_entity = (((event.get("payload") or {}).get("order") or {}).get("entity") or {})
    order_id = payment_entity.get("order_id") or order_entity.get("id")
    payment = session.scalar(select(FeePayment).where(FeePayment.gateway_provider == "razorpay", FeePayment.gateway_order_id == order_id)) if order_id else None
    if event_name not in {"payment.captured", "order.paid", "payment.failed"}:
        event_row.status = "ignored"
    elif not payment:
        event_row.status = "unmatched"
        event_row.error = "No internal payment matches the provider order"
    elif event_name == "payment.failed":
        if payment.status != "paid":
            payment.status = "failed"
        event_row.status = "processed"
    else:
        fee = session.get(FeeRecord, payment.fee_record_id)
        transaction_id = payment_entity.get("id")
        if not fee or fee.institution_id != payment.institution_id:
            event_row.status = "failed"
            event_row.error = "Fee ledger record is unavailable"
        else:
            if transaction_id and session.scalar(select(FeePayment.id).where(FeePayment.gateway_provider == "razorpay", FeePayment.gateway_transaction_id == transaction_id, FeePayment.id != payment.id)):
                event_row.status = "failed"
                event_row.error = "Provider transaction is already linked to another payment"
            else:
                payment.gateway_transaction_id = transaction_id or payment.gateway_transaction_id
                apply_verified_payment(fee, payment)
                event_row.status = "processed"
                student = session.get(Student, payment.student_id)
                db.add_notification(session, institution_id=payment.institution_id, user_id=student.user_id if student else None, student_id=student.id if student else None, template="fee_payment_confirmed", message=f"Your online fee payment of {float(payment.amount):.2f} is confirmed.", dedup_key=f"fee-paid:{payment.id}", deep_link="/portal?view=fees")
    event_row.processed_at = datetime.now(timezone.utc)
    session.commit()
    return {"status": event_row.status, "duplicate": False}


@app.post("/api/student/teacher-feedback", status_code=status.HTTP_201_CREATED)
def submit_teacher_feedback(payload: TeacherFeedbackCreate, user: User = Depends(current_user), session: Session = Depends(db.get_db)):
    if user.role != "student":
        raise HTTPException(status.HTTP_403_FORBIDDEN, "Student access is required")
    student = session.scalar(select(Student).where(Student.user_id == user.id))
    assignment = session.get(TeacherAssignment, payload.assignment_id)
    membership = session.scalar(select(SectionMembership.id).where(SectionMembership.student_id == student.id, SectionMembership.section_id == assignment.section_id, SectionMembership.term_id == assignment.term_id, SectionMembership.active.is_(True))) if student and assignment else None
    if not student or not assignment or not membership:
        raise HTTPException(status.HTTP_403_FORBIDDEN, "This teacher is not assigned to your class and term")
    row = TeacherFeedback(institution_id=user.institution_id, teacher_id=assignment.teacher_id, student_id=student.id, course_id=assignment.course_id, section_id=assignment.section_id, term_id=assignment.term_id, **payload.model_dump(exclude={"assignment_id"}))
    session.add(row)
    flush_or_conflict(session, "Feedback for this teacher, subject and week was already submitted")
    db.write_audit(session, user.id, "teacher_feedback.submitted", "teacher_feedback", row.id, {"assignment_id": assignment.id})
    session.commit()
    return {"id": row.id, "message": "Anonymous weekly feedback submitted"}


def career_identity_for(session: Session, user: User) -> StaffCareerIdentity:
    identity = session.scalar(select(StaffCareerIdentity).where(StaffCareerIdentity.user_id == user.id))
    if not identity:
        identity = StaffCareerIdentity(user_id=user.id, platform_staff_id=f"OURS-{user.institution_id:04d}-{user.id:06d}-{secrets.token_hex(2).upper()}")
        session.add(identity)
        session.flush()
        profile = session.scalar(select(StaffProfile).where(StaffProfile.user_id == user.id))
        session.add(StaffCareerServiceRecord(identity_id=identity.id, institution_id=user.institution_id, local_employee_id=profile.employee_no if profile else None, designation=profile.designation if profile else user.role, started_on=campus_today(), active=True))
    return identity


@app.patch("/api/staff/career")
def update_career_profile(payload: CareerProfileUpdate, actor: User = Depends(teacher_only), session: Session = Depends(db.get_db)):
    identity = career_identity_for(session, actor)
    identity.professional_summary = (payload.professional_summary or "").strip() or None
    db.write_audit(session, actor.id, "career_profile.updated", "staff_career_identity", identity.id)
    session.commit()
    return db.portal_for_user(session, actor)["career"]


@app.post("/api/staff/career/transfers", status_code=status.HTTP_201_CREATED)
def request_career_transfer(payload: CareerTransferCreate, actor: User = Depends(teacher_only), session: Session = Depends(db.get_db)):
    if not payload.consent:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "Teacher consent is required")
    identity = session.scalar(select(StaffCareerIdentity).where(StaffCareerIdentity.platform_staff_id == payload.platform_staff_id.strip()))
    previous = session.scalar(select(Institution).where(Institution.slug == payload.previous_institution_slug.strip().lower()))
    if not identity or not previous:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Verified career identity or previous institution not found")
    if identity.user_id != actor.id:
        raise HTTPException(status.HTTP_403_FORBIDDEN, "Sign in to the account that owns this career identity")
    service = session.scalar(select(StaffCareerServiceRecord).where(StaffCareerServiceRecord.identity_id == identity.id, StaffCareerServiceRecord.institution_id == previous.id, StaffCareerServiceRecord.local_employee_id == payload.previous_employee_id.strip()))
    if not service:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "The previous employee identifier could not be verified")
    row = StaffTransferRequest(identity_id=identity.id, from_institution_id=previous.id, to_institution_id=actor.institution_id, previous_employee_id=payload.previous_employee_id.strip(), notes=payload.notes)
    session.add(row)
    session.flush()
    db.write_audit(session, actor.id, "career_transfer.requested", "staff_transfer_request", row.id, {"from_institution_id": previous.id, "to_institution_id": actor.institution_id})
    session.commit()
    return {"id": row.id, "status": row.status}


@app.get("/api/admin/career/transfers")
def career_transfers(actor: User = Depends(admin_only), session: Session = Depends(db.get_db)):
    rows = session.scalars(select(StaffTransferRequest).where((StaffTransferRequest.from_institution_id == actor.institution_id) | (StaffTransferRequest.to_institution_id == actor.institution_id)).order_by(StaffTransferRequest.created_at.desc())).all()
    return [{"id": row.id, "platform_staff_id": session.get(StaffCareerIdentity, row.identity_id).platform_staff_id, "from_institution_id": row.from_institution_id, "to_institution_id": row.to_institution_id, "previous_employee_id": row.previous_employee_id, "status": row.status, "created_at": row.created_at.isoformat()} for row in rows]


@app.patch("/api/admin/career/transfers/{transfer_id}")
def review_career_transfer(transfer_id: int, payload: CareerTransferReview, actor: User = Depends(admin_only), session: Session = Depends(db.get_db)):
    row = session.get(StaffTransferRequest, transfer_id)
    if not row or actor.institution_id not in {row.from_institution_id, row.to_institution_id}:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Transfer request not found")
    now = datetime.now(timezone.utc)
    if actor.institution_id == row.from_institution_id and row.status == "pending_previous_approval":
        if payload.approve:
            row.status = "pending_new_approval"
            row.previous_approved_by = actor.id
            row.previous_approved_at = now
        else:
            row.status = "rejected"
    elif actor.institution_id == row.to_institution_id and row.status == "pending_new_approval":
        if payload.approve:
            row.status = "approved"
            row.new_approved_by = actor.id
            row.new_approved_at = now
        else:
            row.status = "rejected"
    else:
        raise HTTPException(status.HTTP_409_CONFLICT, "This institution cannot review the request at its current stage")
    row.notes = payload.notes or row.notes
    db.write_audit(session, actor.id, "career_transfer.reviewed", "staff_transfer_request", row.id, {"status": row.status})
    session.commit()
    return {"id": row.id, "status": row.status}


@app.post("/api/admin/email/test")
def test_email_configuration(actor: User = Depends(admin_only), session: Session = Depends(db.get_db)):
    delivered, error = email_service.send_configuration_test(actor.email, actor.institution.name)
    session.add(NotificationDelivery(institution_id=actor.institution_id, user_id=actor.id, channel="email", destination_masked=f"***{actor.email[-6:]}", template="configuration_test", status="sent" if delivered else "failed", error=error, sent_at=datetime.now(timezone.utc) if delivered else None))
    db.write_audit(session, actor.id, "email.configuration_test", "institution", actor.institution_id, {"status": "sent" if delivered else "failed"})
    session.commit()
    if not delivered:
        raise HTTPException(status.HTTP_503_SERVICE_UNAVAILABLE, error or "Email configuration test failed")
    return {"message": f"Test email sent to {actor.email}"}


@app.get("/api/me/export")
def export_my_data(user: User = Depends(current_user), session: Session = Depends(db.get_db)):
    return {"account": db.public_user(user), "dashboard": db.dashboard_for(session, user)}


@app.patch("/api/me/biometric-consent")
def update_my_consent(payload: ConsentUpdate, user: User = Depends(current_user), session: Session = Depends(db.get_db)):
    student = session.scalar(select(Student).where(Student.user_id == user.id))
    if not student:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "This account is not a student")
    institution = session.get(Institution, user.institution_id)
    profile = session.scalar(select(StudentProfile).where(StudentProfile.student_id == student.id))
    minor_status = student_is_minor(profile)
    if payload.granted and institution and institution.minor_guardian_consent_required:
        if minor_status is None:
            raise HTTPException(status.HTTP_409_CONFLICT, "Save your date of birth before enabling face recognition. A college administrator must verify guardian consent for a minor.")
        if minor_status and not profile.guardian_biometric_consent:
            raise HTTPException(status.HTTP_409_CONFLICT, "A verified guardian consent record is required before a minor can enable face recognition")
    student.biometric_consent = payload.granted
    student.consent_recorded_at = datetime.now(timezone.utc) if payload.granted else None
    student.consent_recorded_by = user.id
    if not payload.granted:
        embedding = session.scalar(select(FaceEmbedding).where(FaceEmbedding.student_id == student.id))
        if embedding:
            session.delete(embedding)
    db.write_audit(session, user.id, "biometric.self_consent_updated", "student", student.id, {"granted": payload.granted})
    session.commit()
    return {"biometric_consent": payload.granted, "face_data_deleted": not payload.granted}


@app.get("/api/students")
def students(actor: User = Depends(teacher_only), session: Session = Depends(db.get_db)):
    return db.list_students(session, actor.institution_id, actor.id if actor.role == "teacher" else None)


@app.get("/api/admin/courses")
def courses(actor: User = Depends(admin_only), session: Session = Depends(db.get_db)):
    return db.list_courses(session, actor.institution_id)


@app.get("/api/admin/enrollments")
def enrollments(actor: User = Depends(admin_only), session: Session = Depends(db.get_db)):
    return db.list_enrollments(session, actor.institution_id)


@app.get("/api/admin/audit-logs")
def audit_logs(limit: int = 100, actor: User = Depends(admin_only), session: Session = Depends(db.get_db)):
    return db.list_audit_logs(session, actor.institution_id, limit)


@app.get("/api/admin/institution")
def admin_institution(actor: User = Depends(admin_only), session: Session = Depends(db.get_db)):
    institution = session.get(Institution, actor.institution_id)
    return db.public_institution(institution) if institution else {}


@app.patch("/api/admin/institution")
def update_admin_institution(payload: InstitutionUpdate, actor: User = Depends(admin_only), session: Session = Depends(db.get_db)):
    institution = session.get(Institution, actor.institution_id)
    if not institution:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Institution not found")
    institution.name = payload.name.strip()
    institution.brand_name = payload.brand_name.strip()
    institution.primary_color = payload.primary_color
    db.write_audit(session, actor.id, "institution.profile_updated", "institution", institution.id, payload.model_dump())
    session.commit()
    return db.public_institution(institution)


@app.patch("/api/admin/institution/monitoring-policy")
def update_monitoring_policy(payload: InstitutionMonitoringPolicy, actor: User = Depends(admin_only), session: Session = Depends(db.get_db)):
    institution = session.get(Institution, actor.institution_id)
    if not institution:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Institution not found")
    for key, value in payload.model_dump().items():
        setattr(institution, key, value)
    db.write_audit(session, actor.id, "institution.monitoring_policy_updated", "institution", institution.id, payload.model_dump())
    session.commit()
    return db.public_institution(institution)


@app.get("/api/admin/users")
def admin_users(actor: User = Depends(admin_only), session: Session = Depends(db.get_db)):
    rows = session.scalars(select(User).where(User.institution_id == actor.institution_id, User.active.is_(True)).order_by(User.role, User.name)).all()
    return [db.public_user(row) for row in rows]


@app.get("/api/staff")
def staff(actor: User = Depends(teacher_only), session: Session = Depends(db.get_db)):
    return db.list_staff(session, actor.institution_id)


@app.post("/api/admin/staff", status_code=status.HTTP_201_CREATED)
def create_staff(payload: StaffCreate, actor: User = Depends(admin_only), session: Session = Depends(db.get_db)):
    department = None
    if payload.department_id:
        department = session.get(Department, payload.department_id)
        if not department or department.institution_id != actor.institution_id:
            raise HTTPException(status.HTTP_404_NOT_FOUND, "Department not found for this institution")
    account = User(institution_id=actor.institution_id, name=payload.name.strip(), email=str(payload.email).lower(), password_hash=hash_password(payload.password), role=payload.access_role, must_change_password=True)
    session.add(account)
    flush_or_conflict(session, "A staff account already uses this email")
    profile = StaffProfile(institution_id=actor.institution_id, user_id=account.id, designation=payload.designation, employee_no=(payload.employee_no or "").strip().upper() or None, department_id=department.id if department else None, biometric_consent=payload.biometric_consent, consent_recorded_at=datetime.now(timezone.utc) if payload.biometric_consent else None, consent_recorded_by=actor.id if payload.biometric_consent else None)
    session.add(profile)
    db.write_audit(session, actor.id, "staff.created", "user", account.id, {"designation": payload.designation, "access_role": payload.access_role})
    commit_or_conflict(session, "Staff email or employee number already exists")
    return next(row for row in db.list_staff(session, actor.institution_id) if row["id"] == account.id)


@app.patch("/api/admin/staff/{user_id}/biometric-consent")
def update_staff_consent(user_id: int, payload: ConsentUpdate, actor: User = Depends(admin_only), session: Session = Depends(db.get_db)):
    account = session.get(User, user_id)
    if not account or account.institution_id != actor.institution_id or account.role not in ("teacher", "admin", "staff"):
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Staff account not found")
    profile = staff_profile_for(session, account)
    profile.biometric_consent = payload.granted
    profile.consent_recorded_at = datetime.now(timezone.utc) if payload.granted else None
    profile.consent_recorded_by = actor.id
    if not payload.granted:
        embedding = session.scalar(select(StaffFaceEmbedding).where(StaffFaceEmbedding.user_id == account.id))
        if embedding:
            session.delete(embedding)
    db.write_audit(session, actor.id, "staff.biometric_consent_updated", "user", account.id, {"granted": payload.granted})
    session.commit()
    return {"user_id": account.id, "biometric_consent": profile.biometric_consent, "face_data_deleted": not payload.granted}


@app.post("/api/admin/administrators", status_code=status.HTTP_201_CREATED)
def create_institution_administrator(payload: AdministratorCreate, actor: User = Depends(admin_only), session: Session = Depends(db.get_db)):
    administrator = User(institution_id=actor.institution_id, name=payload.name.strip(), email=str(payload.email).lower(), password_hash=hash_password(payload.password), role="admin", must_change_password=True)
    session.add(administrator)
    flush_or_conflict(session, "An account already uses this email")
    session.add(StaffProfile(institution_id=actor.institution_id, user_id=administrator.id, designation="administrator"))
    db.write_audit(session, actor.id, "administrator.created", "user", administrator.id)
    session.commit()
    return db.public_user(administrator)


@app.post("/api/admin/users/{user_id}/reset-password")
def admin_reset_user_password(user_id: int, payload: TemporaryPasswordReset, actor: User = Depends(admin_only), session: Session = Depends(db.get_db)):
    target = session.get(User, user_id)
    if not target or target.institution_id != actor.institution_id:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Account not found")
    if target.id == actor.id:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "Use Change password for your own account")
    if target.platform_admin and not actor.platform_admin:
        raise HTTPException(status.HTTP_403_FORBIDDEN, "Only a platform owner can reset this account")
    target.password_hash = hash_password(payload.temporary_password)
    target.must_change_password = True
    target.failed_login_count = 0
    target.locked_until = None
    target.token_version += 1
    db.write_audit(session, actor.id, "account.password_reset", "user", target.id, {"role": target.role})
    session.commit()
    return {"message": "Temporary password issued; all previous sessions were signed out"}


@app.get("/api/academics/workspace")
def academics_workspace(user: User = Depends(current_user), session: Session = Depends(db.get_db)):
    if user.platform_admin:
        raise HTTPException(status.HTTP_403_FORBIDDEN, "Platform accounts cannot open college academic workspaces")
    return db.academic_workspace(session, user)


@app.get("/api/admin/academics/catalog")
def academics_catalog(actor: User = Depends(allocation_manager), session: Session = Depends(db.get_db)):
    return db.academic_catalog(session, actor.institution_id)


@app.post("/api/admin/departments", status_code=status.HTTP_201_CREATED)
def create_department(payload: DepartmentCreate, actor: User = Depends(admin_only), session: Session = Depends(db.get_db)):
    row = Department(institution_id=actor.institution_id, name=payload.name.strip(), code=payload.code.upper().strip(), email_domain=payload.email_domain)
    session.add(row)
    flush_or_conflict(session, "Department code or email domain already exists")
    db.write_audit(session, actor.id, "department.created", "department", row.id, payload.model_dump())
    session.commit()
    return {"id": row.id, "name": row.name, "code": row.code, "email_domain": row.email_domain}


@app.patch("/api/admin/departments/{department_id}")
def update_department(department_id: int, payload: DepartmentCreate, actor: User = Depends(admin_only), session: Session = Depends(db.get_db)):
    row = department_for_student(session, actor.institution_id, department_id)
    old_name, old_code = row.name, row.code
    new_name, new_code = payload.name.strip(), payload.code.upper().strip()
    row.name = new_name
    row.code = new_code
    row.email_domain = payload.email_domain
    flush_or_conflict(session, "Department code or email domain already exists")
    session.query(ClassSection).filter(
        ClassSection.institution_id == actor.institution_id,
        func.lower(ClassSection.department).in_([old_name.casefold(), old_code.casefold()]),
    ).update({ClassSection.department: new_code}, synchronize_session=False)
    session.query(Student).filter(Student.institution_id == actor.institution_id, Student.department_id == row.id).update({Student.department: new_name}, synchronize_session=False)
    db.write_audit(session, actor.id, "department.updated", "department", row.id, {**payload.model_dump(), "previous_name": old_name, "previous_code": old_code})
    session.commit()
    return {"id": row.id, "name": row.name, "code": row.code, "email_domain": row.email_domain}


@app.post("/api/admin/academics/terms", status_code=status.HTTP_201_CREATED)
def create_academic_term(payload: AcademicTermCreate, actor: User = Depends(admin_only), session: Session = Depends(db.get_db)):
    if payload.end_date <= payload.start_date:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "Term end date must be after its start date")
    row = AcademicTerm(institution_id=actor.institution_id, **payload.model_dump())
    session.add(row)
    flush_or_conflict(session, "Academic term could not be created")
    db.write_audit(session, actor.id, "academic_term.created", "academic_term", row.id, payload.model_dump(mode="json"))
    session.commit()
    return {"id": row.id, **payload.model_dump(mode="json")}


@app.patch("/api/admin/academics/terms/{term_id}")
def update_academic_term(term_id: int, payload: AcademicTermUpdate, actor: User = Depends(admin_only), session: Session = Depends(db.get_db)):
    row = session.get(AcademicTerm, term_id)
    if not row or row.institution_id != actor.institution_id:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Academic term not found")
    if payload.end_date <= payload.start_date:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "Term end date must be after its start date")
    for key, value in payload.model_dump().items():
        setattr(row, key, value)
    flush_or_conflict(session, "Another academic term already uses this name and year")
    db.write_audit(session, actor.id, "academic_term.updated", "academic_term", row.id, payload.model_dump(mode="json"))
    session.commit()
    return {"id": row.id, **payload.model_dump(mode="json")}


@app.post("/api/admin/academics/sections", status_code=status.HTTP_201_CREATED)
def create_class_section(payload: ClassSectionCreate, actor: User = Depends(admin_only), session: Session = Depends(db.get_db)):
    department = session.scalar(select(Department).where(Department.institution_id == actor.institution_id, Department.active.is_(True), func.lower(Department.code) == payload.department.strip().lower()))
    if not department:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "Select an existing department before creating its class section")
    values = payload.model_dump()
    values["department"] = department.code
    row = ClassSection(institution_id=actor.institution_id, **values)
    session.add(row)
    flush_or_conflict(session, "This class section already exists for the selected batch")
    db.write_audit(session, actor.id, "class_section.created", "class_section", row.id, payload.model_dump())
    session.commit()
    return {"id": row.id, **payload.model_dump()}


@app.post("/api/admin/academics/assignments", status_code=status.HTTP_201_CREATED)
def create_teacher_assignment(payload: TeacherAssignmentCreate, actor: User = Depends(allocation_manager), session: Session = Depends(db.get_db)):
    teacher = session.get(User, payload.teacher_id)
    course = session.get(Course, payload.course_id)
    section = session.get(ClassSection, payload.section_id)
    term = session.get(AcademicTerm, payload.term_id)
    if not teacher or teacher.role != "teacher" or teacher.institution_id != actor.institution_id or not course or course.institution_id != actor.institution_id or not section or section.institution_id != actor.institution_id or not term or term.institution_id != actor.institution_id:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Teacher, course, section or term not found")
    current = session.scalar(select(TeacherAssignment).where(TeacherAssignment.course_id == payload.course_id, TeacherAssignment.section_id == payload.section_id, TeacherAssignment.term_id == payload.term_id, TeacherAssignment.active.is_(True)))
    if current:
        current_teacher = session.get(User, current.teacher_id)
        raise HTTPException(status.HTTP_409_CONFLICT, f"{current_teacher.name if current_teacher else 'A teacher'} is already assigned. Use Replace teacher for a mid-term change.")
    row = TeacherAssignment(**payload.model_dump())
    session.add(row)
    flush_or_conflict(session, "This teacher is already assigned to that course and section")
    members = session.scalars(select(SectionMembership).where(SectionMembership.section_id == payload.section_id, SectionMembership.term_id == payload.term_id, SectionMembership.active.is_(True))).all()
    for membership in members:
        if not session.scalar(select(Enrollment.id).where(Enrollment.student_id == membership.student_id, Enrollment.course_id == payload.course_id)):
            session.add(Enrollment(student_id=membership.student_id, course_id=payload.course_id))
    db.write_audit(session, actor.id, "teacher_assignment.created", "teacher_assignment", row.id, payload.model_dump())
    session.commit()
    return {"id": row.id, **payload.model_dump()}


@app.patch("/api/admin/academics/assignments/{assignment_id}/replace")
def replace_teacher_assignment(assignment_id: int, payload: TeacherReplacement, actor: User = Depends(allocation_manager), session: Session = Depends(db.get_db)):
    current = session.get(TeacherAssignment, assignment_id)
    current_course = session.get(Course, current.course_id) if current else None
    current_section = session.get(ClassSection, current.section_id) if current else None
    if not current or not current_course or not current_section or current_course.institution_id != actor.institution_id or current_section.institution_id != actor.institution_id:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Teacher assignment not found")
    if not current.active:
        raise HTTPException(status.HTTP_409_CONFLICT, "This teaching assignment has already ended")
    replacement = session.get(User, payload.new_teacher_id)
    if not replacement or replacement.institution_id != actor.institution_id or replacement.role != "teacher" or not replacement.active:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Replacement teacher not found")
    if replacement.id == current.teacher_id:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "Select a different teacher")
    now = datetime.now(timezone.utc)
    current.active = False
    current.ended_at = now
    next_assignment = session.scalar(select(TeacherAssignment).where(TeacherAssignment.teacher_id == replacement.id, TeacherAssignment.course_id == current.course_id, TeacherAssignment.section_id == current.section_id, TeacherAssignment.term_id == current.term_id))
    if next_assignment:
        next_assignment.active = True
        next_assignment.started_at = now
        next_assignment.ended_at = None
    else:
        next_assignment = TeacherAssignment(teacher_id=replacement.id, course_id=current.course_id, section_id=current.section_id, term_id=current.term_id, started_at=now)
        session.add(next_assignment)
        session.flush()
    session.query(TimetableSlot).filter(TimetableSlot.assignment_id == current.id, TimetableSlot.effective_to >= campus_today()).update({TimetableSlot.assignment_id: next_assignment.id}, synchronize_session=False)
    session.query(ClassSession).filter(ClassSession.assignment_id == current.id, ClassSession.session_date >= campus_today(), ClassSession.status == "scheduled").update({ClassSession.assignment_id: next_assignment.id, ClassSession.created_by: replacement.id}, synchronize_session=False)
    session.query(Assessment).filter(Assessment.teacher_assignment_id == current.id, Assessment.status.in_(["draft", "published"])).update({Assessment.teacher_assignment_id: next_assignment.id, Assessment.created_by: replacement.id}, synchronize_session=False)
    session.query(LearningAssignment).filter(LearningAssignment.teacher_assignment_id == current.id, LearningAssignment.status.in_(["draft", "published"])).update({LearningAssignment.teacher_assignment_id: next_assignment.id, LearningAssignment.created_by: replacement.id}, synchronize_session=False)
    db.write_audit(session, actor.id, "teacher_assignment.replaced", "teacher_assignment", next_assignment.id, {"previous_assignment_id": current.id, "previous_teacher_id": current.teacher_id, "new_teacher_id": replacement.id, "reason": payload.reason})
    session.commit()
    return {"id": next_assignment.id, "teacher_id": replacement.id, "previous_assignment_id": current.id, "status": "active"}


@app.post("/api/admin/academics/section-memberships", status_code=status.HTTP_201_CREATED)
def create_section_membership(payload: SectionMembershipCreate, actor: User = Depends(admin_only), session: Session = Depends(db.get_db)):
    student = session.get(Student, payload.student_id)
    section = session.get(ClassSection, payload.section_id)
    term = session.get(AcademicTerm, payload.term_id)
    if not student or student.institution_id != actor.institution_id or not section or section.institution_id != actor.institution_id or not term or term.institution_id != actor.institution_id:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Student, section or term not found")
    current_rows = session.scalars(select(SectionMembership).where(SectionMembership.student_id == payload.student_id, SectionMembership.term_id == payload.term_id, SectionMembership.active.is_(True))).all()
    if any(row.section_id == payload.section_id for row in current_rows):
        raise HTTPException(status.HTTP_409_CONFLICT, "Student is already assigned to this class for the selected term")
    for current in current_rows:
        current.active = False
    row = session.scalar(select(SectionMembership).where(SectionMembership.student_id == payload.student_id, SectionMembership.section_id == payload.section_id, SectionMembership.term_id == payload.term_id))
    if row:
        row.active = True
    else:
        row = SectionMembership(**payload.model_dump())
        session.add(row)
        session.flush()
    enrolled_courses = ensure_section_enrollments(session, payload.student_id, payload.section_id, payload.term_id)
    db.write_audit(session, actor.id, "section_membership.assigned", "section_membership", row.id, {**payload.model_dump(), "previous_memberships_deactivated": len(current_rows)})
    session.commit()
    return {"id": row.id, **payload.model_dump(), "moved_from_previous_class": bool(current_rows), "courses_enrolled": enrolled_courses}


@app.post("/api/admin/academics/timetable", status_code=status.HTTP_201_CREATED)
def create_timetable_slot(payload: TimetableSlotCreate, actor: User = Depends(allocation_manager), session: Session = Depends(db.get_db)):
    assignment = timetable_access(session, actor, payload.assignment_id)
    payload = semester_dates(session, assignment, payload)
    if payload.end_time <= payload.start_time or payload.effective_to < payload.effective_from:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "Check timetable times and effective dates")
    ensure_timetable_has_no_conflict(session, assignment, payload)
    row = TimetableSlot(section_id=assignment.section_id, publication_status="published", proposed_by=actor.id, approved_by=actor.id, approved_at=datetime.now(timezone.utc), **payload.model_dump())
    session.add(row)
    flush_or_conflict(session, "This section already has a class in that period")
    generated = generate_timetable_sessions(session, row, actor.id)
    db.write_audit(session, actor.id, "timetable.created", "timetable_slot", row.id, {**payload.model_dump(mode="json"), "sessions_generated": generated})
    session.commit()
    return {"id": row.id, "sessions_generated": generated}


@app.post("/api/academics/timetable/proposals", status_code=status.HTTP_201_CREATED)
def propose_timetable_slot(payload: TimetableSlotCreate, actor: User = Depends(teacher_only), session: Session = Depends(db.get_db)):
    if actor.role != "teacher":
        raise HTTPException(status.HTTP_403_FORBIDDEN, "Teacher access is required for timetable proposals")
    assignment = timetable_access(session, actor, payload.assignment_id)
    payload = semester_dates(session, assignment, payload)
    if payload.end_time <= payload.start_time or payload.effective_to < payload.effective_from:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "Check timetable times and effective dates")
    ensure_timetable_has_no_conflict(session, assignment, payload)
    row = TimetableSlot(section_id=assignment.section_id, publication_status="pending", proposed_by=actor.id, **payload.model_dump())
    session.add(row)
    flush_or_conflict(session, "This section already has a timetable entry for that period and start date")
    db.write_audit(session, actor.id, "timetable.proposed", "timetable_slot", row.id, payload.model_dump(mode="json"))
    session.commit()
    return {"id": row.id, "publication_status": row.publication_status, "message": "Timetable proposal sent for administrator approval"}


@app.patch("/api/admin/academics/timetable/{slot_id}/review")
def review_timetable_slot(slot_id: int, payload: TimetableReview, actor: User = Depends(allocation_manager), session: Session = Depends(db.get_db)):
    row = session.get(TimetableSlot, slot_id)
    assignment = session.get(TeacherAssignment, row.assignment_id) if row else None
    course = session.get(Course, assignment.course_id) if assignment else None
    if not row or not assignment or not course or course.institution_id != actor.institution_id:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Timetable proposal not found")
    if row.publication_status != "pending":
        return {"id": row.id, "publication_status": row.publication_status, "sessions_generated": 0}
    if payload.decision == "rejected" and not (payload.reason or "").strip():
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "A rejection reason is required")
    generated = 0
    if payload.decision == "approved":
        timetable_access(session, actor, assignment.id)
        candidate = TimetableSlotCreate(assignment_id=row.assignment_id, weekday=row.weekday, period_number=row.period_number, start_time=row.start_time, end_time=row.end_time, classroom=row.classroom, effective_from=row.effective_from, effective_to=row.effective_to)
        ensure_timetable_has_no_conflict(session, assignment, candidate, exclude_slot_id=row.id)
        row.publication_status = "published"
        row.approved_by = actor.id
        row.approved_at = datetime.now(timezone.utc)
        generated = generate_timetable_sessions(session, row, actor.id)
    else:
        row.publication_status = "rejected"
        row.active = False
    row.review_reason = (payload.reason or "").strip() or None
    db.write_audit(session, actor.id, f"timetable.{payload.decision}", "timetable_slot", row.id, {"reason": row.review_reason, "sessions_generated": generated})
    session.commit()
    return {"id": row.id, "publication_status": row.publication_status, "sessions_generated": generated}


@app.post("/api/admin/academics/syllabus/units", status_code=status.HTTP_201_CREATED)
def create_syllabus_unit(payload: SyllabusUnitCreate, actor: User = Depends(admin_only), session: Session = Depends(db.get_db)):
    course = session.get(Course, payload.course_id)
    if not course or course.institution_id != actor.institution_id:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Course not found")
    row = SyllabusUnit(**payload.model_dump())
    session.add(row)
    flush_or_conflict(session, "This unit sequence number already exists for the course")
    db.write_audit(session, actor.id, "syllabus_unit.created", "syllabus_unit", row.id, payload.model_dump())
    session.commit()
    return {"id": row.id, **payload.model_dump()}


@app.post("/api/admin/academics/syllabus/topics", status_code=status.HTTP_201_CREATED)
def create_syllabus_topic(payload: SyllabusTopicCreate, actor: User = Depends(admin_only), session: Session = Depends(db.get_db)):
    unit = session.get(SyllabusUnit, payload.unit_id)
    course = session.get(Course, unit.course_id) if unit else None
    if not unit or not course or course.institution_id != actor.institution_id:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Syllabus unit not found")
    row = SyllabusTopic(**payload.model_dump())
    session.add(row)
    flush_or_conflict(session, "This topic sequence number already exists in the unit")
    db.write_audit(session, actor.id, "syllabus_topic.created", "syllabus_topic", row.id, payload.model_dump(mode="json"))
    session.commit()
    return {"id": row.id, **payload.model_dump(mode="json"), "status": row.status}


@app.post("/api/academics/sessions", status_code=status.HTTP_201_CREATED)
def create_class_session(payload: ClassSessionCreate, actor: User = Depends(teacher_only), session: Session = Depends(db.get_db)):
    assignment = require_assignment_access(session, actor, payload.assignment_id)
    if payload.timetable_slot_id:
        timetable_slot = session.get(TimetableSlot, payload.timetable_slot_id)
        if not timetable_slot or timetable_slot.assignment_id != assignment.id or timetable_slot.section_id != assignment.section_id:
            raise HTTPException(status.HTTP_404_NOT_FOUND, "Timetable slot not found for this class")
    row = ClassSession(section_id=assignment.section_id, created_by=actor.id, **payload.model_dump())
    session.add(row)
    flush_or_conflict(session, "This section already has a recorded class for that period")
    db.write_audit(session, actor.id, "class_session.created", "class_session", row.id, payload.model_dump(mode="json"))
    session.commit()
    return {"id": row.id, **payload.model_dump(mode="json"), "status": row.status}


@app.patch("/api/academics/sessions/{class_session_id}/coverage")
def record_topic_coverage(class_session_id: int, payload: TopicCoverageUpdate, actor: User = Depends(teacher_only), session: Session = Depends(db.get_db)):
    class_session = session.get(ClassSession, class_session_id)
    if not class_session:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Class session not found")
    assignment = require_assignment_access(session, actor, class_session.assignment_id)
    topic = session.get(SyllabusTopic, payload.topic_id)
    unit = session.get(SyllabusUnit, topic.unit_id) if topic else None
    if not topic or not unit or unit.course_id != assignment.course_id:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "Topic does not belong to this class subject")
    if payload.status == "completed" and payload.coverage_percent != 100:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "Completed topics must have 100% coverage")
    if payload.status == "postponed" and payload.coverage_percent != 0:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "Postponed topics must have 0% coverage")
    coverage = session.scalar(select(TopicCoverage).where(TopicCoverage.class_session_id == class_session.id, TopicCoverage.topic_id == topic.id))
    if not coverage:
        coverage = TopicCoverage(class_session_id=class_session.id, topic_id=topic.id, recorded_by=actor.id, status=payload.status, coverage_percent=payload.coverage_percent)
        session.add(coverage)
    coverage.status = payload.status
    coverage.coverage_percent = payload.coverage_percent
    coverage.notes = payload.notes
    coverage.recorded_by = actor.id
    coverage.recorded_at = datetime.now(timezone.utc)
    topic.status = payload.status
    topic.coverage_percent = payload.coverage_percent
    topic.completed_at = datetime.now(timezone.utc) if payload.status == "completed" else None
    class_session.status = payload.class_status
    class_session.teacher_notes = payload.notes
    class_session.actual_start = class_session.actual_start or datetime.now(timezone.utc)
    class_session.actual_end = datetime.now(timezone.utc)
    db.write_audit(session, actor.id, "topic_coverage.recorded", "class_session", class_session.id, payload.model_dump())
    session.commit()
    return {"class_session_id": class_session.id, "topic_id": topic.id, "topic_status": topic.status, "coverage_percent": topic.coverage_percent}


@app.post("/api/academics/sessions/{class_session_id}/attendance/{student_id}/override")
def teacher_attendance_override(class_session_id: int, student_id: int, payload: AttendanceOverride, actor: User = Depends(teacher_only), session: Session = Depends(db.get_db)):
    class_session = session.get(ClassSession, class_session_id)
    student = session.get(Student, student_id)
    if not class_session or not student or student.institution_id != actor.institution_id:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Class session or student not found")
    assignment = require_assignment_access(session, actor, class_session.assignment_id)
    membership = session.scalar(select(SectionMembership.id).where(SectionMembership.student_id == student_id, SectionMembership.section_id == class_session.section_id, SectionMembership.term_id == assignment.term_id, SectionMembership.active.is_(True)))
    if not membership:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "Student is not assigned to this class and academic term")
    result = db.correct_session_attendance(session, class_session, student_id, "Present", payload.reason.strip(), actor.id, payload.decision_source)
    session.commit()
    return result


@app.get("/api/academics/sessions/{class_session_id}/attendance/preview")
def class_attendance_preview(class_session_id: int, actor: User = Depends(teacher_only), session: Session = Depends(db.get_db)):
    class_session = session.get(ClassSession, class_session_id)
    if not class_session:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Class session not found")
    require_assignment_access(session, actor, class_session.assignment_id)
    return {"class_session_id": class_session.id, "status": class_session.status, "attendance_state": class_session.attendance_state, "camera_state": class_session.camera_state, "failure_reason": class_session.failure_reason, "students": db.attendance_preview(session, class_session)}


@app.patch("/api/academics/attendance/raised-hand/{request_id}")
def review_raised_hand_attendance(request_id: int, payload: RaisedHandReview, actor: User = Depends(teacher_only), session: Session = Depends(db.get_db)):
    row = session.get(RaisedHandAttendanceRequest, request_id)
    class_session = session.get(ClassSession, row.class_session_id) if row else None
    if not row or not class_session or row.institution_id != actor.institution_id:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Raised-hand attendance request not found")
    require_assignment_access(session, actor, class_session.assignment_id)
    if row.status != "pending":
        raise HTTPException(status.HTTP_409_CONFLICT, "This raised-hand request was already reviewed")
    row.status = payload.decision
    row.reviewed_by = actor.id
    row.reviewed_at = datetime.now(timezone.utc)
    row.review_reason = (payload.reason or "").strip() or None
    if payload.decision == "approved":
        db.correct_session_attendance(session, class_session, row.student_id, "Present", payload.reason or "Teacher approved the camera-detected raised-hand attendance request", actor.id, "approved_exit")
    db.write_audit(session, actor.id, f"attendance.raised_hand_{payload.decision}", "raised_hand_request", row.id, {"student_id": row.student_id, "class_session_id": class_session.id})
    session.commit()
    return {"id": row.id, "status": row.status}


@app.post("/api/academics/sessions/{class_session_id}/attendance/finalize")
def finalize_class_attendance(class_session_id: int, actor: User = Depends(teacher_only), session: Session = Depends(db.get_db)):
    class_session = session.get(ClassSession, class_session_id)
    if not class_session:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Class session not found")
    require_assignment_access(session, actor, class_session.assignment_id)
    if class_session.attendance_state == "published":
        return {"class_session_id": class_session_id, "attendance_state": "published", "students": db.attendance_preview(session, class_session)}
    pending = session.scalar(select(RaisedHandAttendanceRequest.id).where(RaisedHandAttendanceRequest.class_session_id == class_session.id, RaisedHandAttendanceRequest.status == "pending"))
    if pending:
        raise HTTPException(status.HTTP_409_CONFLICT, "Review all raised-hand requests before finalising attendance")
    rows = db.finalize_session_attendance(session, class_session, actor.id)
    session.commit()
    return {"class_session_id": class_session_id, "attendance_state": class_session.attendance_state, "students": rows}


@app.patch("/api/academics/sessions/{class_session_id}/attendance/{student_id}/correct")
def correct_class_attendance(class_session_id: int, student_id: int, payload: AttendanceCorrectionCreate, actor: User = Depends(teacher_only), session: Session = Depends(db.get_db)):
    class_session = session.get(ClassSession, class_session_id)
    student = session.get(Student, student_id)
    if not class_session or not student or student.institution_id != actor.institution_id:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Class session or student not found")
    require_assignment_access(session, actor, class_session.assignment_id)
    try:
        result = db.correct_session_attendance(session, class_session, student_id, payload.status, payload.reason.strip(), actor.id)
    except ValueError as exc:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, str(exc)) from exc
    session.commit()
    return result


@app.post("/api/academics/sessions/{class_session_id}/classroom-change", status_code=status.HTTP_201_CREATED)
def change_classroom(class_session_id: int, payload: ClassroomChangeCreate, actor: User = Depends(teacher_only), session: Session = Depends(db.get_db)):
    class_session = session.get(ClassSession, class_session_id)
    if not class_session:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Class session not found")
    require_assignment_access(session, actor, class_session.assignment_id)
    change = db.record_classroom_change(session, class_session, payload.classroom.strip(), payload.reason.strip(), actor.id, actor.institution_id)
    session.commit()
    return {"id": change.id, "class_session_id": class_session.id, "classroom": class_session.temporary_classroom, "camera_state": class_session.camera_state}


@app.get("/api/monitor/cameras")
def camera_status(actor: User = Depends(teacher_only), session: Session = Depends(db.get_db)):
    return {"cameras": db.camera_status_for_institution(session, actor.institution_id)}


@app.post("/api/academics/coursework", status_code=status.HTTP_201_CREATED)
def create_coursework(payload: LearningAssignmentCreate, actor: User = Depends(teacher_only), session: Session = Depends(db.get_db)):
    teaching = require_assignment_access(session, actor, payload.assignment_id)
    due_at = payload.due_at if payload.due_at.tzinfo else payload.due_at.replace(tzinfo=timezone.utc)
    if due_at <= datetime.now(timezone.utc):
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "Assignment due date must be in the future")
    row = LearningAssignment(teacher_assignment_id=teaching.id, course_id=teaching.course_id, section_id=teaching.section_id, term_id=teaching.term_id, title=payload.title.strip(), instructions=payload.instructions.strip(), due_at=payload.due_at, max_marks=payload.max_marks, created_by=actor.id)
    session.add(row)
    session.flush()
    db.write_audit(session, actor.id, "coursework.created", "learning_assignment", row.id, payload.model_dump(mode="json"))
    session.commit()
    return {"id": row.id, "status": row.status}


@app.patch("/api/academics/coursework/{coursework_id}/publish")
def publish_coursework(coursework_id: int, actor: User = Depends(teacher_only), session: Session = Depends(db.get_db)):
    row = require_coursework_access(session, actor, coursework_id)
    if row.status != "draft":
        raise HTTPException(status.HTTP_409_CONFLICT, "Only a draft assignment can be published")
    row.status = "published"
    notified = notify_section_students(
        session,
        institution_id=actor.institution_id,
        section_id=row.section_id,
        term_id=row.term_id,
        template="coursework_published",
        message=f"New assignment: {row.title}. Due {row.due_at.strftime('%d %b %Y, %H:%M')}.",
        dedup_prefix=f"coursework:{row.id}:published",
        deep_link="/portal?view=assignments",
    )
    db.write_audit(session, actor.id, "coursework.published", "learning_assignment", row.id, {"notified": notified})
    session.commit()
    return {"id": row.id, "status": row.status}


@app.post("/api/student/coursework/{coursework_id}/submit", status_code=status.HTTP_201_CREATED)
def submit_coursework(coursework_id: int, payload: LearningSubmissionCreate, user: User = Depends(current_user), session: Session = Depends(db.get_db)):
    if user.role != "student":
        raise HTTPException(status.HTTP_403_FORBIDDEN, "Student access is required")
    student = session.scalar(select(Student).where(Student.user_id == user.id))
    row = session.get(LearningAssignment, coursework_id)
    course = session.get(Course, row.course_id) if row else None
    if not student or not row or not course or course.institution_id != user.institution_id or row.status != "published":
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Published assignment not found")
    membership = session.scalar(select(SectionMembership.id).where(SectionMembership.student_id == student.id, SectionMembership.section_id == row.section_id, SectionMembership.term_id == row.term_id, SectionMembership.active.is_(True)))
    if not membership:
        raise HTTPException(status.HTTP_403_FORBIDDEN, "This assignment is not for your class and term")
    if session.scalar(select(LearningSubmission.id).where(LearningSubmission.learning_assignment_id == row.id, LearningSubmission.student_id == student.id)):
        raise HTTPException(status.HTTP_409_CONFLICT, "This assignment was already finally submitted")
    submitted_at = datetime.now(timezone.utc)
    due_at = row.due_at if row.due_at.tzinfo else row.due_at.replace(tzinfo=timezone.utc)
    if submitted_at > due_at:
        raise HTTPException(status.HTTP_409_CONFLICT, "The assignment deadline has passed. Ask the subject teacher for an audited extension.")
    submission = LearningSubmission(learning_assignment_id=row.id, student_id=student.id, answer_text=payload.answer_text.strip(), submitted_at=submitted_at, status="submitted")
    session.add(submission)
    session.flush()
    teaching = session.get(TeacherAssignment, row.teacher_assignment_id)
    if teaching:
        db.add_notification(session, institution_id=user.institution_id, user_id=teaching.teacher_id, student_id=None, template="coursework_submitted", message=f"{user.name} submitted {row.title}.", dedup_key=f"coursework-submission:{submission.id}:{teaching.teacher_id}", deep_link="/portal?view=academics")
    db.write_audit(session, user.id, "coursework.submitted", "learning_submission", submission.id, {"coursework_id": row.id, "late": False})
    session.commit()
    return {"id": submission.id, "status": submission.status, "submitted_at": submission.submitted_at.isoformat(), "forwarded_to_teacher_assignment_id": row.teacher_assignment_id}


def valid_assignment_file(file_name: str, content_type: str | None, payload: bytes) -> bool:
    suffix = Path(file_name).suffix.lower()
    if suffix == ".pdf" and content_type == "application/pdf":
        return payload.startswith(b"%PDF-")
    if suffix in {".jpg", ".jpeg"} and content_type in {"image/jpeg", "image/jpg"}:
        return payload.startswith(b"\xff\xd8\xff")
    if suffix == ".png" and content_type == "image/png":
        return payload.startswith(b"\x89PNG\r\n\x1a\n")
    return False


@app.post("/api/student/coursework/{coursework_id}/submit-files", status_code=status.HTTP_201_CREATED)
async def submit_coursework_files(coursework_id: int, answer_text: str = Form(""), files: list[UploadFile] = File(default=[]), user: User = Depends(current_user), session: Session = Depends(db.get_db)):
    if user.role != "student":
        raise HTTPException(status.HTTP_403_FORBIDDEN, "Student access is required")
    if files and not fixed_window_allow(f"upload:assignment:{user.id}", 20, 3600):
        raise HTTPException(status.HTTP_429_TOO_MANY_REQUESTS, "Too many assignment uploads. Try again later.")
    if not answer_text.strip() and not files:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "Enter an answer or attach handwritten work")
    if len(files) > settings.assignment_max_files:
        raise HTTPException(status.HTTP_413_REQUEST_ENTITY_TOO_LARGE, f"Attach no more than {settings.assignment_max_files} files")
    student = session.scalar(select(Student).where(Student.user_id == user.id))
    row = session.get(LearningAssignment, coursework_id)
    course = session.get(Course, row.course_id) if row else None
    if not student or not row or not course or course.institution_id != user.institution_id or row.status != "published":
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Published assignment not found")
    membership = session.scalar(select(SectionMembership.id).where(SectionMembership.student_id == student.id, SectionMembership.section_id == row.section_id, SectionMembership.term_id == row.term_id, SectionMembership.active.is_(True)))
    if not membership:
        raise HTTPException(status.HTTP_403_FORBIDDEN, "This assignment is not for your class and term")
    if session.scalar(select(LearningSubmission.id).where(LearningSubmission.learning_assignment_id == row.id, LearningSubmission.student_id == student.id)):
        raise HTTPException(status.HTTP_409_CONFLICT, "This assignment was already finally submitted")
    now = datetime.now(timezone.utc)
    due_at = row.due_at if row.due_at.tzinfo else row.due_at.replace(tzinfo=timezone.utc)
    if now > due_at:
        raise HTTPException(status.HTTP_409_CONFLICT, "The assignment deadline has passed. Ask the subject teacher for an audited extension.")
    prepared: list[tuple[UploadFile, bytes]] = []
    total_size = 0
    for uploaded in files:
        payload = await uploaded.read(settings.max_upload_mb * 1024 * 1024 + 1)
        total_size += len(payload)
        if total_size > settings.max_upload_mb * 1024 * 1024:
            raise HTTPException(status.HTTP_413_REQUEST_ENTITY_TOO_LARGE, f"Attachments exceed the combined {settings.max_upload_mb} MB limit")
        if not valid_assignment_file(uploaded.filename or "", uploaded.content_type, payload):
            raise HTTPException(status.HTTP_415_UNSUPPORTED_MEDIA_TYPE, "Only genuine PDF, JPG and PNG files are accepted")
        prepared.append((uploaded, payload))
    submission = LearningSubmission(learning_assignment_id=row.id, student_id=student.id, answer_text=answer_text.strip(), submitted_at=now, status="submitted")
    session.add(submission)
    session.flush()
    stored_keys: list[str] = []
    replaced_photo_keys: list[str] = []
    try:
        for uploaded, payload in prepared:
            suffix = Path(uploaded.filename or "file").suffix.lower()
            storage_name = f"institutions/{user.institution_id}/assignments/{coursework_id}/submissions/{submission.id}/{uuid4().hex}{suffix}"
            try:
                stored_object = await run_in_threadpool(storage_service.put_private, storage_name, payload, uploaded.content_type or "application/octet-stream")
            except ValueError as exc:
                raise HTTPException(status.HTTP_422_UNPROCESSABLE_ENTITY, str(exc)) from exc
            except Exception as exc:
                logger.exception("Assignment attachment storage failed")
                raise HTTPException(status.HTTP_503_SERVICE_UNAVAILABLE, "Secure upload scanning or storage is unavailable") from exc
            stored_keys.append(stored_object.key)
            safe_name = Path(uploaded.filename or f"attachment{suffix}").name[:255]
            session.add(LearningSubmissionAttachment(submission_id=submission.id, file_name=safe_name, storage_name=stored_object.key, mime_type=uploaded.content_type or "application/octet-stream", size_bytes=len(payload), checksum_sha256=hashlib.sha256(payload).hexdigest()))
        teaching = session.get(TeacherAssignment, row.teacher_assignment_id)
        if teaching:
            db.add_notification(session, institution_id=user.institution_id, user_id=teaching.teacher_id, student_id=None, template="coursework_submitted", message=f"{user.name} submitted {row.title} with {len(prepared)} attachment(s).", dedup_key=f"coursework-submission:{submission.id}:{teaching.teacher_id}", deep_link="/portal?view=academics")
        db.write_audit(session, user.id, "coursework.submitted", "learning_submission", submission.id, {"coursework_id": row.id, "attachment_count": len(prepared), "late": False})
        session.commit()
    except Exception:
        session.rollback()
        for key in stored_keys:
            try:
                storage_service.delete_private(key)
            except Exception:
                logger.exception("Failed to remove orphaned assignment attachment after transaction rollback")
        raise
    return {"id": submission.id, "status": submission.status, "submitted_at": submission.submitted_at.isoformat(), "attachment_count": len(prepared), "forwarded_to_teacher_assignment_id": row.teacher_assignment_id}


@app.get("/api/academics/coursework/{coursework_id}/submissions")
def coursework_submissions(coursework_id: int, actor: User = Depends(teacher_only), session: Session = Depends(db.get_db)):
    row = require_coursework_access(session, actor, coursework_id)
    submissions = session.execute(select(LearningSubmission, Student.roll_no, User.name).join(Student, Student.id == LearningSubmission.student_id).join(User, User.id == Student.user_id).where(LearningSubmission.learning_assignment_id == row.id).order_by(LearningSubmission.submitted_at.desc())).all()
    return {"assignment_id": row.id, "title": row.title, "max_marks": row.max_marks, "submissions": [{"id": submission.id, "student_id": submission.student_id, "student_name": name, "roll_no": roll_no, "answer_text": submission.answer_text, "submitted_at": submission.submitted_at.isoformat(), "status": submission.status, "marks": submission.marks, "feedback": submission.feedback, "attachments": [{"id": attachment.id, "file_name": attachment.file_name, "mime_type": attachment.mime_type, "size_bytes": attachment.size_bytes} for attachment in session.scalars(select(LearningSubmissionAttachment).where(LearningSubmissionAttachment.submission_id == submission.id).order_by(LearningSubmissionAttachment.id)).all()]} for submission, roll_no, name in submissions]}


@app.get("/api/academics/coursework/attachments/{attachment_id}")
def download_coursework_attachment(attachment_id: int, actor: User = Depends(current_user), session: Session = Depends(db.get_db)):
    attachment = session.get(LearningSubmissionAttachment, attachment_id)
    submission = session.get(LearningSubmission, attachment.submission_id) if attachment else None
    coursework = session.get(LearningAssignment, submission.learning_assignment_id) if submission else None
    if not attachment or not submission or not coursework:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Attachment not found")
    allowed = False
    if actor.role == "student":
        student = session.scalar(select(Student).where(Student.user_id == actor.id))
        allowed = bool(student and submission.student_id == student.id)
    elif actor.role == "admin":
        course = session.get(Course, coursework.course_id)
        allowed = bool(course and course.institution_id == actor.institution_id)
    else:
        teaching = session.get(TeacherAssignment, coursework.teacher_assignment_id)
        allowed = bool(teaching and teaching.teacher_id == actor.id and teaching.active)
    if not allowed:
        raise HTTPException(status.HTTP_403_FORBIDDEN, "You cannot access this attachment")
    expected_prefix = f"institutions/{actor.institution_id}/assignments/{coursework.id}/submissions/{submission.id}/"
    if not attachment.storage_name.startswith(expected_prefix):
        # Legacy v8.1 filenames are served only from the fixed submissions root.
        legacy_name = Path(attachment.storage_name).name
        legacy = (SUBMISSION_ROOT / legacy_name).resolve()
        if legacy.parent == SUBMISSION_ROOT.resolve() and legacy.is_file():
            return FileResponse(legacy, media_type=attachment.mime_type, filename=attachment.file_name)
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Attachment file is unavailable")
    local = storage_service.local_path(attachment.storage_name)
    if local:
        return FileResponse(local, media_type=attachment.mime_type, filename=attachment.file_name)
    return {"url": storage_service.download_url(attachment.storage_name, 300), "expires_in": 300}


@app.patch("/api/academics/coursework/submissions/{submission_id}/grade")
def grade_coursework(submission_id: int, payload: LearningSubmissionGrade, actor: User = Depends(teacher_only), session: Session = Depends(db.get_db)):
    submission = session.get(LearningSubmission, submission_id)
    coursework = require_coursework_access(session, actor, submission.learning_assignment_id) if submission else None
    if not submission or not coursework:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Assignment submission not found")
    if payload.marks > coursework.max_marks:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, f"Marks cannot exceed {coursework.max_marks}")
    submission.marks = payload.marks
    submission.feedback = (payload.feedback or "").strip() or None
    submission.status = "graded"
    submission.graded_by = actor.id
    submission.graded_at = datetime.now(timezone.utc)
    student = session.get(Student, submission.student_id)
    if student:
        db.add_notification(session, institution_id=actor.institution_id, user_id=student.user_id, student_id=student.id, template="coursework_graded", message=f"{coursework.title} was graded: {payload.marks:g}/{coursework.max_marks:g}.", dedup_key=f"coursework-grade:{submission.id}:{payload.marks}", deep_link="/portal?view=assignments")
    db.write_audit(session, actor.id, "coursework.graded", "learning_submission", submission.id, payload.model_dump())
    session.commit()
    return {"id": submission.id, "status": submission.status, "marks": submission.marks, "max_marks": coursework.max_marks}


@app.post("/api/academics/assessments", status_code=status.HTTP_201_CREATED)
def create_assessment(payload: AssessmentCreate, actor: User = Depends(teacher_only), session: Session = Depends(db.get_db)):
    course = session.get(Course, payload.course_id)
    section = session.get(ClassSection, payload.section_id)
    if not course or course.institution_id != actor.institution_id or not section or section.institution_id != actor.institution_id:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Course or class section not found for this institution")
    teaching_assignment = require_assignment_access(session, actor, payload.assignment_id) if payload.assignment_id else session.scalar(select(TeacherAssignment).where(TeacherAssignment.teacher_id == actor.id, TeacherAssignment.course_id == payload.course_id, TeacherAssignment.section_id == payload.section_id, TeacherAssignment.active.is_(True)).order_by(TeacherAssignment.id.desc()))
    if actor.role == "admin" and not teaching_assignment:
        teaching_assignment = session.scalar(select(TeacherAssignment).where(TeacherAssignment.course_id == payload.course_id, TeacherAssignment.section_id == payload.section_id, TeacherAssignment.active.is_(True)).order_by(TeacherAssignment.id.desc()))
    if actor.role != "admin":
        if not teaching_assignment:
            raise HTTPException(status.HTTP_403_FORBIDDEN, "You are not assigned to this course and section")
    if teaching_assignment and (teaching_assignment.course_id != payload.course_id or teaching_assignment.section_id != payload.section_id):
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "Selected teaching assignment does not match this course and section")
    topics = session.scalars(select(SyllabusTopic).join(SyllabusUnit, SyllabusUnit.id == SyllabusTopic.unit_id).where(SyllabusTopic.id.in_(payload.topic_ids), SyllabusUnit.course_id == payload.course_id)).all()
    if len(topics) != len(set(payload.topic_ids)) or any(topic.status != "completed" for topic in topics):
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "Assessments can use only completed topics from this course")
    if payload.mode == "online" and payload.question_count < 1:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "Enter the number of questions for an online assessment")
    if payload.mode == "online" and payload.marking_mode == "common" and not payload.common_marks:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "Enter common marks for every question")
    if payload.available_until and aware_utc(payload.available_until) <= aware_utc(payload.scheduled_at):
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "The closing time must be after the scheduled opening time")
    values = payload.model_dump(exclude={"topic_ids", "assignment_id"})
    if payload.mode == "online" and payload.marking_mode == "common":
        values["total_marks"] = payload.question_count * float(payload.common_marks or 0)
    row = Assessment(created_by=actor.id, teacher_assignment_id=teaching_assignment.id if teaching_assignment else None, **values)
    session.add(row)
    session.flush()
    session.add_all([AssessmentTopic(assessment_id=row.id, topic_id=topic_id) for topic_id in set(payload.topic_ids)])
    db.write_audit(session, actor.id, "assessment.created", "assessment", row.id, {**payload.model_dump(mode="json"), "topic_ids": payload.topic_ids})
    session.commit()
    return {"id": row.id, "status": row.status, **payload.model_dump(mode="json")}


@app.patch("/api/academics/assessments/{assessment_id}")
def update_assessment(assessment_id: int, payload: AssessmentUpdate, actor: User = Depends(teacher_only), session: Session = Depends(db.get_db)):
    assessment = require_assessment_access(session, actor, assessment_id)
    if assessment.status != "draft" or session.scalar(select(AssessmentAttempt.id).where(AssessmentAttempt.assessment_id == assessment.id).limit(1)):
        raise HTTPException(status.HTTP_409_CONFLICT, "Only an unattempted draft assessment can be edited")
    course = session.get(Course, payload.course_id)
    section = session.get(ClassSection, payload.section_id)
    teaching = require_assignment_access(session, actor, payload.assignment_id) if payload.assignment_id else None
    if not course or not section or course.institution_id != actor.institution_id or section.institution_id != actor.institution_id:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Course or class section not found")
    if actor.role != "admin" and not teaching:
        raise HTTPException(status.HTTP_403_FORBIDDEN, "Select one of your active teaching assignments")
    if teaching and (teaching.course_id != payload.course_id or teaching.section_id != payload.section_id):
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "Selected teaching assignment does not match this course and section")
    topics = session.scalars(select(SyllabusTopic).join(SyllabusUnit, SyllabusUnit.id == SyllabusTopic.unit_id).where(SyllabusTopic.id.in_(payload.topic_ids), SyllabusUnit.course_id == payload.course_id)).all()
    if len(topics) != len(set(payload.topic_ids)) or any(topic.status != "completed" for topic in topics):
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "Assessments can use only completed topics from this course")
    if payload.mode == "online" and payload.question_count < 1:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "Enter the number of questions for an online assessment")
    if payload.mode == "online" and payload.marking_mode == "common" and not payload.common_marks:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "Enter common marks for every question")
    if payload.available_until and aware_utc(payload.available_until) <= aware_utc(payload.scheduled_at):
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "The closing time must be after the scheduled opening time")
    for key, value in payload.model_dump(exclude={"topic_ids", "assignment_id"}).items():
        setattr(assessment, key, value)
    assessment.teacher_assignment_id = teaching.id if teaching else assessment.teacher_assignment_id
    if payload.mode == "online" and payload.marking_mode == "common":
        assessment.total_marks = payload.question_count * float(payload.common_marks or 0)
    session.query(AssessmentTopic).filter(AssessmentTopic.assessment_id == assessment.id).delete(synchronize_session=False)
    session.add_all([AssessmentTopic(assessment_id=assessment.id, topic_id=topic_id) for topic_id in set(payload.topic_ids)])
    db.write_audit(session, actor.id, "assessment.updated", "assessment", assessment.id, payload.model_dump(mode="json"))
    session.commit()
    return {"id": assessment.id, "status": assessment.status}


@app.patch("/api/academics/assessments/{assessment_id}/status")
def update_assessment_status(assessment_id: int, payload: AssessmentStatusUpdate, actor: User = Depends(teacher_only), session: Session = Depends(db.get_db)):
    assessment = require_assessment_access(session, actor, assessment_id)
    attempts_exist = bool(session.scalar(select(AssessmentAttempt.id).where(AssessmentAttempt.assessment_id == assessment.id).limit(1)))
    if payload.action == "publish":
        return publish_assessment(assessment_id, actor, session)
    if payload.action == "unpublish":
        if attempts_exist:
            raise HTTPException(status.HTTP_409_CONFLICT, "An assessment with student attempts cannot be unpublished")
        assessment.status = "draft"
    else:
        assessment.status = "completed"
    db.write_audit(session, actor.id, f"assessment.{payload.action}", "assessment", assessment.id, {"reason": payload.reason})
    session.commit()
    return {"id": assessment.id, "status": assessment.status}


@app.post("/api/academics/assessments/{assessment_id}/questions", status_code=status.HTTP_201_CREATED)
def add_question(assessment_id: int, payload: QuestionCreate, actor: User = Depends(teacher_only), session: Session = Depends(db.get_db)):
    assessment = require_assessment_access(session, actor, assessment_id)
    if assessment.status != "draft":
        raise HTTPException(status.HTTP_409_CONFLICT, "Questions can be changed only while the assessment is a draft")
    if payload.question_type == "mcq" and (not payload.options or len(payload.options) < 2 or payload.correct_answer not in payload.options):
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "MCQ requires at least two options and a matching correct answer")
    if payload.question_type == "true_false" and (payload.correct_answer or "").casefold() not in {"true", "false"}:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "True/False questions require True or False as the correct answer")
    if not (payload.correct_answer or "").strip():
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "Every online question requires a correct answer for automatic evaluation")
    existing_count = session.scalar(select(func.count()).select_from(Question).where(Question.assessment_id == assessment.id)) or 0
    if assessment.question_count and existing_count >= assessment.question_count:
        raise HTTPException(status.HTTP_409_CONFLICT, f"This assessment already contains its required {assessment.question_count} questions")
    marks = float(assessment.common_marks) if assessment.marking_mode == "common" and assessment.common_marks else payload.marks
    row = Question(assessment_id=assessment.id, question_text=payload.question_text, question_type=payload.question_type, options_json=json.dumps(payload.options) if payload.options else None, correct_answer=payload.correct_answer, marks=marks)
    session.add(row)
    session.flush()
    db.write_audit(session, actor.id, "question.created", "question", row.id, {"assessment_id": assessment.id, "marks": marks})
    session.commit()
    return {"id": row.id, **payload.model_dump(exclude={"marks"}), "marks": marks, "questions_added": existing_count + 1, "questions_required": assessment.question_count}


@app.patch("/api/academics/assessments/{assessment_id}/publish")
def publish_assessment(assessment_id: int, actor: User = Depends(teacher_only), session: Session = Depends(db.get_db)):
    assessment = require_assessment_access(session, actor, assessment_id)
    questions = session.scalars(select(Question).where(Question.assessment_id == assessment.id)).all()
    if assessment.mode == "online" and not questions:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "Add at least one question before publishing an online assessment")
    if assessment.mode == "online" and assessment.question_count and len(questions) != assessment.question_count:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, f"Add exactly {assessment.question_count} questions before publishing; currently {len(questions)}")
    if questions:
        assessment.total_marks = sum(question.marks for question in questions)
    assessment.status = "published"
    teaching = session.get(TeacherAssignment, assessment.teacher_assignment_id) if assessment.teacher_assignment_id else None
    notified = notify_section_students(
        session,
        institution_id=actor.institution_id,
        section_id=assessment.section_id,
        term_id=teaching.term_id,
        template="assessment_published",
        message=f"New {assessment.mode} test: {assessment.title}. Opens {assessment.scheduled_at.strftime('%d %b %Y, %H:%M')}.",
        dedup_prefix=f"assessment:{assessment.id}:published",
        deep_link="/portal?view=exams",
    ) if teaching else 0
    db.write_audit(session, actor.id, "assessment.published", "assessment", assessment.id, {"total_marks": assessment.total_marks, "notified": notified})
    session.commit()
    return {"id": assessment.id, "status": assessment.status, "total_marks": assessment.total_marks}


@app.post("/api/academics/assessments/{assessment_id}/offline-results", status_code=status.HTTP_201_CREATED)
def record_offline_result(assessment_id: int, payload: OfflineResultCreate, actor: User = Depends(teacher_only), session: Session = Depends(db.get_db)):
    assessment = require_assessment_access(session, actor, assessment_id)
    if assessment.mode != "offline" or payload.marks > assessment.total_marks:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "Check assessment mode and maximum marks")
    student_row = session.get(Student, payload.student_id)
    teaching_assignment = session.get(TeacherAssignment, assessment.teacher_assignment_id) if assessment.teacher_assignment_id else None
    membership_conditions = [SectionMembership.student_id == payload.student_id, SectionMembership.section_id == assessment.section_id, SectionMembership.active.is_(True)]
    if teaching_assignment:
        membership_conditions.append(SectionMembership.term_id == teaching_assignment.term_id)
    enrolled = session.scalar(select(SectionMembership.id).where(*membership_conditions))
    if not student_row or student_row.institution_id != actor.institution_id:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Student not found for this institution")
    if not enrolled:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "Student is not assigned to this class section")
    result = session.scalar(select(AssessmentResult).where(AssessmentResult.assessment_id == assessment.id, AssessmentResult.student_id == payload.student_id))
    if not result:
        result = AssessmentResult(assessment_id=assessment.id, student_id=payload.student_id, marks=payload.marks, evaluated_by=actor.id)
        session.add(result)
    result.marks = payload.marks
    result.evaluated_by = actor.id
    assessment.status = "completed"
    db.add_notification(
        session,
        institution_id=actor.institution_id,
        user_id=student_row.user_id,
        student_id=student_row.id,
        template="assessment_result",
        message=f"Your result for {assessment.title} is {payload.marks:g}/{assessment.total_marks:g}.",
        dedup_key=f"offline-assessment-result:{assessment.id}:{student_row.id}:{payload.marks:g}",
        deep_link="/portal?view=results",
    )
    db.write_audit(session, actor.id, "offline_result.recorded", "assessment", assessment.id, payload.model_dump())
    session.commit()
    return {"assessment_id": assessment.id, **payload.model_dump()}


@app.get("/api/academics/assessments/{assessment_id}/results")
def assessment_results(assessment_id: int, actor: User = Depends(teacher_only), session: Session = Depends(db.get_db)):
    assessment = require_assessment_access(session, actor, assessment_id)
    rows = session.execute(
        select(AssessmentResult, Student.roll_no, User.name)
        .join(Student, Student.id == AssessmentResult.student_id)
        .join(User, User.id == Student.user_id)
        .where(AssessmentResult.assessment_id == assessment.id, Student.institution_id == actor.institution_id)
        .order_by(User.name)
    ).all()
    return {
        "assessment_id": assessment.id,
        "title": assessment.title,
        "total_marks": assessment.total_marks,
        "results": [{"student_id": result.student_id, "student_name": name, "roll_no": roll_no, "marks": result.marks, "percentage": round(100 * result.marks / assessment.total_marks, 1) if assessment.total_marks else 0, "status": result.status} for result, roll_no, name in rows],
    }


@app.get("/api/academics/assessments/{assessment_id}/results-import/template")
def result_import_template(assessment_id: int, actor: User = Depends(teacher_only), session: Session = Depends(db.get_db)):
    assessment = require_assessment_access(session, actor, assessment_id)
    teaching = session.get(TeacherAssignment, assessment.teacher_assignment_id) if assessment.teacher_assignment_id else None
    query = select(Student.roll_no, User.name).join(User, User.id == Student.user_id).join(SectionMembership, SectionMembership.student_id == Student.id).where(
        Student.institution_id == actor.institution_id,
        SectionMembership.section_id == assessment.section_id,
        SectionMembership.active.is_(True),
    )
    if teaching:
        query = query.where(SectionMembership.term_id == teaching.term_id)
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(["roll_no", "student_name", "marks", "absent"])
    for roll_no, name in session.execute(query.order_by(Student.roll_no)).all():
        writer.writerow([roll_no, name, "", "false"])
    return StreamingResponse(iter([output.getvalue()]), media_type="text/csv", headers={"Content-Disposition": f"attachment; filename=assessment-{assessment.id}-results.csv"})


def validate_result_rows(session: Session, actor: User, assessment: Assessment, rows: list[dict[str, str]]) -> list[dict]:
    teaching = session.get(TeacherAssignment, assessment.teacher_assignment_id) if assessment.teacher_assignment_id else None
    reviewed, seen = [], set()
    for line, raw in enumerate(rows, 2):
        roll = str(raw.get("roll_no", "")).strip().upper()
        errors: list[str] = []
        student = session.scalar(select(Student).where(Student.institution_id == actor.institution_id, Student.roll_no == roll)) if roll else None
        if not student:
            errors.append("roll_no was not found")
        elif not session.scalar(select(SectionMembership.id).where(
            SectionMembership.student_id == student.id,
            SectionMembership.section_id == assessment.section_id,
            SectionMembership.term_id == (teaching.term_id if teaching else SectionMembership.term_id),
            SectionMembership.active.is_(True),
        )):
            errors.append("student is not in this assessment class")
        if roll in seen:
            errors.append("duplicate roll_no in file")
        seen.add(roll)
        absent = str(raw.get("absent", "false")).strip().lower() in {"true", "yes", "1", "absent"}
        marks = 0.0
        if not absent:
            try:
                marks = float(str(raw.get("marks", "")).strip())
                if marks < 0 or marks > assessment.total_marks:
                    errors.append(f"marks must be between 0 and {assessment.total_marks:g}")
            except ValueError:
                errors.append("marks must be a number, or set absent=true")
        reviewed.append({"row": line, "roll_no": roll, "student_id": student.id if student else None, "student_name": student.user.name if student else str(raw.get("student_name", "")), "marks": marks, "absent": absent, "valid": not errors, "errors": errors})
    return reviewed


@app.post("/api/academics/assessments/{assessment_id}/results-import/preview", status_code=status.HTTP_201_CREATED)
async def preview_result_import(assessment_id: int, file: UploadFile = File(...), actor: User = Depends(teacher_only), session: Session = Depends(db.get_db)):
    assessment = require_assessment_access(session, actor, assessment_id)
    if assessment.mode != "offline":
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "Bulk result import is for offline assessments")
    payload = await file.read(settings.max_upload_mb * 1024 * 1024 + 1)
    if len(payload) > settings.max_upload_mb * 1024 * 1024:
        raise HTTPException(status.HTTP_413_REQUEST_ENTITY_TOO_LARGE, "Import file is too large")
    replay = UploadFile(filename=file.filename or "results.csv", file=io.BytesIO(payload))
    rows = validate_result_rows(session, actor, assessment, await import_rows(replay))
    batch = ResultImportBatch(
        institution_id=actor.institution_id,
        assessment_id=assessment.id,
        uploaded_by=actor.id,
        file_name=Path(file.filename or "results.csv").name,
        checksum_sha256=hashlib.sha256(payload).hexdigest(),
        row_count=len(rows),
        valid_count=sum(row["valid"] for row in rows),
        error_count=sum(not row["valid"] for row in rows),
        preview_json=json.dumps(rows),
    )
    session.add(batch)
    session.flush()
    db.write_audit(session, actor.id, "results.import_previewed", "result_import_batch", batch.id, {"assessment_id": assessment.id, "rows": len(rows), "errors": batch.error_count})
    session.commit()
    return {"batch_id": batch.id, "rows": rows, "commit_allowed": bool(rows) and batch.error_count == 0}


@app.post("/api/academics/results-import/{batch_id}/publish")
def publish_result_import(batch_id: int, actor: User = Depends(teacher_only), session: Session = Depends(db.get_db)):
    batch = session.get(ResultImportBatch, batch_id)
    if not batch or batch.institution_id != actor.institution_id:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Result import was not found")
    assessment = require_assessment_access(session, actor, batch.assessment_id)
    if batch.status != "draft" or batch.error_count:
        raise HTTPException(status.HTTP_409_CONFLICT, "Only an error-free draft can be published")
    for row in json.loads(batch.preview_json):
        result = session.scalar(select(AssessmentResult).where(AssessmentResult.assessment_id == assessment.id, AssessmentResult.student_id == row["student_id"]))
        if not result:
            result = AssessmentResult(assessment_id=assessment.id, student_id=row["student_id"], marks=row["marks"], evaluated_by=actor.id)
            session.add(result)
        result.marks = row["marks"]
        result.status = "absent" if row["absent"] else "published"
        result.evaluated_by = actor.id
    batch.status = "published"
    batch.approved_by = actor.id
    batch.approved_at = batch.published_at = datetime.now(timezone.utc)
    assessment.status = "completed"
    db.write_audit(session, actor.id, "results.import_published", "result_import_batch", batch.id, {"assessment_id": assessment.id, "rows": batch.row_count})
    session.commit()
    return {"batch_id": batch.id, "status": batch.status, "published": batch.valid_count}


@app.post("/api/student/assessments/{assessment_id}/start", status_code=status.HTTP_201_CREATED)
def start_assessment(assessment_id: int, user: User = Depends(current_user), session: Session = Depends(db.get_db)):
    student = session.scalar(select(Student).where(Student.user_id == user.id))
    assessment = session.get(Assessment, assessment_id)
    course = session.get(Course, assessment.course_id) if assessment else None
    if not student or not assessment or not course or course.institution_id != user.institution_id or assessment.mode != "online" or assessment.status != "published":
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Available online assessment not found")
    teaching_assignment = session.get(TeacherAssignment, assessment.teacher_assignment_id) if assessment.teacher_assignment_id else None
    membership_conditions = [SectionMembership.student_id == student.id, SectionMembership.section_id == assessment.section_id, SectionMembership.active.is_(True)]
    if teaching_assignment:
        membership_conditions.append(SectionMembership.term_id == teaching_assignment.term_id)
    membership = session.scalar(select(SectionMembership.id).where(*membership_conditions))
    if not membership:
        if teaching_assignment:
            raise HTTPException(status.HTTP_403_FORBIDDEN, "Assessment is not assigned to your class and academic term")
        enrollment = session.scalar(select(Enrollment.id).where(Enrollment.student_id == student.id, Enrollment.course_id == assessment.course_id))
        if not enrollment:
            raise HTTPException(status.HTTP_403_FORBIDDEN, "Assessment is not assigned to your section or enrolled course")
    scheduled_at = assessment.scheduled_at if assessment.scheduled_at.tzinfo else assessment.scheduled_at.replace(tzinfo=timezone.utc)
    now = datetime.now(timezone.utc)
    if now < scheduled_at:
        raise HTTPException(status.HTTP_409_CONFLICT, f"This test opens at {scheduled_at.isoformat()}")
    if assessment.available_until and now > aware_utc(assessment.available_until):
        raise HTTPException(status.HTTP_409_CONFLICT, "This test is closed")
    attempt = session.scalar(select(AssessmentAttempt).where(AssessmentAttempt.assessment_id == assessment.id, AssessmentAttempt.student_id == student.id))
    if not attempt:
        attempt = AssessmentAttempt(assessment_id=assessment.id, student_id=student.id)
        session.add(attempt)
        session.flush()
    elif attempt.status != "in_progress":
        raise HTTPException(status.HTTP_409_CONFLICT, "This assessment has already been submitted. Open its result instead.")
    questions = session.scalars(select(Question).where(Question.assessment_id == assessment.id).order_by(Question.id)).all()
    if not questions:
        raise HTTPException(status.HTTP_409_CONFLICT, "This assessment has no questions. Ask the teacher to correct it.")
    started_at = attempt.started_at if attempt.started_at.tzinfo else attempt.started_at.replace(tzinfo=timezone.utc)
    ends_at = attempt_deadline(attempt, assessment)
    session.commit()
    saved_answers = {row.question_id: row.answer_text or "" for row in session.scalars(select(StudentAnswer).where(StudentAnswer.attempt_id == attempt.id)).all()}
    return {"attempt_id": attempt.id, "ends_at": ends_at.isoformat(), "seconds_remaining": max(0, int((ends_at - datetime.now(timezone.utc)).total_seconds())), "assessment": {"id": assessment.id, "title": assessment.title, "instructions": assessment.instructions, "duration_minutes": assessment.duration_minutes, "total_marks": assessment.total_marks}, "questions": [{"id": row.id, "question_text": row.question_text, "question_type": row.question_type, "options": json.loads(row.options_json) if row.options_json else None, "marks": row.marks} for row in questions], "saved_answers": saved_answers}


@app.patch("/api/student/attempts/{attempt_id}/autosave")
def autosave_assessment(attempt_id: int, payload: AttemptAutosave, user: User = Depends(current_user), session: Session = Depends(db.get_db)):
    student = session.scalar(select(Student).where(Student.user_id == user.id))
    attempt = session.get(AssessmentAttempt, attempt_id)
    if not student or not attempt or attempt.student_id != student.id:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Assessment attempt not found")
    if attempt.status != "in_progress":
        raise HTTPException(status.HTTP_409_CONFLICT, "Assessment was already submitted")
    assessment = session.get(Assessment, attempt.assessment_id)
    if datetime.now(timezone.utc) > attempt_deadline(attempt, assessment):
        saved = {row.question_id: row.answer_text or "" for row in session.scalars(select(StudentAnswer).where(StudentAnswer.attempt_id == attempt.id)).all()}
        result = score_assessment_attempt(session, attempt, assessment, student, saved, user.id, True)
        commit_or_conflict(session, "Assessment result already exists")
        return result
    valid_question_ids = set(session.scalars(select(Question.id).where(Question.assessment_id == assessment.id)).all())
    for submitted in payload.answers:
        if submitted.question_id not in valid_question_ids:
            raise HTTPException(status.HTTP_400_BAD_REQUEST, "An answer does not belong to this assessment")
        row = session.scalar(select(StudentAnswer).where(StudentAnswer.attempt_id == attempt.id, StudentAnswer.question_id == submitted.question_id))
        if not row:
            row = StudentAnswer(attempt_id=attempt.id, question_id=submitted.question_id)
            session.add(row)
        row.answer_text = submitted.answer
        row.awarded_marks = None
    session.commit()
    return {"attempt_id": attempt.id, "saved": len(payload.answers), "saved_at": datetime.now(timezone.utc).isoformat()}


@app.post("/api/student/attempts/{attempt_id}/submit")
def submit_assessment(attempt_id: int, payload: AttemptSubmit, user: User = Depends(current_user), session: Session = Depends(db.get_db)):
    student = session.scalar(select(Student).where(Student.user_id == user.id))
    attempt = session.get(AssessmentAttempt, attempt_id)
    if not student or not attempt or attempt.student_id != student.id:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Assessment attempt not found")
    if attempt.status != "in_progress":
        raise HTTPException(status.HTTP_409_CONFLICT, "Assessment was already submitted")
    assessment = session.get(Assessment, attempt.assessment_id)
    submitted_at = datetime.now(timezone.utc)
    timed_out = submitted_at > attempt_deadline(attempt, assessment)
    saved = {row.question_id: row.answer_text or "" for row in session.scalars(select(StudentAnswer).where(StudentAnswer.attempt_id == attempt.id)).all()}
    if timed_out:
        submitted = saved
    else:
        submitted = {row.question_id: row.answer for row in payload.answers}
        valid_question_ids = set(session.scalars(select(Question.id).where(Question.assessment_id == assessment.id)).all())
        if any(question_id not in valid_question_ids for question_id in submitted):
            raise HTTPException(status.HTTP_400_BAD_REQUEST, "An answer does not belong to this assessment")
    response = score_assessment_attempt(session, attempt, assessment, student, submitted, user.id, timed_out)
    commit_or_conflict(session, "Assessment result already exists")
    return response


@app.get("/api/student/assessments/{assessment_id}/result")
def student_assessment_result(assessment_id: int, user: User = Depends(current_user), session: Session = Depends(db.get_db)):
    student = session.scalar(select(Student).where(Student.user_id == user.id))
    assessment = session.get(Assessment, assessment_id)
    course = session.get(Course, assessment.course_id) if assessment else None
    if not student or not assessment or not course or course.institution_id != user.institution_id:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Assessment result not found")
    teaching = session.get(TeacherAssignment, assessment.teacher_assignment_id) if assessment.teacher_assignment_id else None
    if teaching and not session.scalar(select(SectionMembership.id).where(SectionMembership.student_id == student.id, SectionMembership.section_id == assessment.section_id, SectionMembership.term_id == teaching.term_id, SectionMembership.active.is_(True))):
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Assessment result not found")
    attempt = session.scalar(select(AssessmentAttempt).where(AssessmentAttempt.assessment_id == assessment_id, AssessmentAttempt.student_id == student.id))
    result = session.scalar(select(AssessmentResult).where(AssessmentResult.assessment_id == assessment_id, AssessmentResult.student_id == student.id))
    if not attempt or not result or attempt.status == "in_progress":
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Assessment result is not available yet")
    questions = session.scalars(select(Question).where(Question.assessment_id == assessment_id).order_by(Question.id)).all()
    answers = {row.question_id: row for row in session.scalars(select(StudentAnswer).where(StudentAnswer.attempt_id == attempt.id)).all()}
    review = []
    for question in questions:
        answer = answers.get(question.id)
        selected = answer.answer_text if answer else ""
        is_correct = selected.strip().casefold() == (question.correct_answer or "").strip().casefold()
        review.append({"question_id": question.id, "question_text": question.question_text, "selected_answer": selected, "correct_answer": question.correct_answer if assessment.show_correct_answers else None, "is_correct": is_correct, "awarded_marks": answer.awarded_marks if answer else None, "marks": question.marks})
    return {"assessment_id": assessment.id, "title": assessment.title, "status": attempt.status, "score": result.marks, "total_marks": assessment.total_marks, "answers": review}


STUDENT_IMPORT_COLUMNS = [
    "institution", "student_name", "admission_no", "reg_roll_no", "date_of_birth", "gender",
    "admission_year", "program", "department", "admission_batch", "current_academic_year",
    "year_of_study", "current_semester", "class", "section", "blood_group", "student_mobile_no",
    "email", "father_name", "mother_name", "guardian_name", "parent_guardian_mobile_no",
    "address", "admission_status", "profile_photo",
]
TIMETABLE_IMPORT_COLUMNS = ["action", "department_code", "section_name", "academic_year", "term_name", "course_code", "teacher_email", "weekday", "period_number", "start_time", "end_time", "classroom"]


def _header_key(value: object) -> str:
    return "".join(ch for ch in str(value or "").strip().casefold() if ch.isalnum())


_STUDENT_HEADER_ALIASES = {
    "institution": {"institution", "college", "collegeinstitution", "institutionname", "collegename"},
    "student_name": {"studentname", "name", "nameofstudent", "fullname"},
    "admission_no": {"admissionno", "admissionnumber", "admissionid", "admnno"},
    "reg_roll_no": {"regrollno", "regno", "registrationno", "registrationnumber", "rollno", "rollnumber", "registerno", "registernumber"},
    "date_of_birth": {"dateofbirth", "dob", "birthdate"},
    "gender": {"gender", "sex"},
    "admission_year": {"admissionyear", "yearofadmission"},
    "program": {"program", "programme", "course", "degree"},
    "department": {"department", "departmentname", "departmentcode", "dept", "deptcode", "branch", "branchname"},
    "admission_batch": {"admissionbatch", "batch", "batchyear", "batchname"},
    "current_academic_year": {"currentacademicyear", "academicyear", "academicyearcurrent"},
    "year_of_study": {"yearofstudy", "studyyear", "currentyear", "year"},
    "current_semester": {"currentsemester", "semester", "term", "termname"},
    "class": {"class", "classname", "classgroup", "yearclass"},
    "section": {"section", "sectionname", "division"},
    "blood_group": {"bloodgroup", "bloodtype"},
    "student_mobile_no": {"studentmobileno", "studentmobile", "mobile", "mobileno", "phone", "studentphone"},
    "email": {"email", "studentemail", "emailid", "emailaddress"},
    "personal_email": {"personalemail", "personalemailaddress"},
    "father_name": {"fathername", "fathersname", "parentfathername"},
    "mother_name": {"mothername", "mothersname", "parentmothername"},
    "guardian_name": {"guardianname", "nameofguardian"},
    "parent_guardian_mobile_no": {"parentguardianmobileno", "parentsguardianmobileno", "parentmobile", "parentmobileno", "guardianmobile", "guardianphone", "parentphone", "parentcontact"},
    "address": {"address", "residentialaddress", "studentaddress", "permanentaddress"},
    "admission_status": {"admissionstatus", "status", "studentstatus"},
    "profile_photo": {"profilephoto", "photograph", "photo", "image", "photofilename", "imagefilename"},
    "action": {"action", "importaction"},
    "temporary_password": {"temporarypassword", "initialpassword"},
}


def normalize_student_import_row(raw: dict[str, str]) -> dict[str, str]:
    normalized = {_header_key(key): str(value or "").strip() for key, value in raw.items() if key is not None}
    result = {key: "" for key in [*STUDENT_IMPORT_COLUMNS, "action", "temporary_password", "personal_email"]}
    for target, aliases in _STUDENT_HEADER_ALIASES.items():
        for alias in aliases | {_header_key(target)}:
            if alias in normalized and normalized[alias] != "":
                result[target] = normalized[alias]
                break
    return result


async def import_rows(upload: UploadFile) -> list[dict[str, str]]:
    raw = await upload.read(settings.max_upload_mb * 1024 * 1024 + 1)
    if len(raw) > settings.max_upload_mb * 1024 * 1024:
        raise HTTPException(413, "Import file is too large")
    suffix = Path(upload.filename or "").suffix.lower()
    try:
        if suffix == ".csv":
            reader = csv.DictReader(io.StringIO(raw.decode("utf-8-sig")))
            if not reader.fieldnames:
                raise HTTPException(400, "The CSV file has no header row")
            return [{str(key).strip(): str(value or "").strip() for key, value in row.items() if key is not None} for row in reader if any(str(v or "").strip() for v in row.values())]
        if suffix == ".xlsx":
            from openpyxl import load_workbook
            workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
            sheet = workbook.active
            cells = sheet.iter_rows()
            headers = [str(cell.value or "").strip() for cell in next(cells)]
            rows = []
            for row in cells:
                raw_row = {}
                for index, cell in enumerate(row):
                    if index >= len(headers) or not headers[index]:
                        continue
                    value = cell.value
                    if value is None: text_value = ""
                    elif isinstance(value, (datetime, date)): text_value = value.date().isoformat() if isinstance(value, datetime) else value.isoformat()
                    elif isinstance(value, int) and str(cell.number_format or "").strip() and set(str(cell.number_format).strip()) == {"0"}:
                        text_value = str(value).zfill(len(str(cell.number_format).strip()))
                    else: text_value = str(value).strip()
                    raw_row[headers[index]] = text_value
                if any(raw_row.values()): rows.append(raw_row)
            workbook.close()
            return rows
        if suffix == ".xls":
            import xlrd
            workbook = xlrd.open_workbook(file_contents=raw, on_demand=True)
            sheet = workbook.sheet_by_index(0)
            headers = [str(value or "").strip() for value in sheet.row_values(0)]
            rows = []
            for row_index in range(1, sheet.nrows):
                raw_row = {}
                for col_index, header in enumerate(headers):
                    if not header: continue
                    value = sheet.cell_value(row_index, col_index)
                    if sheet.cell_type(row_index, col_index) == xlrd.XL_CELL_DATE:
                        value = datetime(*xlrd.xldate_as_tuple(value, workbook.datemode)).date().isoformat()
                    elif isinstance(value, float) and value.is_integer(): value = str(int(value))
                    raw_row[header] = str(value or "").strip()
                if any(raw_row.values()): rows.append(raw_row)
            workbook.release_resources()
            return rows
    except HTTPException:
        raise
    except Exception as exc:
        logger.exception("Student import file could not be parsed")
        raise HTTPException(400, "The spreadsheet could not be read. Check the file format and header row.") from exc
    raise HTTPException(400, "Upload a CSV, XLSX, or XLS file")


def parse_import_date(value: str) -> date | None:
    value = (value or "").strip()
    if not value: return None
    for fmt in ("%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y", "%m/%d/%Y"):
        try: return datetime.strptime(value, fmt).date()
        except ValueError: pass
    raise ValueError("Use YYYY-MM-DD or DD-MM-YYYY")


def _year_key(value: str) -> str:
    digits = "".join(ch for ch in (value or "") if ch.isdigit())
    if len(digits) >= 6:
        return digits[:4] + digits[-2:]
    return digits


def _semester_number(value: str) -> int | None:
    digits = "".join(ch for ch in (value or "") if ch.isdigit())
    return int(digits) if digits else None


def apply_photo_filename_inference(rows: list[dict[str, str]], photos: dict[str, tuple[bytes, str]]) -> None:
    extensions = (".jpg", ".jpeg", ".png", ".webp")
    for item in rows:
        if item.get("profile_photo", "").strip():
            continue
        for identifier in (item.get("admission_no", ""), item.get("reg_roll_no", "")):
            identifier = identifier.strip()
            if not identifier:
                continue
            match = next((f"{identifier}{ext}".casefold() for ext in extensions if f"{identifier}{ext}".casefold() in photos), None)
            if match:
                item["profile_photo"] = match
                break


def validate_student_import(session: Session, actor: User, rows: list[dict[str, str]]) -> list[dict]:
    output, seen_email, seen_roll, seen_admission = [], set(), set(), set()
    institution = session.get(Institution, actor.institution_id)
    for number, item in enumerate(rows, 2):
        errors = []
        action = (item.get("action") or "upsert").lower()
        item["action"] = action
        if action not in {"create", "update", "upsert"}: errors.append("action must be create, update, or upsert")
        if item.get("institution") and institution and item["institution"].strip().casefold() not in {institution.name.casefold(), institution.slug.casefold()}:
            errors.append("institution/college does not match your signed-in institution")
        department_value = item.get("department", "").strip()
        department = session.scalar(select(Department).where(Department.institution_id == actor.institution_id, Department.active.is_(True), or_(func.lower(Department.code) == department_value.lower(), func.lower(Department.name) == department_value.lower()))) if department_value else None
        section_value = item.get("section", "").strip()
        section = None
        if section_value:
            candidates = session.scalars(select(ClassSection).where(ClassSection.institution_id == actor.institution_id, ClassSection.active.is_(True))).all()
            matching = [candidate for candidate in candidates if candidate.name.casefold() == section_value.casefold() or candidate.name.casefold().endswith("-" + section_value.casefold()) or candidate.name.casefold().endswith("/" + section_value.casefold())]
            if department:
                matching = [candidate for candidate in matching if candidate.department.casefold() in {department.name.casefold(), department.code.casefold()}]
            if item.get("current_semester"):
                sem_digits = "".join(ch for ch in item["current_semester"] if ch.isdigit())
                if sem_digits:
                    matching = [candidate for candidate in matching if candidate.semester == int(sem_digits)]
            if item.get("admission_batch"):
                batch_digits = "".join(ch for ch in item["admission_batch"] if ch.isdigit())[:4]
                if batch_digits and len(batch_digits) == 4:
                    matching = [candidate for candidate in matching if candidate.batch_year == int(batch_digits)]
            section = matching[0] if len(matching) == 1 else None
            if len(matching) > 1: errors.append("section is ambiguous; use the full existing section name")
        academic_year = item.get("current_academic_year", "").strip()
        term_name = item.get("current_semester", "").strip()
        term = None
        if academic_year and term_name:
            terms = session.scalars(select(AcademicTerm).where(AcademicTerm.institution_id == actor.institution_id, AcademicTerm.active.is_(True))).all()
            requested_semester = _semester_number(term_name)
            term = next((candidate for candidate in terms if _year_key(candidate.academic_year) == _year_key(academic_year) and candidate.name.casefold() == term_name.casefold()), None)
            if term is None and requested_semester is not None:
                term = next((candidate for candidate in terms if _year_key(candidate.academic_year) == _year_key(academic_year) and _semester_number(candidate.name) == requested_semester), None)
        name = item.get("student_name", "").strip()
        roll = item.get("reg_roll_no", "").strip().upper()
        admission_no = item.get("admission_no", "").strip().upper()
        email_raw = item.get("email", "").strip().lower()
        if not name: errors.append("student_name is required")
        if not roll: errors.append("reg_roll_no is required")
        if not department: errors.append("department was not found; use its existing code or name")
        if not section: errors.append("section was not found or does not uniquely match an existing class section")
        if not term: errors.append("current_academic_year/current_semester did not match an existing academic term")
        existing = None
        if admission_no:
            existing = session.scalar(select(Student).join(User, User.id == Student.user_id).where(Student.institution_id == actor.institution_id, func.lower(Student.admission_number) == admission_no.lower()))
        if existing is None and roll:
            existing = session.scalar(select(Student).join(User, User.id == Student.user_id).where(Student.institution_id == actor.institution_id, func.lower(Student.roll_no) == roll.lower()))
        if existing is None and email_raw:
            existing = session.scalar(select(Student).join(User, User.id == Student.user_id).where(Student.institution_id == actor.institution_id, func.lower(User.email) == email_raw))
        if email_raw:
            email = email_raw
        elif existing:
            # Blank email on an update means preserve the existing login email.
            email = existing.user.email.lower()
        elif department and roll and not department.email_domain.endswith(".invalid"):
            email = f"{roll.lower()}@{department.email_domain}"
        else:
            email = ""
        if not email or "@" not in email: errors.append("valid email is required for new students, or configure a department email domain")
        if email in seen_email: errors.append("duplicate email inside file")
        if roll in seen_roll: errors.append("duplicate registration/roll number inside file")
        if admission_no and admission_no in seen_admission: errors.append("duplicate admission number inside file")
        if email: seen_email.add(email)
        if roll: seen_roll.add(roll)
        if admission_no: seen_admission.add(admission_no)
        if existing and email and roll:
            conflict_query = select(Student.id).join(User, User.id == Student.user_id).where(
                Student.institution_id == actor.institution_id,
                or_(func.lower(User.email) == email, func.lower(Student.roll_no) == roll),
                Student.id != existing.id,
            )
            if session.scalar(conflict_query): errors.append("email or registration/roll number belongs to a different student")
        if action == "create" and existing: errors.append("student already exists; use action=update or upsert")
        if action == "update" and not existing: errors.append("student to update was not found")
        if action == "upsert": item["action"] = "update" if existing else "create"
        try: dob = parse_import_date(item.get("date_of_birth", ""))
        except ValueError as exc: dob = None; errors.append(f"date_of_birth: {exc}")
        if item["action"] == "create" and not dob: errors.append("date_of_birth is required for a new account because it is used for the initial temporary password")
        status_value = item.get("admission_status", "").strip().casefold()
        if status_value and status_value not in {"active", "inactive", "graduated", "withdrawn", "cancelled", "deferred", "on hold", "on_hold"}:
            errors.append("admission_status must be Active, Inactive, Graduated, Withdrawn, Cancelled, Deferred, or On Hold")
        photo_name = item.get("profile_photo", "").strip()
        if photo_name and (Path(photo_name).name != photo_name or ".." in photo_name): errors.append("profile_photo must be a filename inside the uploaded photo ZIP")
        resolved = {"department_id": department.id if department else None, "section_id": section.id if section else None, "term_id": term.id if term else None, "existing_student_id": existing.id if existing else None, "date_of_birth": dob.isoformat() if dob else None, "email": email}
        output.append({"row": number, "action": item["action"], "data": item, "valid": not errors, "errors": errors, "resolved": resolved})
    return output


def normalize_student_import_rows(raw_rows: list[dict[str, str]]) -> list[dict[str, str]]:
    return [normalize_student_import_row(row) for row in raw_rows]


def _student_template_content(institution_name: str = "") -> str:
    labels = ["College/Institution", "Student Name", "Admission No", "Reg/Roll No", "Date of Birth", "Gender", "Admission Year", "Program", "Department", "Admission Batch", "Current Academic Year", "Year of Study", "Current Semester", "Class", "Section", "Blood Group", "Student Mobile No", "Email", "Father Name", "Mother Name", "Guardian Name", "Parent's/Guardian Mobile No", "Address", "Admission Status", "Profile Photo"]
    output = io.StringIO()
    csv.writer(output).writerow(labels)
    return output.getvalue()


@app.get("/api/admin/import/students/template")
def student_import_template(actor: User = Depends(admin_only)):
    content = _student_template_content(actor.institution.name if actor.institution else "")
    return StreamingResponse(iter([content]), media_type="text/csv; charset=utf-8", headers={"Content-Disposition": "attachment; filename=ours-student-master-template.csv"})


async def read_photo_bundle(upload: UploadFile | None) -> dict[str, tuple[bytes, str]]:
    if not upload or not upload.filename: return {}
    if Path(upload.filename).suffix.lower() != ".zip": raise HTTPException(400, "Student photographs must be provided as a ZIP file")
    raw = await upload.read(settings.max_upload_mb * 1024 * 1024 + 1)
    if len(raw) > settings.max_upload_mb * 1024 * 1024: raise HTTPException(413, "Photo ZIP is too large")
    photos: dict[str, tuple[bytes, str]] = {}
    try:
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            members = [item for item in archive.infolist() if not item.is_dir()]
            if len(members) > 2000: raise HTTPException(413, "Photo ZIP contains too many files")
            for member in members:
                filename = Path(member.filename).name
                if not filename or filename != member.filename.replace("\\", "/").split("/")[-1] or member.file_size > 8 * 1024 * 1024: continue
                ext = Path(filename).suffix.lower()
                if ext not in {".jpg", ".jpeg", ".png", ".webp"}: continue
                payload = archive.read(member)
                if not payload or len(payload) > 8 * 1024 * 1024: continue
                image_array = np.frombuffer(payload, dtype=np.uint8)
                decoded = cv2.imdecode(image_array, cv2.IMREAD_COLOR)
                if decoded is None or decoded.shape[0] < 32 or decoded.shape[1] < 32: continue
                mime = {".jpg": "image/jpeg", ".jpeg": "image/jpeg", ".png": "image/png", ".webp": "image/webp"}[ext]
                if filename.casefold() in photos:
                    raise HTTPException(400, f"Photo ZIP contains duplicate filename '{filename}'")
                photos[filename.casefold()] = (payload, mime)
    except HTTPException: raise
    except Exception as exc: raise HTTPException(400, "Could not read the photograph ZIP. Recreate it with standard ZIP compression.") from exc
    return photos


@app.post("/api/admin/import/students/preview")
async def preview_student_import(file: UploadFile = File(...), photos_zip: UploadFile | None = File(default=None), actor: User = Depends(admin_only), session: Session = Depends(db.get_db)):
    rows = normalize_student_import_rows(await import_rows(file))
    photos = await read_photo_bundle(photos_zip)
    apply_photo_filename_inference(rows, photos)
    result = validate_student_import(session, actor, rows)
    for row in result:
        photo_name = row["data"].get("profile_photo", "").strip()
        if photo_name and photo_name.casefold() not in photos:
            row["valid"] = False; row["errors"].append(f"profile photo '{photo_name}' was not found or is not a valid JPG/PNG/WEBP in the uploaded ZIP")
    return {"rows": result, "valid": sum(row["valid"] for row in result), "invalid": sum(not row["valid"] for row in result), "commit_allowed": bool(result) and all(row["valid"] for row in result)}


@app.post("/api/admin/import/students/commit")
async def commit_student_import(file: UploadFile = File(...), photos_zip: UploadFile | None = File(default=None), confirm: bool = Form(False), actor: User = Depends(admin_only), session: Session = Depends(db.get_db)):
    if not confirm: raise HTTPException(400, "Preview the file and explicitly confirm the import")
    rows = normalize_student_import_rows(await import_rows(file))
    photos = await read_photo_bundle(photos_zip)
    apply_photo_filename_inference(rows, photos)
    validated = validate_student_import(session, actor, rows)
    for row in validated:
        photo_name = row["data"].get("profile_photo", "").strip()
        if photo_name and photo_name.casefold() not in photos:
            row["valid"] = False; row["errors"].append(f"profile photo '{photo_name}' was not found in the uploaded ZIP")
    invalid = [row for row in validated if not row["valid"]]
    if invalid: raise HTTPException(400, {"message": "Correct all row errors before importing", "rows": invalid})
    created = updated = 0
    stored_keys: list[str] = []
    replaced_photo_keys: list[str] = []
    try:
        for row in validated:
            data, resolved = row["data"], row["resolved"]
            student = session.get(Student, resolved["existing_student_id"]) if resolved["existing_student_id"] else None
            department = session.get(Department, resolved["department_id"])
            section = session.get(ClassSection, resolved["section_id"])
            term = session.get(AcademicTerm, resolved["term_id"])
            roll = data["reg_roll_no"].strip().upper()
            email = normalized_student_email(resolved["email"] or roll, department)
            if student:
                student.user.name, student.user.email, student.roll_no, student.department, student.department_id = data["student_name"], email, roll, department.name, department.id
                if data.get("admission_no"): student.admission_number = data["admission_no"].strip().upper()
                updated += 1
            else:
                dob = date.fromisoformat(resolved["date_of_birth"])
                initial_password = data.get("temporary_password") or dob.strftime("%d%m%Y")
                account = User(institution_id=actor.institution_id, name=data["student_name"], email=email, password_hash=hash_password(initial_password), role="student", must_change_password=True)
                session.add(account); session.flush()
                student = Student(institution_id=actor.institution_id, user=account, roll_no=roll, admission_number=data.get("admission_no", "").strip().upper() or None, department=department.name, department_id=department.id)
                session.add(student); session.flush(); created += 1
            profile = session.scalar(select(StudentProfile).where(StudentProfile.student_id == student.id))
            if not profile:
                profile = StudentProfile(student_id=student.id); session.add(profile)
            dob = parse_import_date(data.get("date_of_birth", ""))
            if dob: profile.date_of_birth = dob
            personal_email = data.get("personal_email") or data.get("email")
            if personal_email: profile.personal_email = personal_email.strip().lower()
            if data.get("student_mobile_no"): profile.phone = data["student_mobile_no"]
            if data.get("parent_guardian_mobile_no"): profile.guardian_phone = data["parent_guardian_mobile_no"]
            if data.get("guardian_name"): profile.guardian_name = data["guardian_name"]
            if data.get("address"): profile.address = data["address"]
            if data.get("blood_group"): profile.blood_group = data["blood_group"]
            for field, key in (("gender", "gender"), ("admission_year", "admission_year"), ("program", "program"), ("admission_batch", "admission_batch"), ("current_academic_year", "current_academic_year"), ("year_of_study", "year_of_study"), ("current_semester", "current_semester"), ("class_name", "class"), ("section_name", "section"), ("father_name", "father_name"), ("mother_name", "mother_name"), ("admission_status", "admission_status")):
                value = data.get(key, "").strip()
                if value: setattr(profile, field, value)
            photo_name = data.get("profile_photo", "").strip()
            if photo_name:
                payload, mime = photos[photo_name.casefold()]
                key = f"institutions/{actor.institution_id}/students/{student.id}/profile/{uuid4().hex}{Path(photo_name).suffix.lower()}"
                try:
                    stored = storage_service.put_private(key, payload, mime)
                except ValueError as exc:
                    raise HTTPException(status.HTTP_422_UNPROCESSABLE_ENTITY, str(exc)) from exc
                stored_keys.append(stored.key)
                if profile.profile_photo_storage_key and profile.profile_photo_storage_key != stored.key:
                    replaced_photo_keys.append(profile.profile_photo_storage_key)
                profile.profile_photo_storage_key = stored.key
            if data.get("admission_status", "").strip().casefold() in {"inactive", "withdrawn", "cancelled", "graduated"}:
                student.user.active = False
            elif data.get("admission_status", "").strip():
                student.user.active = True
            session.query(SectionMembership).filter(SectionMembership.student_id == student.id, SectionMembership.active.is_(True)).update({SectionMembership.active: False})
            membership = session.scalar(select(SectionMembership).where(SectionMembership.student_id == student.id, SectionMembership.section_id == section.id, SectionMembership.term_id == term.id))
            if membership: membership.active = True
            else: session.add(SectionMembership(student_id=student.id, section_id=section.id, term_id=term.id, active=True))
            ensure_section_enrollments(session, student.id, section.id, term.id)
        db.write_audit(session, actor.id, "students.bulk_imported", "student", None, {"created": created, "updated": updated, "profile_photos": len(stored_keys)})
        commit_or_conflict(session, "Import conflicts with an existing email, roll number, admission number, or membership")
        for old_key in replaced_photo_keys:
            try: storage_service.delete_private(old_key)
            except Exception: logger.exception("Failed to remove replaced student profile photo after successful import")
    except Exception:
        session.rollback()
        for key in stored_keys:
            try: storage_service.delete_private(key)
            except Exception: logger.exception("Failed to remove orphaned student profile photo after import rollback")
        raise
    return {"created": created, "updated": updated, "profile_photos": len(stored_keys), "message": "Student import completed"}


@app.get("/api/students/{student_id}/profile-photo")
def get_student_profile_photo(student_id: int, actor: User = Depends(current_user), session: Session = Depends(db.get_db)):
    student = session.get(Student, student_id)
    if not student or student.institution_id != actor.institution_id:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Student not found")
    allowed = actor.role == "admin"
    if actor.role == "student": allowed = student.user_id == actor.id
    elif actor.role == "teacher":
        profile = session.scalar(select(StaffProfile).where(StaffProfile.user_id == actor.id, StaffProfile.active.is_(True)))
        membership_ids = select(SectionMembership.section_id).where(SectionMembership.student_id == student.id, SectionMembership.active.is_(True))
        allowed = bool(
            profile
            and (
                (profile.designation in {"hod", "dean"} and (not profile.department_id or profile.department_id == student.department_id))
                or session.scalar(select(TeacherAssignment.id).where(
                    TeacherAssignment.teacher_id == actor.id,
                    TeacherAssignment.section_id.in_(membership_ids),
                    TeacherAssignment.active.is_(True),
                ))
            )
        )
    if not allowed:
        raise HTTPException(status.HTTP_403_FORBIDDEN, "You cannot access this student photograph")
    profile = session.scalar(select(StudentProfile).where(StudentProfile.student_id == student.id))
    if not profile or not profile.profile_photo_storage_key: raise HTTPException(status.HTTP_404_NOT_FOUND, "Student photograph not found")
    expected_prefix = f"institutions/{actor.institution_id}/students/{student.id}/profile/"
    if not profile.profile_photo_storage_key.startswith(expected_prefix):
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Student photograph is unavailable")
    local = storage_service.local_path(profile.profile_photo_storage_key)
    if local: return FileResponse(local, media_type={".jpg":"image/jpeg", ".jpeg":"image/jpeg", ".png":"image/png", ".webp":"image/webp"}.get(local.suffix.lower(), "application/octet-stream"), headers={"Cache-Control":"private, no-store"})
    return RedirectResponse(storage_service.download_url(profile.profile_photo_storage_key, 120), status_code=307)

def validate_timetable_import(session: Session, actor: User, rows: list[dict[str, str]]) -> list[dict]:
    weekdays = {"monday": 0, "tuesday": 1, "wednesday": 2, "thursday": 3, "friday": 4, "saturday": 5, "sunday": 6}
    output, seen = [], set()
    for number, raw in enumerate(rows, 2):
        item = {key: raw.get(key, "").strip() for key in TIMETABLE_IMPORT_COLUMNS}; errors = []
        action = (item["action"] or "create").lower()
        if action not in {"create", "update"}: errors.append("action must be create or update")
        department = session.scalar(select(Department).where(Department.institution_id == actor.institution_id, func.lower(Department.code) == item["department_code"].lower()))
        section = session.scalar(select(ClassSection).where(ClassSection.institution_id == actor.institution_id, func.lower(ClassSection.name) == item["section_name"].lower(), ClassSection.active.is_(True)))
        term = session.scalar(select(AcademicTerm).where(AcademicTerm.institution_id == actor.institution_id, AcademicTerm.academic_year == item["academic_year"], func.lower(AcademicTerm.name) == item["term_name"].lower()))
        course = session.scalar(select(Course).where(Course.institution_id == actor.institution_id, func.lower(Course.code) == item["course_code"].lower(), Course.active.is_(True)))
        teacher = session.scalar(select(User).where(User.institution_id == actor.institution_id, func.lower(User.email) == item["teacher_email"].lower(), User.active.is_(True)))
        if not department: errors.append("department_code was not found")
        if not section: errors.append("section_name was not found")
        elif department and section.department.casefold() not in {department.code.casefold(), department.name.casefold()}: errors.append("section belongs to another department")
        if not term: errors.append("academic year/term was not found")
        if not course: errors.append("course_code was not found")
        if not teacher or teacher.role != "teacher": errors.append("teacher_email was not found")
        assignment = session.scalar(select(TeacherAssignment).where(TeacherAssignment.teacher_id == (teacher.id if teacher else -1), TeacherAssignment.course_id == (course.id if course else -1), TeacherAssignment.section_id == (section.id if section else -1), TeacherAssignment.term_id == (term.id if term else -1), TeacherAssignment.active.is_(True)))
        if not assignment: errors.append("matching teacher allocation was not found")
        weekday_text = item["weekday"].lower()
        try: weekday = weekdays.get(weekday_text, int(weekday_text) if weekday_text.isdigit() else -1); period = int(item["period_number"]); start = datetime.strptime(item["start_time"], "%H:%M").time(); end = datetime.strptime(item["end_time"], "%H:%M").time()
        except ValueError: weekday, period, start, end = -1, -1, None, None; errors.append("weekday, period or HH:MM time is invalid")
        if weekday not in range(7): errors.append("weekday must be Monday-Sunday or 0-6")
        if period < 1 or period > 30: errors.append("period_number must be 1-30")
        if start and end and end <= start: errors.append("end_time must be after start_time")
        if not item["classroom"]: errors.append("classroom is required")
        key = (section.id if section else None, weekday, period, term.id if term else None)
        if key in seen: errors.append("duplicate class/day/period inside file")
        seen.add(key)
        existing = session.scalar(select(TimetableSlot).join(TeacherAssignment, TeacherAssignment.id == TimetableSlot.assignment_id).where(TimetableSlot.section_id == (section.id if section else -1), TeacherAssignment.term_id == (term.id if term else -1), TimetableSlot.weekday == weekday, TimetableSlot.period_number == period, TimetableSlot.active.is_(True)))
        if action == "create" and existing: errors.append("published period already exists; use update")
        if action == "update" and not existing: errors.append("period to update was not found")
        output.append({"row": number, "action": action, "data": item, "valid": not errors, "errors": errors, "resolved": {"assignment_id": assignment.id if assignment else None, "section_id": section.id if section else None, "term_id": term.id if term else None, "weekday": weekday, "period_number": period, "start_time": start.strftime("%H:%M") if start else None, "end_time": end.strftime("%H:%M") if end else None, "slot_id": existing.id if existing else None}})
    return output


@app.get("/api/admin/import/timetable/template")
def timetable_import_template(_: User = Depends(admin_only)):
    content = ",".join(TIMETABLE_IMPORT_COLUMNS) + "\ncreate,CSE,CSE-A,2026-27,Semester 1,CS101,teacher@cse.example.edu,Monday,1,09:00,10:00,Room 101\n"
    return StreamingResponse(iter([content]), media_type="text/csv", headers={"Content-Disposition": "attachment; filename=ours-timetable-import.csv"})


@app.post("/api/admin/import/timetable/preview")
async def preview_timetable_import(file: UploadFile = File(...), actor: User = Depends(admin_only), session: Session = Depends(db.get_db)):
    rows = validate_timetable_import(session, actor, await import_rows(file))
    return {"rows": rows, "valid": sum(row["valid"] for row in rows), "invalid": sum(not row["valid"] for row in rows), "commit_allowed": bool(rows) and all(row["valid"] for row in rows)}


@app.post("/api/admin/import/timetable/commit")
async def commit_timetable_import(file: UploadFile = File(...), confirm: bool = Form(False), actor: User = Depends(admin_only), session: Session = Depends(db.get_db)):
    if not confirm: raise HTTPException(400, "Preview the file and explicitly confirm the import")
    rows = validate_timetable_import(session, actor, await import_rows(file)); invalid = [row for row in rows if not row["valid"]]
    if invalid: raise HTTPException(400, {"message": "Correct all row errors before importing", "rows": invalid})
    created = updated = generated = 0
    for item in rows:
        data, resolved = item["data"], item["resolved"]
        assignment = session.get(TeacherAssignment, resolved["assignment_id"]); term = session.get(AcademicTerm, resolved["term_id"])
        payload = TimetableSlotCreate(assignment_id=assignment.id, weekday=resolved["weekday"], period_number=resolved["period_number"], start_time=resolved["start_time"], end_time=resolved["end_time"], classroom=data["classroom"], effective_from=term.start_date, effective_to=term.end_date)
        payload = semester_dates(session, assignment, payload)
        slot = session.get(TimetableSlot, resolved["slot_id"]) if resolved["slot_id"] else None
        ensure_timetable_has_no_conflict(session, assignment, payload, exclude_slot_id=slot.id if slot else None)
        if slot:
            for key, value in payload.model_dump().items(): setattr(slot, key, value)
            session.query(ClassSession).filter(ClassSession.timetable_slot_id == slot.id, ClassSession.session_date >= campus_today(), ClassSession.attendance_state != "published", ClassSession.actual_start.is_(None)).update({ClassSession.assignment_id: assignment.id, ClassSession.classroom: payload.classroom}, synchronize_session=False)
            updated += 1
        else:
            slot = TimetableSlot(section_id=assignment.section_id, publication_status="published", proposed_by=actor.id, approved_by=actor.id, approved_at=datetime.now(timezone.utc), **payload.model_dump()); session.add(slot); session.flush(); created += 1
        generated += generate_timetable_sessions(session, slot, actor.id)
    db.write_audit(session, actor.id, "timetable.bulk_imported", "timetable_slot", None, {"created": created, "updated": updated, "sessions_generated": generated})
    session.commit()
    return {"created": created, "updated": updated, "sessions_generated": generated, "message": "Timetable import completed"}


@app.post("/api/admin/students", status_code=status.HTTP_201_CREATED)
def create_student(payload: StudentCreate, request: Request, actor: User = Depends(admin_only), session: Session = Depends(db.get_db)):
    department = department_for_student(session, actor.institution_id, payload.department_id)
    email = normalized_student_email(payload.email or payload.roll_no, department)
    name = payload.name.strip()
    roll_no = payload.roll_no.upper().strip()
    ensure_unique_student_identity(session, actor.institution_id, name, email, roll_no)
    if not payload.password and not payload.date_of_birth:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "Date of birth is required when no temporary password is provided")
    initial_password = payload.password or payload.date_of_birth.strftime("%d%m%Y")
    user = User(institution_id=actor.institution_id, name=name, email=email, password_hash=hash_password(initial_password), role="student", must_change_password=True)
    session.add(user)
    flush_or_conflict(session, "This email address already exists in the institution")
    student = Student(institution_id=actor.institution_id, user=user, roll_no=roll_no, department=department.name if department else payload.department.strip(), department_id=department.id if department else None)
    session.add(student)
    flush_or_conflict(session, "This register/roll number already exists")
    session.add(StudentProfile(student_id=student.id, date_of_birth=payload.date_of_birth))
    section = session.get(ClassSection, payload.section_id)
    term = session.get(AcademicTerm, payload.term_id)
    if not section or section.institution_id != actor.institution_id or not term or term.institution_id != actor.institution_id:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Class section or academic term not found")
    if department and section.department.casefold() not in {department.name.casefold(), department.code.casefold()}:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "The selected class section belongs to a different department")
    session.add(SectionMembership(student_id=student.id, section_id=section.id, term_id=term.id))
    ensure_section_enrollments(session, student.id, section.id, term.id)
    db.write_audit(session, actor.id, "student.created", "student", student.id, {"roll_no": student.roll_no}, request.client.host if request.client else None)
    commit_or_conflict(session, "Student email, register number or section membership already exists")
    return db.student_to_dict(student)


@app.patch("/api/admin/students/{student_id}")
def update_student(student_id: int, payload: StudentUpdate, actor: User = Depends(admin_only), session: Session = Depends(db.get_db)):
    student = session.get(Student, student_id)
    if not student or student.institution_id != actor.institution_id:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Student not found")
    department = department_for_student(session, actor.institution_id, payload.department_id)
    email = normalized_student_email(payload.email, department)
    name = payload.name.strip()
    roll_no = payload.roll_no.upper().strip()
    ensure_unique_student_identity(session, actor.institution_id, name, email, roll_no, student.id)
    student.user.name = name
    student.user.email = email
    student.roll_no = roll_no
    student.department = department.name if department else payload.department.strip()
    student.department_id = department.id if department else None
    profile = session.scalar(select(StudentProfile).where(StudentProfile.student_id == student.id))
    if not profile:
        profile = StudentProfile(student_id=student.id)
        session.add(profile)
    profile.date_of_birth = payload.date_of_birth
    if bool(payload.section_id) != bool(payload.term_id):
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "Choose both the class section and academic term")
    if payload.section_id and payload.term_id:
        section = session.get(ClassSection, payload.section_id)
        term = session.get(AcademicTerm, payload.term_id)
        if not section or section.institution_id != actor.institution_id or not term or term.institution_id != actor.institution_id:
            raise HTTPException(status.HTTP_404_NOT_FOUND, "Class section or academic term not found")
        if department and section.department.casefold() not in {department.name.casefold(), department.code.casefold()}:
            raise HTTPException(status.HTTP_400_BAD_REQUEST, "The selected class section belongs to a different department")
        session.query(SectionMembership).filter(SectionMembership.student_id == student.id, SectionMembership.term_id == term.id, SectionMembership.section_id != section.id, SectionMembership.active.is_(True)).update({SectionMembership.active: False}, synchronize_session=False)
        membership = session.scalar(select(SectionMembership).where(SectionMembership.student_id == student.id, SectionMembership.section_id == section.id, SectionMembership.term_id == term.id))
        if membership:
            membership.active = True
        else:
            session.add(SectionMembership(student_id=student.id, section_id=section.id, term_id=term.id, active=True))
        ensure_section_enrollments(session, student.id, section.id, term.id)
    flush_or_conflict(session, "Student email or register number already exists")
    db.write_audit(session, actor.id, "student.updated", "student", student.id, {"name": name, "email": email, "roll_no": roll_no, "department": student.department, "section_id": payload.section_id, "term_id": payload.term_id})
    session.commit()
    return db.student_to_dict(student)


@app.patch("/api/admin/students/{student_id}/profile-edit-permission")
def set_student_profile_permission(student_id: int, payload: StudentProfilePermissionUpdate, actor: User = Depends(admin_only), session: Session = Depends(db.get_db)):
    student = session.get(Student, student_id)
    if not student or student.institution_id != actor.institution_id:
        raise HTTPException(404, "Student not found")
    profile = session.scalar(select(StudentProfile).where(StudentProfile.student_id == student.id))
    if not profile:
        profile = StudentProfile(student_id=student.id)
        session.add(profile)
    profile.self_edit_enabled = payload.enabled
    db.write_audit(session, actor.id, "student_profile.permission_changed", "student", student.id, {"enabled": payload.enabled})
    session.commit()
    return {"student_id": student.id, "self_edit_enabled": profile.self_edit_enabled}


@app.post("/api/admin/teachers", status_code=status.HTTP_201_CREATED)
def create_teacher(payload: TeacherCreate, request: Request, actor: User = Depends(admin_only), session: Session = Depends(db.get_db)):
    teacher = User(institution_id=actor.institution_id, name=payload.name.strip(), email=str(payload.email).lower(), password_hash=hash_password(payload.password), role="teacher", must_change_password=True)
    session.add(teacher)
    flush_or_conflict(session, "Email already exists")
    session.add(StaffProfile(institution_id=actor.institution_id, user_id=teacher.id, designation="teacher"))
    db.write_audit(session, actor.id, "teacher.created", "user", teacher.id, ip_address=request.client.host if request.client else None)
    commit_or_conflict(session, "Email already exists")
    return db.public_user(teacher)


@app.post("/api/admin/courses", status_code=status.HTTP_201_CREATED)
def create_course(payload: CourseCreate, actor: User = Depends(admin_only), session: Session = Depends(db.get_db)):
    course = Course(institution_id=actor.institution_id, title=payload.title.strip(), code=payload.code.upper().strip(), total_lessons=payload.total_lessons)
    session.add(course)
    flush_or_conflict(session, "Course code already exists")
    db.write_audit(session, actor.id, "course.created", "course", course.id, {"code": course.code})
    commit_or_conflict(session, "Course code already exists")
    return {"id": course.id, "title": course.title, "code": course.code, "total_lessons": course.total_lessons}


@app.patch("/api/admin/courses/{course_id}")
def update_course(course_id: int, payload: CourseUpdate, actor: User = Depends(admin_only), session: Session = Depends(db.get_db)):
    course = session.get(Course, course_id)
    if not course or course.institution_id != actor.institution_id:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Course not found")
    course.title = payload.title.strip()
    course.code = payload.code.upper().strip()
    course.total_lessons = payload.total_lessons
    flush_or_conflict(session, "Another course already uses this code")
    db.write_audit(session, actor.id, "course.updated", "course", course.id, payload.model_dump())
    session.commit()
    return {"id": course.id, "title": course.title, "code": course.code, "total_lessons": course.total_lessons}


@app.post("/api/admin/enrollments", status_code=status.HTTP_201_CREATED)
def create_enrollment(payload: EnrollmentCreate, actor: User = Depends(admin_only), session: Session = Depends(db.get_db)):
    student_row = session.get(Student, payload.student_id)
    course_row = session.get(Course, payload.course_id)
    if not student_row or student_row.institution_id != actor.institution_id or not course_row or course_row.institution_id != actor.institution_id:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Student or course not found")
    enrollment = Enrollment(student_id=payload.student_id, course_id=payload.course_id)
    session.add(enrollment)
    flush_or_conflict(session, "Student is already enrolled in this course")
    db.write_audit(session, actor.id, "enrollment.created", "enrollment", enrollment.id, payload.model_dump())
    commit_or_conflict(session, "Student is already enrolled in this course")
    return {"id": enrollment.id, **payload.model_dump()}


@app.patch("/api/admin/enrollments/{enrollment_id}")
def update_enrollment(enrollment_id: int, payload: EnrollmentUpdate, actor: User = Depends(admin_only), session: Session = Depends(db.get_db)):
    enrollment = session.get(Enrollment, enrollment_id)
    student_row = session.get(Student, enrollment.student_id) if enrollment else None
    course = session.get(Course, enrollment.course_id) if enrollment else None
    if not enrollment or not student_row or student_row.institution_id != actor.institution_id or not course or course.institution_id != actor.institution_id:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Enrollment not found")
    if payload.completed_lessons > course.total_lessons:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "Completed lessons cannot exceed the course total")
    enrollment.completed_lessons = payload.completed_lessons
    enrollment.score = payload.score
    db.write_audit(session, actor.id, "enrollment.updated", "enrollment", enrollment.id, payload.model_dump())
    session.commit()
    return {"id": enrollment.id, **payload.model_dump()}


@app.patch("/api/admin/students/{student_id}/biometric-consent")
def update_consent(student_id: int, payload: StudentBiometricConsentUpdate, actor: User = Depends(admin_only), session: Session = Depends(db.get_db)):
    student = session.get(Student, student_id)
    if not student or student.institution_id != actor.institution_id:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Student not found")
    profile = session.scalar(select(StudentProfile).where(StudentProfile.student_id == student.id))
    if not profile:
        profile = StudentProfile(student_id=student.id)
        session.add(profile)
    if payload.guardian_confirmed:
        profile.guardian_biometric_consent = True
        profile.guardian_consent_recorded_at = datetime.now(timezone.utc)
    institution = session.get(Institution, actor.institution_id)
    minor_status = student_is_minor(profile)
    if payload.granted and institution and institution.minor_guardian_consent_required:
        if minor_status is None:
            raise HTTPException(status.HTTP_409_CONFLICT, "Record the student's date of birth before enabling biometrics")
        if minor_status and not profile.guardian_biometric_consent:
            raise HTTPException(status.HTTP_409_CONFLICT, "Confirm verified guardian consent before enabling biometrics for a minor")
    student.biometric_consent = payload.granted
    student.consent_recorded_at = datetime.now(timezone.utc) if payload.granted else None
    student.consent_recorded_by = actor.id
    if not payload.granted:
        embedding = session.scalar(select(FaceEmbedding).where(FaceEmbedding.student_id == student.id))
        if embedding:
            session.delete(embedding)
    db.write_audit(session, actor.id, "biometric.consent_updated", "student", student.id, {"granted": payload.granted, "guardian_confirmed": payload.guardian_confirmed})
    session.commit()
    return {"student_id": student.id, "biometric_consent": student.biometric_consent, "guardian_biometric_consent": profile.guardian_biometric_consent, "face_data_deleted": not payload.granted}


@app.delete("/api/admin/students/{student_id}/face")
def delete_face(student_id: int, actor: User = Depends(admin_only), session: Session = Depends(db.get_db)):
    student = session.get(Student, student_id)
    if not student or student.institution_id != actor.institution_id:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Student not found")
    embedding = session.scalar(select(FaceEmbedding).where(FaceEmbedding.student_id == student_id))
    if embedding:
        session.delete(embedding)
    db.write_audit(session, actor.id, "biometric.deleted", "student", student_id)
    session.commit()
    return {"message": "Biometric template deleted"}


@app.post("/api/me/face")
async def register_my_face(image: UploadFile = File(...), user: User = Depends(current_user), session: Session = Depends(db.get_db)):
    if user.role != "student":
        raise HTTPException(status.HTTP_403_FORBIDDEN, "Student access is required")
    student = session.scalar(select(Student).where(Student.user_id == user.id))
    if not student or not student.biometric_consent:
        raise HTTPException(status.HTTP_409_CONFLICT, "Record biometric consent before registering your face")
    try:
        result = await run_in_threadpool(vision.register_face, await decode_image(image), student.id, user.institution_id)
        result["name"] = user.name
        db.write_audit(session, user.id, "biometric.self_enrolled", "student", student.id, {"samples": result["samples"]})
        session.commit()
        return result
    except ValueError as exc:
        code = status.HTTP_409_CONFLICT if "already registered" in str(exc).lower() else status.HTTP_400_BAD_REQUEST
        raise HTTPException(code, str(exc)) from exc
    except Exception as exc:
        raise HTTPException(status.HTTP_503_SERVICE_UNAVAILABLE, f"Face model could not start: {exc}") from exc


@app.post("/api/students/{student_id}/face")
async def register_face(student_id: int, image: UploadFile = File(...), actor: User = Depends(teacher_only), session: Session = Depends(db.get_db)):
    student = session.get(Student, student_id)
    if not student or student.institution_id != actor.institution_id:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Student not found")
    if actor.role == "teacher":
        assigned = session.scalar(
            select(SectionMembership.id)
            .join(TeacherAssignment, (TeacherAssignment.section_id == SectionMembership.section_id) & (TeacherAssignment.term_id == SectionMembership.term_id))
            .where(SectionMembership.student_id == student.id, SectionMembership.active.is_(True), TeacherAssignment.teacher_id == actor.id, TeacherAssignment.active.is_(True))
            .limit(1)
        )
        if not assigned:
            raise HTTPException(status.HTTP_403_FORBIDDEN, "You may enrol faces only for students in your active classes")
    if not student.biometric_consent:
        raise HTTPException(status.HTTP_409_CONFLICT, "Biometric consent must be recorded by an administrator first")
    try:
        result = await run_in_threadpool(vision.register_face, await decode_image(image), student_id, actor.institution_id)
        result["name"] = student.user.name
        result["message"] = f"{student.user.name}: {result['message']}"
        db.write_audit(session, actor.id, "biometric.enrolled", "student", student_id)
        session.commit()
        return result
    except ValueError as exc:
        code = status.HTTP_409_CONFLICT if "already registered" in str(exc).lower() else status.HTTP_400_BAD_REQUEST
        raise HTTPException(code, str(exc)) from exc
    except Exception as exc:
        raise HTTPException(status.HTTP_503_SERVICE_UNAVAILABLE, f"Face model could not start: {exc}") from exc


@app.post("/api/staff/{user_id}/face")
async def register_staff_face(user_id: int, image: UploadFile = File(...), actor: User = Depends(teacher_only), session: Session = Depends(db.get_db)):
    account = session.get(User, user_id)
    if not account or account.institution_id != actor.institution_id or account.role not in ("teacher", "admin", "staff"):
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Staff account not found")
    if actor.role != "admin" and actor.id != account.id:
        raise HTTPException(status.HTTP_403_FORBIDDEN, "Staff may register only their own face; an administrator can assist other staff")
    profile = staff_profile_for(session, account)
    if not profile.biometric_consent:
        raise HTTPException(status.HTTP_409_CONFLICT, "Staff biometric consent must be recorded by an administrator first")
    try:
        result = await run_in_threadpool(vision.register_staff_face, await decode_image(image), account.id, actor.institution_id)
        result["name"] = account.name
        result["message"] = f"{account.name}: {result['message']}"
        db.write_audit(session, actor.id, "staff.biometric_enrolled", "user", account.id, {"designation": profile.designation})
        session.commit()
        return result
    except ValueError as exc:
        code = status.HTTP_409_CONFLICT if "already registered" in str(exc).lower() else status.HTTP_400_BAD_REQUEST
        raise HTTPException(code, str(exc)) from exc
    except Exception as exc:
        raise HTTPException(status.HTTP_503_SERVICE_UNAVAILABLE, f"Face model could not start: {exc}") from exc


@app.patch("/api/me/staff-biometric-consent")
def update_my_staff_consent(payload: ConsentUpdate, user: User = Depends(current_user), session: Session = Depends(db.get_db)):
    if user.role not in {"teacher", "admin", "staff"}:
        raise HTTPException(status.HTTP_403_FORBIDDEN, "Staff access is required")
    profile = staff_profile_for(session, user)
    profile.biometric_consent = payload.granted
    profile.consent_recorded_at = datetime.now(timezone.utc) if payload.granted else None
    profile.consent_recorded_by = user.id
    if not payload.granted:
        embedding = session.scalar(select(StaffFaceEmbedding).where(StaffFaceEmbedding.user_id == user.id))
        if embedding:
            session.delete(embedding)
    db.write_audit(session, user.id, "staff.biometric_self_consent_updated", "user", user.id, {"granted": payload.granted})
    session.commit()
    return {"biometric_consent": payload.granted, "face_data_deleted": not payload.granted}


@app.post("/api/me/staff-face")
async def register_my_staff_face(image: UploadFile = File(...), user: User = Depends(current_user), session: Session = Depends(db.get_db)):
    if user.role not in {"teacher", "admin", "staff"}:
        raise HTTPException(status.HTTP_403_FORBIDDEN, "Staff access is required")
    return await register_staff_face(user.id, image, user, session)


@app.post("/api/admin/biometrics/apply-retention")
def apply_biometric_retention(actor: User = Depends(admin_only), session: Session = Depends(db.get_db)):
    result = db.apply_biometric_retention(session, actor.institution_id)
    db.write_audit(session, actor.id, "biometric.retention_applied", "institution", actor.institution_id, result)
    session.commit()
    return result


@app.post("/api/monitor/frame")
async def monitor_frame(image: UploadFile = File(...), classroom_id: str = Form("CSE-A"), sample_label: str = Form("none"), sample_track_id: str | None = Form(None), class_session_id: int | None = Form(None), actor: User = Depends(teacher_only), session: Session = Depends(db.get_db)):
    allowed = {"none", "attentive", "sleeping", "talking", "fighting"}
    if sample_label not in allowed:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "Unsupported behaviour label")
    if not classroom_id.strip() or len(classroom_id) > 80:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "Invalid classroom identifier")
    resolved_from_room = False
    try:
        if class_session_id:
            class_session = session.get(ClassSession, class_session_id)
            if not class_session:
                raise HTTPException(status.HTTP_404_NOT_FOUND, "Class session not found")
            require_assignment_access(session, actor, class_session.assignment_id)
        else:
            # The stream is linked to a physical room. Resolve the active period
            # on the server using the configured campus timezone.
            class_session = db.active_class_session_for_room(
                session,
                actor.institution_id,
                classroom_id.strip(),
                db.campus_now(),
                actor.id,
            )
            if class_session:
                require_assignment_access(session, actor, class_session.assignment_id)
                class_session_id = class_session.id
                resolved_from_room = True
        # Record browser/mobile health even when the room currently has no class.
        # This lets operators distinguish an online camera from an unscheduled
        # room, while the server still links attendance only during a timetable
        # window.
        session.commit()
        db.update_camera_status(
            actor.institution_id,
            f"browser:{actor.id}:{classroom_id.strip()}",
            classroom_id.strip(),
            "browser",
            "online",
            class_session_id=class_session_id,
            processing_fps=0,
        )
    except HTTPException:
        raise
    except Exception as exc:
        session.rollback()
        logger.exception("Camera room resolution failed for %s", classroom_id)
        raise HTTPException(status.HTTP_503_SERVICE_UNAVAILABLE, f"Room/timetable resolution failed: {exc}") from exc
    try:
        teaching = session.get(TeacherAssignment, class_session.assignment_id) if class_session_id and class_session else None
        gesture_teacher_id = teaching.teacher_id if teaching else (actor.id if actor.role == "teacher" else None)
        result = await run_in_threadpool(vision.analyse, await decode_image(image), classroom_id.strip(), sample_label, sample_track_id, actor.institution_id, actor.id, class_session_id, gesture_teacher_id)
        result["resolved_class_session_id"] = class_session_id
        result["room_linked"] = True
        result["resolved_from_room"] = resolved_from_room
        return result
    except ValueError as exc:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, str(exc)) from exc
    except Exception as exc:
        try:
            db.update_camera_status(
                actor.institution_id,
                f"browser:{actor.id}:{classroom_id.strip()}",
                classroom_id.strip(),
                "browser",
                "failed",
                class_session_id=class_session_id,
                error_code="processing_failed",
                error_message=str(exc),
            )
        except Exception:
            logger.exception("Could not persist failed browser camera status")
        raise HTTPException(status.HTTP_503_SERVICE_UNAVAILABLE, f"Vision service could not analyse this frame: {exc}") from exc


@app.post("/api/model/train")
async def train_model(actor: User = Depends(platform_admin_only)):
    try:
        return await run_in_threadpool(vision.train_behaviour_model, actor.id)
    except ValueError as exc:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, str(exc)) from exc


@app.get("/api/model/versions")
def behaviour_model_versions(_: User = Depends(platform_admin_only), session: Session = Depends(db.get_db)):
    return db.list_behaviour_model_versions(session)


@app.post("/api/model/versions/{version_id}/promote")
async def promote_behaviour_model(version_id: int, payload: BehaviourModelPromotion, actor: User = Depends(platform_admin_only), session: Session = Depends(db.get_db)):
    row = session.get(BehaviourModelVersion, version_id)
    if not row:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Behaviour-model candidate not found")
    if row.status != "candidate":
        raise HTTPException(status.HTTP_409_CONFLICT, "Only a candidate can be approved or rejected")
    if not payload.approve:
        row.status = "rejected"
        db.write_audit(session, actor.id, "model.candidate_rejected", "behaviour_model_version", row.id, {"version": row.version, "reason": payload.reason})
        session.commit()
        return {"id": row.id, "version": row.version, "status": row.status}
    metrics = json.loads(row.metrics_json)
    class_report = metrics.get("classification_report", {})
    weakest_recall = min((float(class_report.get(label, {}).get("recall", 0)) for label in metrics.get("classes", [])), default=0)
    if row.macro_f1 < 0.70 or weakest_recall < 0.60:
        raise HTTPException(status.HTTP_409_CONFLICT, "Candidate is below the minimum macro-F1 or per-class recall gate")
    try:
        await run_in_threadpool(vision.activate_candidate, row.model_path, row.sha256)
    except ValueError as exc:
        raise HTTPException(status.HTTP_409_CONFLICT, str(exc)) from exc
    session.query(BehaviourModelVersion).filter(BehaviourModelVersion.status == "active", BehaviourModelVersion.id != row.id).update({BehaviourModelVersion.status: "superseded"}, synchronize_session=False)
    row.status = "active"
    row.promoted_by = actor.id
    row.promoted_at = datetime.now(timezone.utc)
    db.write_audit(session, actor.id, "model.candidate_promoted", "behaviour_model_version", row.id, {"version": row.version, "reason": payload.reason, "checksum": row.sha256})
    session.commit()
    db.update_model_meta(row.samples, row.accuracy, metrics)
    return {"id": row.id, "version": row.version, "status": row.status, "notice": "Promoted after metric gate. Continue monitored shadow/pilot review and retain rollback evidence."}


@app.patch("/api/incidents/{incident_id}/review")
def review_incident(incident_id: int, actor: User = Depends(teacher_only), session: Session = Depends(db.get_db)):
    incident = session.get(Incident, incident_id)
    if not incident or incident.institution_id != actor.institution_id:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Incident not found")
    incident.reviewed = True
    incident.reviewed_by = actor.id
    db.write_audit(session, actor.id, "incident.reviewed", "incident", incident.id)
    session.commit()
    return {"id": incident.id, "reviewed": True}


register_timetable_routes(app, current_user, allocation_manager, ensure_timetable_has_no_conflict, generate_timetable_sessions)
